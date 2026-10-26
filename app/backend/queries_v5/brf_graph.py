"""BRF studio as chat + flowchart (PLAN_R10 Session B, item 13).

A return is a graph the user can edit:

    source (domain.dataset) → transform (period, filter, join) → calc (Σ, count, ratio …) → check → output (template line)

`from_template(lines)` reads the lines of an uploaded BRF template and maps each one onto the Data Domains with
keyword rules (stage, past-due buckets, sector, collateral, restructured, coverage, NPL ratio, totals). Shared
intermediate calculations are built once and reused, so "Stage 3 coverage" draws its edges from the
"Stage 3 provisions" and "Stage 3 gross exposure" nodes.

`evaluate(graph, period)` walks the graph in topological order on the books at the period and returns a value,
display string and status per node. The user can add or remove nodes and edges (in the canvas or by chat) and the
outputs recompute. `chat(message, graph, template, period)` is the conversational box.

Pure functions; the API (backend.queries_v5.api) only wires them to HTTP.
"""
from __future__ import annotations

import io
import re
import uuid
from typing import Any, Callable

import pandas as pd

from backend.common import format as F
from backend.queries_v5 import engine as E

# ------------------------------------------------------------------ data domains the studio can read
FAC = "books.corp_facility_quarter"
BOR = "books.corp_borrower_quarter"
COL = "collateral.corp_collateral_quarter"

DATASETS: dict[str, dict[str, Any]] = {
    FAC: {"label": "corp_facility_quarter", "domain": "books", "load": lambda: E.facilities(),
          "measures": {"exposure": "ead_sar_mn", "ead": "ead_sar_mn", "provision": "ecl_sar_mn", "ecl": "ecl_sar_mn",
                       "drawn": "drawn_sar_mn", "limit": "limit_sar_mn", "undrawn": "undrawn_sar_mn",
                       "write-off": "write_off_sar_mn", "write off": "write_off_sar_mn", "recover": "recovery_sar_mn"}},
    BOR: {"label": "corp_borrower_quarter", "domain": "books", "load": lambda: E.borrowers(),
          "measures": {"debt": "total_debt_sar_mn", "revenue": "revenue_sar_mn", "ebitda": "ebitda_sar_mn"}},
    COL: {"label": "corp_collateral_quarter", "domain": "collateral", "load": lambda: E.collateral(),
          "measures": {"allocated": "allocated_value_sar_mn", "market": "market_value_sar_mn", "collateral": "allocated_value_sar_mn",
                       "value": "allocated_value_sar_mn"}},
}
KINDS = ("source", "transform", "calc", "check", "output")
AGG = {"sum", "count", "mean", "max"}
ARITH = {"ratio", "pct", "diff", "add"}
UNIT = {"const": "sar", "sum": "sar", "count": "count", "mean": "sar", "max": "sar", "ratio": "ratio", "pct": "pct", "diff": "sar", "add": "sar"}
_WHERE_OK = re.compile(r"^[\w\s.<>=!&|()'\"\-+,]*$")


def _sectors() -> list[str]:
    return sorted(str(s) for s in E.facilities()["sector"].dropna().unique())


def nid(prefix: str) -> str:
    return f"{prefix}:{uuid.uuid4().hex[:6]}"


def node(id_: str, kind: str, label: str, **spec: Any) -> dict[str, Any]:
    return {"id": id_, "kind": kind, "label": label, "spec": {k: v for k, v in spec.items() if v is not None}}


def edge(a: str, b: str) -> dict[str, Any]:
    return {"id": f"e:{a}>{b}", "source": a, "target": b}


# ------------------------------------------------------------------ template reading
def template_lines(name: str, data: bytes) -> list[str]:
    """The line labels of a BRF template: column A text of every sheet (xlsx), or one line per row (csv/txt)."""
    low = (name or "").lower()
    out: list[str] = []
    if low.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [c for c in row if c is not None and str(c).strip()]
                if cells and isinstance(cells[0], str):
                    out.append(cells[0].strip())
    else:
        text = data.decode("utf-8", errors="ignore")
        for ln in text.splitlines():
            first = ln.split(",")[0].split("\t")[0].strip().strip('"')
            if first:
                out.append(first)
    return [l for l in out if _is_line(l)][:80]


_HEADER = re.compile(r"^(line|item|description|code|value|amount|row|column|bank|return|period|template|reporting|prepared|notes?)\b", re.I)


def _is_line(label: str) -> bool:
    """Skip headers, titles and notes: a line names a measure (exposure, provisions, ratio …)."""
    if len(label) < 4 or len(label) > 120 or _HEADER.match(label):
        return False
    return bool(re.search(r"exposure|provision|ecl|coverage|ratio|past due|npl|non-performing|collateral|restructur|"
                          r"facilit|drawn|limit|write|recover|stage|total|count|number", label, re.I))


# ------------------------------------------------------------------ mapping lines → graph
class Builder:
    """Builds a graph, reusing identical sources, filters and calculations."""

    def __init__(self) -> None:
        self.nodes: list[dict[str, Any]] = []
        self.edges: list[dict[str, Any]] = []
        self._keys: dict[tuple, str] = {}

    def _get(self, key: tuple, make: Callable[[], dict[str, Any]], parents: list[str]) -> str:
        if key in self._keys:
            return self._keys[key]
        n = make()
        self.nodes.append(n)
        for p in parents:
            self.edges.append(edge(p, n["id"]))
        self._keys[key] = n["id"]
        return n["id"]

    def source(self, ds: str) -> str:
        return self._get(("src", ds), lambda: node(f"src:{DATASETS[ds]['label']}", "source", DATASETS[ds]["label"], dataset=ds), [])

    def period(self, ds: str) -> str:
        s = self.source(ds)
        return self._get(("period", ds), lambda: node(f"tx:period:{DATASETS[ds]['label']}", "transform", "Filter to period", op="period"), [s])

    def join_borrower(self) -> str:
        p = self.period(FAC)
        b = self.period(BOR)
        return self._get(("join", "bor"), lambda: node("tx:join:borrower", "transform", "Join facility → borrower", op="join", on="borrower_id",
                                                        fields=["restructured_flag", "watchlist_flag", "group_id"]), [p, b])

    def filt(self, base: str, where: str, label: str) -> str:
        if not where:
            return base
        return self._get(("filter", base, where), lambda: node(nid("tx"), "transform", label, op="filter", where=where), [base])

    def agg(self, frame: str, op: str, field: str | None, label: str) -> str:
        return self._get(("agg", frame, op, field), lambda: node(nid("calc"), "calc", label, op=op, field=field, unit=UNIT[op]), [frame])

    def arith(self, op: str, a: str, b: str, label: str) -> str:
        return self._get(("ar", op, a, b), lambda: node(nid("calc"), "calc", label, op=op, inputs=[a, b], unit=UNIT[op]), [a, b])

    def output(self, label: str, src: str | None, code: str, note: str = "") -> str:
        n = node(f"out:{code}", "output", label, cell=code, note=note or None)
        self.nodes.append(n)
        if src:
            self.edges.append(edge(src, n["id"]))
        return n["id"]

    def check(self, label: str, op: str, a: str, b: str | None = None, rule: str = "") -> str:
        cid = f"chk:{len([n for n in self.nodes if n['kind'] == 'check']) + 1}"
        self.nodes.append(node(cid, "check", label, op=op, inputs=[x for x in (a, b) if x], tol=0.01, rule=rule))
        self.edges.append(edge(a, cid))
        if b:
            self.edges.append(edge(b, cid))
        return cid

    def graph(self) -> dict[str, Any]:
        return {"nodes": self.nodes, "edges": self.edges}


def _filters(label: str) -> tuple[str, str]:
    """(pandas where, short words) for the label's filters."""
    low = label.lower()
    parts: list[tuple[str, str]] = []
    m = re.search(r"stage\s*([123])", low)
    if m:
        parts.append((f"stage == {m.group(1)}", f"Stage {m.group(1)}"))
    if re.search(r"non-?performing|\bnpl\b", low) and "ratio" not in low:
        parts.append(("stage == 3", "Stage 3"))
    if re.search(r"(over|more than|above|>)\s*90", low):
        parts.append(("dpd_days > 90", "DPD > 90"))
    elif re.search(r"31\s*[–\-to]+\s*90", low):
        parts.append(("dpd_days >= 31 and dpd_days <= 90", "DPD 31–90"))
    elif re.search(r"1\s*[–\-to]+\s*30", low):
        parts.append(("dpd_days >= 1 and dpd_days <= 30", "DPD 1–30"))
    elif re.search(r"past due", low) and not re.search(r"not past due", low):
        parts.append(("dpd_days > 0", "DPD > 0"))
    for s in _sectors():
        if s.lower() in low:
            parts.append((f"sector == '{s}'", s))
            break
    seen: set[str] = set()
    parts = [p for p in parts if not (p[0] in seen or seen.add(p[0]))]
    return " and ".join(p[0] for p in parts), " · ".join(p[1] for p in parts)


def _measure(label: str) -> tuple[str, str | None, str]:
    """(dataset, field, op) of the line's measure."""
    low = label.lower()
    if "collateral" in low:
        return COL, ("market_value_sar_mn" if "market" in low else "allocated_value_sar_mn"), "sum"
    if re.search(r"number of|count|facilities\b(?! exposure)", low) and "exposure" not in low:
        return FAC, None, "count"
    for word, field in DATASETS[FAC]["measures"].items():
        if word in low:
            return FAC, field, "sum"
    return FAC, "ead_sar_mn", "sum"


def map_line(b: Builder, label: str, code: str) -> dict[str, Any]:
    """Map one template line onto the domains; returns {code, label, mapped, how}."""
    low = label.lower()
    where, words = _filters(label)

    def total_of(field: str, nm: str) -> str:
        return b.agg(b.period(FAC), "sum", field, nm)

    def subset(field: str, w: str, ws: str, nm: str) -> str:
        return b.agg(b.filt(b.period(FAC), w, f"Filter {ws}" if ws else "All"), "sum", field, nm)

    if "coverage" in low:
        ecl = subset("ecl_sar_mn", where, words, f"{words} provisions".strip()) if where else total_of("ecl_sar_mn", "Total provisions")
        ead = subset("ead_sar_mn", where, words, f"{words} gross exposure".strip()) if where else total_of("ead_sar_mn", "Total gross exposure")
        src = b.arith("pct", ecl, ead, label)
        b.output(label, src, code)
        return {"code": code, "label": label, "mapped": True, "how": "Σ ECL ÷ Σ EAD" + (f" where {words}" if words else "")}
    if re.search(r"npl ratio|non-?performing.*ratio|ratio.*non-?performing", low):
        npl = subset("ead_sar_mn", "stage == 3", "Stage 3", "Stage 3 gross exposure")
        tot = total_of("ead_sar_mn", "Total gross exposure")
        src = b.arith("pct", npl, tot, label)
        b.output(label, src, code)
        return {"code": code, "label": label, "mapped": True, "how": "Stage 3 EAD ÷ total EAD"}
    if "restructur" in low:
        j = b.join_borrower()
        f = b.filt(j, "restructured_flag == 1", "Filter restructured")
        src = b.agg(f, "count" if re.search(r"number|count", low) else "sum", None if re.search(r"number|count", low) else "ead_sar_mn", label)
        b.output(label, src, code)
        return {"code": code, "label": label, "mapped": True, "how": "join borrower; restructured_flag = 1; Σ EAD"}
    ds, field, op = _measure(label)
    if ds == COL:
        src = b.agg(b.period(COL), "sum", field, label)
        b.output(label, src, code)
        return {"code": code, "label": label, "mapped": True, "how": f"Σ {field} on {DATASETS[COL]['label']}"}
    base = b.period(FAC)
    frame = b.filt(base, where, f"Filter {words}") if where else base
    nm = label if where or "total" not in low else label
    src = b.agg(frame, op, field, nm)
    b.output(label, src, code)
    return {"code": code, "label": label, "mapped": True, "how": f"{'count' if op == 'count' else 'Σ ' + str(field)}" + (f" where {words}" if words else "")}


def from_template(lines: list[str]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    """Graph for the template lines, plus the standard checks; returns (graph, mapping)."""
    b = Builder()
    mapping = [map_line(b, l, f"L{i + 1:02d}") for i, l in enumerate(lines)]
    _standard_checks(b, lines)
    return b.graph(), mapping


def _standard_checks(b: Builder, lines: list[str]) -> None:
    low = [l.lower() for l in lines]
    tot = b.agg(b.period(FAC), "sum", "ead_sar_mn", "Total gross exposure")
    stage_ead = [b._keys.get(("agg", b._keys.get(("filter", b._keys[("period", FAC)], f"stage == {s}")), "sum", "ead_sar_mn")) for s in (1, 2, 3)]
    if all(stage_ead):
        s12 = b.arith("add", stage_ead[0], stage_ead[1], "Stage 1 + 2 exposure")
        s123 = b.arith("add", s12, stage_ead[2], "Σ stage exposure")
        b.check("Stages add up to the total", "eq", s123, tot, "Σ Stage 1–3 EAD = total EAD")
    if any("provision" in l or "ecl" in l for l in low):
        stage_ecl = [b._keys.get(("agg", b._keys.get(("filter", b._keys[("period", FAC)], f"stage == {s}")), "sum", "ecl_sar_mn")) for s in (1, 2, 3)]
        tot_ecl = b.agg(b.period(FAC), "sum", "ecl_sar_mn", "Total provisions")
        if all(stage_ecl):
            e12 = b.arith("add", stage_ecl[0], stage_ecl[1], "Stage 1 + 2 provisions")
            e123 = b.arith("add", e12, stage_ecl[2], "Σ stage provisions")
            b.check("Stage provisions add up to ECL", "eq", e123, tot_ecl, "Σ Stage 1–3 ECL = book ECL")
    bad = b.agg(b.filt(b.period(FAC), "dpd_days > 90 and stage < 3", "Filter DPD > 90 outside Stage 3"), "count", None, "DPD > 90 outside Stage 3")
    b.check("DPD consistent with stage", "zero", bad, None, "no facility over 90 days past due outside Stage 3")


# ------------------------------------------------------------------ quick starts (the four prepared returns)
QUICK: dict[str, list[str]] = {
    "R1": ["Stage 1 gross exposure", "Stage 1 provisions (ECL)", "Stage 1 coverage", "Stage 2 gross exposure", "Stage 2 provisions (ECL)",
           "Stage 2 coverage", "Stage 3 gross exposure", "Stage 3 provisions (ECL)", "Stage 3 coverage", "Total gross exposure", "Total provisions"],
    "R2": ["Total gross exposure", "Total drawn balance", "Total limits", "Number of facilities", "Collateral allocated value"],
    "R3": ["Construction exposure", "Real Estate exposure", "Manufacturing exposure", "Wholesale Trade exposure", "Construction Stage 3 exposure",
           "Real Estate Stage 3 exposure", "Total gross exposure"],
    "R4": ["Past due 1–30 days exposure", "Past due 31–90 days exposure", "Past due over 90 days exposure", "Restructured exposure",
           "Non-performing (Stage 3) exposure", "NPL ratio", "Total gross exposure"],
}


# ------------------------------------------------------------------ evaluation
def _topo(graph: dict[str, Any]) -> list[str]:
    ids = [n["id"] for n in graph["nodes"]]
    idset = set(ids)
    indeg = {i: 0 for i in ids}
    kids: dict[str, list[str]] = {i: [] for i in ids}
    for e in graph["edges"]:
        if e["source"] in idset and e["target"] in idset:
            indeg[e["target"]] += 1
            kids[e["source"]].append(e["target"])
    q = [i for i in ids if indeg[i] == 0]
    out: list[str] = []
    while q:
        i = q.pop(0)
        out.append(i)
        for k in kids[i]:
            indeg[k] -= 1
            if indeg[k] == 0:
                q.append(k)
    return out + [i for i in ids if i not in out]   # cycles: evaluated last, will error


def _display(v: Any, unit: str) -> str:
    if v is None:
        return "—"
    return {"sar": F.sar, "count": F.count, "ratio": F.ratio}.get(unit, lambda x: F.pct(x, points=True))(v)


def _query(df: pd.DataFrame, where: str) -> pd.DataFrame:
    if not where.strip():
        return df
    if not _WHERE_OK.match(where) or "__" in where or "@" in where:
        raise ValueError("Filters use field names, numbers, quotes and comparisons only.")
    return df.query(where, engine="python")


def evaluate(graph: dict[str, Any], period: str | None = None) -> dict[str, Any]:
    period = period or E.as_of()
    byid = {n["id"]: n for n in graph["nodes"]}
    parents: dict[str, list[str]] = {i: [] for i in byid}
    for e in graph["edges"]:
        if e["source"] in byid and e["target"] in byid:
            parents[e["target"]].append(e["source"])
    frames: dict[str, pd.DataFrame] = {}
    vals: dict[str, dict[str, Any]] = {}

    def scalars(nid_: str) -> list[float]:
        order = byid[nid_]["spec"].get("inputs") or []
        ps = [p for p in order if p in parents[nid_]] + [p for p in parents[nid_] if p not in order]
        return [vals[p]["value"] for p in ps if p in vals and isinstance(vals[p].get("value"), (int, float))]

    for i in _topo(graph):
        n = byid[i]
        k, sp = n["kind"], n.get("spec") or {}
        try:
            if k == "source":
                ds = sp.get("dataset")
                if ds not in DATASETS:
                    raise ValueError(f"Unknown dataset {ds}.")
                frames[i] = DATASETS[ds]["load"]()
                vals[i] = {"rows": len(frames[i]), "display": f"{F.count(len(frames[i]))} rows"}
            elif k == "transform":
                ins = [frames[p] for p in parents[i] if p in frames]
                if not ins:
                    raise ValueError("Connect a source or another step to this node.")
                op = sp.get("op", "filter")
                if op == "period":
                    df = ins[0][ins[0]["reporting_quarter"] == period]
                elif op == "join":
                    on = sp.get("on", "borrower_id")
                    if len(ins) < 2:
                        raise ValueError("A join needs two inputs.")
                    right = ins[1]
                    keep = [on] + [c for c in (sp.get("fields") or []) if c in right.columns and c not in ins[0].columns]
                    df = ins[0].merge(right[keep].drop_duplicates(on), on=on, how="left")
                else:
                    df = _query(ins[0], str(sp.get("where", "")))
                frames[i] = df
                vals[i] = {"rows": len(df), "display": f"{F.count(len(df))} rows"}
            elif k == "calc":
                op = sp.get("op", "sum")
                unit = sp.get("unit") or UNIT.get(op, "sar")
                if op == "const":
                    v = float(sp.get("value", 0))
                elif op in AGG:
                    ins = [frames[p] for p in parents[i] if p in frames]
                    if not ins:
                        raise ValueError("Connect a dataset or filter to this calculation.")
                    df = _query(ins[0], str(sp.get("where", "")))
                    f = sp.get("field")
                    if op == "count":
                        v = float(len(df))
                    else:
                        if f not in df.columns:
                            raise ValueError(f"Field {f} is not in the input.")
                        v = float(getattr(df[f], op)()) if len(df) else 0.0
                else:
                    xs = scalars(i)
                    if len(xs) < 2:
                        raise ValueError("Connect two values to this calculation.")
                    if op in ("ratio", "pct"):
                        v = xs[0] / xs[1] * (100 if op == "pct" else 1) if xs[1] else None
                    elif op == "diff":
                        v = xs[0] - xs[1]
                    else:
                        v = float(sum(xs))
                vals[i] = {"value": v, "unit": unit, "display": _display(v, unit)}
            elif k == "check":
                xs = scalars(i)
                op, tol = sp.get("op", "eq"), float(sp.get("tol", 0.01))
                if op == "zero":
                    if not xs:
                        raise ValueError("Connect a value to this check.")
                    ok, det = abs(xs[0]) <= tol, f"{F.count(xs[0])} found"
                else:
                    if len(xs) < 2:
                        raise ValueError("Connect two values to this check.")
                    a, c = xs[0], xs[1]
                    ok = abs(a - c) <= tol if op == "eq" else (a <= c + tol if op == "le" else a + tol >= c)
                    u = next((vals[p].get("unit") for p in parents[i] if p in vals and vals[p].get("unit")), "sar")
                    det = f"{_display(a, u)} against {_display(c, u)}"
                vals[i] = {"status": "PASS" if ok else "FAIL", "display": "PASS" if ok else "FAIL", "detail": det}
            elif k == "output":
                xs = [vals[p] for p in parents[i] if p in vals and "value" in vals[p]]
                if not xs:
                    vals[i] = {"value": None, "display": "Not mapped", "status": "UNMAPPED"}
                else:
                    vals[i] = {"value": xs[0]["value"], "unit": xs[0]["unit"], "display": xs[0]["display"]}
        except Exception as e:  # noqa: BLE001 - every node reports its own error
            vals[i] = {"error": str(e)[:200], "display": "Error"}
    outs = [n for n in graph["nodes"] if n["kind"] == "output"]
    chks = [n for n in graph["nodes"] if n["kind"] == "check"]
    filled = sum(1 for n in outs if vals.get(n["id"], {}).get("value") is not None)
    passed = sum(1 for n in chks if vals.get(n["id"], {}).get("status") == "PASS")
    return {
        "period": period, "values": vals,
        "outputs": [{"id": n["id"], "code": n["spec"].get("cell", ""), "label": n["label"], **vals.get(n["id"], {})} for n in outs],
        "checks": [{"id": n["id"], "label": n["label"], "rule": n["spec"].get("rule", ""), **vals.get(n["id"], {})} for n in chks],
        "summary": f"**{filled}** of **{len(outs)}** lines filled at {period}; **{passed}** of **{len(chks)}** checks pass.",
        "powered_by": E.powered("BRF studio", "CBUAE credit returns"),
        "trace": E.trace(sorted({DATASETS[n['spec']['dataset']]['label'] for n in graph['nodes'] if n['kind'] == 'source' and n['spec'].get('dataset') in DATASETS}),
                         [n["label"] for n in graph["nodes"] if n["kind"] == "transform"], "BRF studio", "CBUAE credit returns",
                         int(len(E.book(period))), period),
    }


# ------------------------------------------------------------------ chat
HELP = ("Upload a BRF template and say “complete all data for this template”, or pick a quick start. Then edit the flow: "
        "“take collateral from the collateral dataset”, “add a reconciliation to the GL”, “add a check that Stage 3 coverage is at least 40%”, "
        "“remove Stage 1 coverage”, “filter Stage 2 gross exposure to sector Construction”, “lock 2026Q2”, “run”, “publish”.")


def _find(graph: dict[str, Any], text: str, kinds: tuple[str, ...] = KINDS) -> dict[str, Any] | None:
    t = text.lower().strip(" .“”\"'")
    cands = [n for n in graph["nodes"] if n["kind"] in kinds]
    for n in cands:
        if n["label"].lower() == t or n["id"].lower() == t:
            return n
    best = [n for n in cands if t and t in n["label"].lower()]
    order = {k: i for i, k in enumerate(("output", "calc", "check", "transform", "source"))}
    if best:
        return sorted(best, key=lambda n: (order[n["kind"]], len(n["label"])))[0]
    # word overlap on stems ("the reconciliation" finds "Reconciles to the GL")
    stop = {"the", "a", "an", "of", "to", "and", "node", "step", "line", "check", "for", "from", "in", "on"}
    def stems(x: str) -> set[str]:
        nums = {"one": "1", "two": "2", "three": "3"}
        return {nums.get(w, w)[:6] for w in re.findall(r"[a-z0-9]+", x.lower()) if w not in stop and (len(w) > 1 or w.isdigit())}
    want = stems(t)
    scored = [(len(want & stems(n["label"])), n) for n in cands]
    scored = [x for x in scored if x[0] > 0]
    if want and scored:
        top = max(x[0] for x in scored)
        if top >= max(1, (len(want) + 1) // 2):
            return sorted([n for k, n in scored if k == top], key=lambda n: (order[n["kind"]], len(n["label"])))[0]
    return None


def remove_node(graph: dict[str, Any], node_id: str) -> dict[str, Any]:
    return {"nodes": [n for n in graph["nodes"] if n["id"] != node_id],
            "edges": [e for e in graph["edges"] if node_id not in (e["source"], e["target"])]}


def _period_node(graph: dict[str, Any], ds: str) -> str | None:
    src = [n["id"] for n in graph["nodes"] if n["kind"] == "source" and n["spec"].get("dataset") == ds]
    for e in graph["edges"]:
        if e["source"] in src:
            t = next((n for n in graph["nodes"] if n["id"] == e["target"]), None)
            if t and t["kind"] == "transform" and t["spec"].get("op") == "period":
                return t["id"]
    return None


def _ensure_period(graph: dict[str, Any], ds: str) -> str:
    p = _period_node(graph, ds)
    if p:
        return p
    sid = f"src:{DATASETS[ds]['label']}"
    if not any(n["id"] == sid for n in graph["nodes"]):
        graph["nodes"].append(node(sid, "source", DATASETS[ds]["label"], dataset=ds))
    pid = f"tx:period:{DATASETS[ds]['label']}"
    graph["nodes"].append(node(pid, "transform", "Filter to period", op="period"))
    graph["edges"].append(edge(sid, pid))
    return pid


def _dataset_in(text: str) -> str | None:
    t = text.lower()
    if "collateral" in t:
        return COL
    if "borrower" in t:
        return BOR
    if "facilit" in t or "book" in t:
        return FAC
    return None


def _next_code(graph: dict[str, Any]) -> str:
    nums = [int(m.group(1)) for n in graph["nodes"] if n["kind"] == "output" for m in [re.match(r"L(\d+)", str(n["spec"].get("cell", "")))] if m]
    return f"L{(max(nums) if nums else 0) + 1:02d}"


def chat(message: str, graph: dict[str, Any] | None, template: list[str] | None, period: str | None) -> dict[str, Any]:
    """One turn of the BRF chat. Returns {reply, graph, template, action, mapping?}."""
    g = {"nodes": list((graph or {}).get("nodes") or []), "edges": list((graph or {}).get("edges") or [])}
    msg = (message or "").strip()
    low = msg.lower()
    out: dict[str, Any] = {"graph": g, "template": template or [], "action": "none"}

    m = re.search(r"\b(r[1-4])\b", low)
    if m and re.search(r"quick|start|open|prepared|load|use", low):
        rid = m.group(1).upper()
        g2, mp = from_template(QUICK[rid])
        return {**out, "graph": g2, "template": QUICK[rid], "mapping": mp, "action": "build",
                "reply": f"Opened quick start **{rid}**: {len(mp)} lines mapped to the Data Domains, with {sum(n['kind'] == 'check' for n in g2['nodes'])} checks. Edit the flow or Run."}

    if re.search(r"\b(complete|fill|populate|map)\b", low):
        if not template:
            return {**out, "reply": "Attach the BRF template first (the paperclip), then ask me to complete it."}
        g2, mp = from_template(template)
        ds = sorted({DATASETS[n["spec"]["dataset"]]["label"] for n in g2["nodes"] if n["kind"] == "source"})
        return {**out, "graph": g2, "mapping": mp, "action": "build",
                "reply": f"Read **{len(template)}** lines. Mapped **{sum(x['mapped'] for x in mp)}** of {len(mp)} to the Data Domains ({', '.join(ds)}), "
                         f"with {sum(n['kind'] == 'calc' for n in g2['nodes'])} intermediate calculations and {sum(n['kind'] == 'check' for n in g2['nodes'])} checks. "
                         "The flow is drawn below; edit any node, then Lock and Run."}

    if re.search(r"\b(remove|delete|drop)\b", low):
        target = re.sub(r"^.*?\b(remove|delete|drop)\b\s*(the\s+)?(node\s+|step\s+|check\s+|line\s+)?", "", msg, flags=re.I)
        n = (_find(g, re.sub(r"\bcheck\b", "", target, flags=re.I), ("check",)) if re.search(r"\bcheck\b", low) else None) or _find(g, target)
        if not n:
            return {**out, "reply": f"I could not find “{target}” in the flow."}
        return {**out, "graph": remove_node(g, n["id"]), "action": "edit", "reply": f"Removed **{n['label']}** ({n['kind']}) and its links. The outputs recompute."}

    if re.search(r"reconcil", low) and re.search(r"\bgl\b|general ledger|ledger", low):
        p = _ensure_period(g, FAC)
        gl = nid("calc")
        g["nodes"].append(node(gl, "calc", "GL control total (drawn balance)", op="sum", field="drawn_sar_mn", unit="sar"))
        g["edges"].append(edge(p, gl))
        ret = next((n for n in g["nodes"] if n["kind"] == "calc" and n["spec"].get("field") == "drawn_sar_mn" and n["id"] != gl), None)
        if not ret:
            ret = node(nid("calc"), "calc", "Return drawn balance", op="sum", field="drawn_sar_mn", unit="sar")
            g["nodes"].append(ret)
            g["edges"].append(edge(p, ret["id"]))
        cid = f"chk:gl{uuid.uuid4().hex[:4]}"
        g["nodes"].append(node(cid, "check", "Reconciles to the GL", op="eq", inputs=[ret["id"], gl], tol=0.01, rule="Return drawn balance = GL control total"))
        g["edges"] += [edge(ret["id"], cid), edge(gl, cid)]
        return {**out, "action": "edit", "reply": "Added a **reconciliation to the GL**: the return's drawn balance against the ledger control total (Σ drawn on the book at the period)."}

    m = re.search(r"\b(?:take|get|source|use)\s+(.+?)\s+from\s+(?:the\s+)?(.+?)(?:\s+dataset|\s+domain)?$", low)
    if m or re.search(r"\badd (a )?source\b", low):
        ds = _dataset_in(m.group(2) if m else low) or FAC
        what = (m.group(1) if m else ds.split(".")[-1]).strip()
        p = _ensure_period(g, ds)
        meas = DATASETS[ds]["measures"]
        field = next((f for w, f in meas.items() if w in what), next(iter(meas.values())))
        tgt = _find(g, what, ("output",))
        cid = nid("calc")
        g["nodes"].append(node(cid, "calc", f"Σ {what} ({DATASETS[ds]['label']})", op="sum", field=field, unit="sar"))
        g["edges"].append(edge(p, cid))
        if tgt:
            g["edges"] = [e for e in g["edges"] if e["target"] != tgt["id"]] + [edge(cid, tgt["id"])]
            where = f"line **{tgt['label']}** now reads it"
        else:
            code = _next_code(g)
            g["nodes"].append(node(f"out:{code}", "output", what[:1].upper() + what[1:], cell=code))
            g["edges"].append(edge(cid, f"out:{code}"))
            where = f"new line **{code}** added"
        return {**out, "action": "edit", "reply": f"Took **{what}** from **{DATASETS[ds]['label']}** (Σ {field}); {where}."}

    m = re.search(r"add (?:a )?check (?:that )?(.+?) (?:is )?(at least|at most|>=|<=|above|below|equals?|=)\s*([\d.]+)\s*(%|m|bn)?", low)
    if m:
        n = _find(g, m.group(1), ("output", "calc"))
        if not n:
            return {**out, "reply": f"I could not find “{m.group(1)}” in the flow."}
        src = n["id"]
        thr = float(m.group(3)) * (1000 if m.group(4) == "bn" else 1)
        k = nid("calc")
        g["nodes"].append(node(k, "calc", f"Threshold {m.group(3)}{m.group(4) or ''}", op="const", value=thr, unit="pct" if m.group(4) == "%" else "sar"))
        op = {"at least": "ge", ">=": "ge", "above": "ge", "at most": "le", "<=": "le", "below": "le"}.get(m.group(2), "eq")
        cid = f"chk:{uuid.uuid4().hex[:4]}"
        g["nodes"].append(node(cid, "check", f"{n['label']} {m.group(2)} {m.group(3)}{m.group(4) or ''}", op=op, inputs=[src, k], tol=0.0001, rule=f"{n['label']} {m.group(2)} {m.group(3)}{m.group(4) or ''}"))
        g["edges"] += [edge(src, cid), edge(k, cid)]
        return {**out, "action": "edit", "reply": f"Added a check: **{n['label']} {m.group(2)} {m.group(3)}{m.group(4) or ''}**."}

    m = re.search(r"filter\s+(.+?)\s+(?:to|by|on)\s+(?:sector\s+)?(.+)$", low)
    if m:
        n = _find(g, m.group(1), ("output", "calc"))
        if n and n["kind"] == "output":
            p = next((e["source"] for e in g["edges"] if e["target"] == n["id"]), None)
            n = next((x for x in g["nodes"] if x["id"] == p), None)
        if not n or n["kind"] != "calc" or n["spec"].get("op") not in AGG:
            return {**out, "reply": f"Name a calculation to filter, for example “filter Stage 2 gross exposure to sector Construction”."}
        sec = next((s for s in _sectors() if s.lower() in m.group(2)), None)
        cond = f"sector == '{sec}'" if sec else m.group(2).strip()
        old = n["spec"].get("where", "")
        n["spec"]["where"] = f"({old}) and {cond}" if old else cond
        return {**out, "action": "edit", "reply": f"**{n['label']}** now also filters `{cond}`."}

    m = re.search(r"\block\b(?:\s+(?:for\s+)?(\d{4}q[1-4]))?", low)
    if m:
        p = (m.group(1) or period or E.as_of()).upper()
        if p not in E.periods():
            return {**out, "reply": f"No book at {p}. Periods: {', '.join(E.periods()[-6:])}."}
        return {**out, "action": "lock", "period": p, "reply": f"Locked for **{p}**. Run to populate the return."}
    if re.fullmatch(r"(please\s+)?run( it| the return)?\.?", low) or low.startswith("run "):
        return {**out, "action": "run", "reply": "Running the return on the books…"}
    if re.search(r"\bpublish\b", low):
        return {**out, "action": "publish", "reply": "Publishing the return."}
    if re.search(r"\b(add|new)\b.*\b(node|step|output|line)\b", low):
        return {**out, "reply": "Use **Add node** above the flow to add a source, filter, calculation, check or output, then drag from its handle to wire it."}
    return {**out, "reply": HELP}
