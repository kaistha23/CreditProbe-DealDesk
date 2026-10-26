"""Queries and BRF studio API: its own app on port 8516 (PLAN_R9 Session F, P7 regulatory response).

Queries (prefix /api/v1/queries)
    GET  /health
    GET  /cases                              the five prepared letters
    GET  /items                              the tracker (Open / Drafted / Reviewed / Sent, owner, due date)
    POST /items/upload                       multipart file (.docx, .pdf, .xlsx) -> numbered questions, answers
    POST /items/prepared/{case}              open a prepared letter
    POST /ask                                {question, item_id?}: a Cockpit question (R10 Q1) -> the matching
                                             letter's answers, with the one that answers it first (422 when it
                                             is not an auditor or regulator query: the Cockpit asks it instead)
    GET  /items/{id}
    POST /items/{id}/answers/{n}             {draft}: edit an answer
    POST /items/{id}/status                  {status?, owner?, due?}
    GET  /items/{id}/export.docx | export.xlsx (the xlsx is "Fill template": the questions with their answers)

BRF studio (prefix /api/v1/brf)
    GET  /returns                            the four credit returns and their state
    GET  /returns/{rid}                      field map, process flow, checks, state
    POST /returns/{rid}/upload               multipart template: kept with the return; fields mapped
    POST /returns/{rid}/field                {id, filters?, calc?, label?}: edit a mapping (flow node)
    POST /returns/{rid}/lock                 {period}
    POST /returns/{rid}/run                  populated grid, check results, a note on movements
    POST /returns/{rid}/publish              {schedule}
    GET  /returns/{rid}/export.xlsx

BRF studio as chat + flowchart (PLAN_R10 Session B; prefix /api/v1/brf/studio)
    GET  /quick                              the four quick starts (prepared returns) and the seeded template
    POST /template                           multipart template -> its lines
    POST /chat                               multipart {message, graph (json), template (json), period, file?}
                                             -> {reply, graph, template, action, evaluation}
    POST /evaluate                           {graph, period} -> values per node, outputs, checks
    POST /publish                            {title, graph, period} -> kept in the studio log

Writes only var/queries_v5/state.json.
"""
from __future__ import annotations

import io
import json
import os
import pathlib
import re
import threading
import uuid
from datetime import date, datetime, timedelta, timezone
from typing import Any

from fastapi import Body, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response

from backend.queries_v5 import brf as B
from backend.queries_v5 import brf_graph as G
from backend.queries_v5 import engine as E

V5 = pathlib.Path(__file__).resolve().parents[2]
LETTERS = V5 / "demo_inputs" / "queries"
STATE = pathlib.Path(os.environ.get("QUERIES_STORE") or V5 / "var" / "queries_v5" / "state.json")
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ORIGINS = [o.strip() for o in (os.environ.get("QUERIES_CORS_ORIGINS") or "http://127.0.0.1:5500,http://localhost:5500").split(",") if o.strip()]
STATUSES = ["Open", "Drafted", "Reviewed", "Sent"]

app = FastAPI(title="CreditProbe v5 Queries and BRF", version="queries-v5-1.0.0")
app.add_middleware(CORSMiddleware, allow_origins=ORIGINS, allow_credentials=True, allow_methods=["GET", "POST", "OPTIONS"],
                   allow_headers=["*"], expose_headers=["Content-Disposition"])
_LOCK = threading.Lock()


def _load() -> dict[str, Any]:
    if STATE.exists():
        try:
            return json.loads(STATE.read_text())
        except Exception:  # noqa: BLE001 - a damaged store starts afresh
            pass
    return {"items": {}, "brf": {}}


def _save(s: dict[str, Any]) -> None:
    STATE.parent.mkdir(parents=True, exist_ok=True)
    tmp = STATE.with_suffix(".tmp")
    tmp.write_text(json.dumps(s, default=str))
    tmp.replace(STATE)


def _now() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


# ------------------------------------------------------------------ letter parsing
def letter_text(name: str, data: bytes) -> str:
    low = name.lower()
    if low.endswith(".docx"):
        from docx import Document
        d = Document(io.BytesIO(data))
        return "\n".join(p.text for p in d.paragraphs)
    if low.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if low.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c not in (None, "")]
                if cells:
                    lines.append(" ".join(cells))
        return "\n".join(lines)
    if low.endswith(".txt"):
        return data.decode("utf-8", "ignore")
    raise HTTPException(415, "Upload a letter as .docx or .pdf, or an Excel template (.xlsx).")


def questions(text: str) -> list[str]:
    out = []
    for line in text.splitlines():
        m = re.match(r"^\s*(?:Q\s*)?(\d{1,2})[.)]\s+(.{12,})$", line.strip())
        if m:
            out.append(m.group(2).strip())
    return out


def meta(text: str) -> dict[str, str]:
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    ref = next((l for l in lines if l.upper().startswith("REF")), "")
    subj = next((l.split(":", 1)[1].strip() for l in lines if l.lower().startswith("subject:")), "")
    return {"sender": lines[0] if lines else "", "reference": ref, "subject": subj}


def _create(name: str, data: bytes) -> dict[str, Any]:
    text = letter_text(name, data)
    qs = questions(text)
    case = E.classify(text)
    if case is None:
        raise HTTPException(422, "The letter's questions do not match a P7 response template yet. Prepared cases cover staging and "
                                 "rebuttals, Stage 3 floors, ECL movement, the 1.5% CRWA floor and restructuring.")
    if not qs:
        qs = [""] * 3
    answers = E.run_case(case, qs)
    for a in answers:
        if not a["question"]:
            a["question"] = E.CASES[case]["title"]
        a["status"] = "Drafted"
    m = meta(text)
    item = {"id": uuid.uuid4().hex[:10], "file": name, "case": case, "title": E.CASES[case]["title"], "from": E.CASES[case]["from"],
            "article": E.CASES[case]["article"], **m, "status": "Drafted", "owner": "Anne", "due": (date.today() + timedelta(days=14)).isoformat(),
            "received": date.today().isoformat(), "created_at": _now(), "period": E.as_of(), "answers": answers,
            "log": [{"at": _now(), "text": f"Read {len(qs)} questions from {name}; matched to {E.CASES[case]['title']}; drafted {len(answers)} answers."}]}
    with _LOCK:
        s = _load()
        s["items"][item["id"]] = item
        _save(s)
    return item


def _card(i: dict[str, Any]) -> dict[str, Any]:
    return {k: i.get(k) for k in ("id", "file", "case", "title", "from", "article", "reference", "subject", "status", "owner", "due", "received")} | {
        "questions": len(i.get("answers") or [])}


def _item(qid: str) -> dict[str, Any]:
    i = _load()["items"].get(qid)
    if not i:
        raise HTTPException(404, "That query does not exist.")
    return i


# ------------------------------------------------------------------ queries routes
Q = "/api/v1/queries"


@app.get(f"{Q}/health")
def health() -> dict[str, Any]:
    return {"status": "ok", "service": "queries_v5", "as_of": E.as_of()}


@app.get(f"{Q}/cases")
def cases() -> list[dict[str, Any]]:
    files = sorted(p.name for p in LETTERS.glob("*.docx")) if LETTERS.exists() else []
    out = []
    for k, c in E.CASES.items():
        f = next((x for x in files if x.startswith(k + "_")), None)
        out.append({"case": k, "title": c["title"], "from": c["from"], "article": c["article"], "file": f})
    return out


@app.get(f"{Q}/items")
def items() -> list[dict[str, Any]]:
    return sorted((_card(i) for i in _load()["items"].values()), key=lambda x: x["received"] or "", reverse=True)


@app.post(f"{Q}/items/upload")
async def upload(file: UploadFile = File(...)) -> dict[str, Any]:
    data = await file.read()
    if len(data) > 15_000_000:
        raise HTTPException(413, "The file is larger than 15 MB.")
    return _create(file.filename or "letter.docx", data)


@app.post(f"{Q}/items/prepared/{{case}}")
def prepared(case: str) -> dict[str, Any]:
    f = next((p for p in LETTERS.glob(f"{case}_*.docx")), None)
    if f is None:
        raise HTTPException(404, "That prepared letter is not on file.")
    return _create(f.name, f.read_bytes())


# ------------------------------------------------------------------ R10 Q1: queries from the Cockpit chat
#: Words that make a Cockpit question an auditor or regulator query rather than a portfolio question. A plain
#: "Why did construction ECL rise?" stays with the Cockpit's own agent; "Answer the auditor's question on ..." comes here.
QUERY_CUES = ("auditor", "audit ", "internal audit", "cbuae", "central bank", "regulator", "regulatory query",
              "letter", "query", "queries", "response pack", "draft answer", "draft a response", "respond to", "reply to")
#: Free-text cues per prepared case, on top of the letter keywords the engine classifies uploads with.
ASK_KEYS = {
    "Q1": ["30+ dpd", "30 dpd", "30 days", "past due", "still in stage 1", "stage 1", "rebuttal", "7.8", "7.10"],
    "Q2": ["stage 3", "top 20", "floor", "9.23", "secured", "unsecured", "time in stage 3"],
    "Q3": ["construction", "stage 2 ecl", "ecl movement", "movement", "waterfall", "increase"],
    "Q4": ["crwa", "1.5%", "risk weighted", "risk-weighted", "9.21", "general provision", "stage 1 + 2", "stage 1 and 2"],
    "Q5": ["restructur", "probation", "forborne", "forbearance", "9.28", "9.35", "art 8"],
}
_CASE_REF = re.compile(r"\b(?:letter|query|case)\s*(?:q)?([1-5])\b|\bq([1-5])\b", re.I)
_STOP = {"the", "a", "an", "of", "on", "in", "for", "and", "or", "to", "is", "are", "with", "by", "its", "their", "that",
         "which", "what", "how", "please", "answer", "question", "auditor", "auditors", "auditor's", "cbuae", "letter",
         "still", "loans", "facilities", "facility", "list", "provide", "each", "all", "we", "our", "us", "this", "be"}


def is_query(question: str) -> bool:
    """Is this Cockpit question an auditor or regulator query (or a reference to a prepared letter)?"""
    t = f" {question.lower()} "
    return bool(_CASE_REF.search(t)) or any(c in t for c in QUERY_CUES)


def case_for(question: str) -> str | None:
    """The prepared case a free-text question is about: an explicit "letter 3" / "Q3" wins, then the cues."""
    t = question.lower()
    m = _CASE_REF.search(t)
    if m:
        return f"Q{m.group(1) or m.group(2)}"
    scores = {k: sum(t.count(w) for w in ASK_KEYS[k]) + sum(t.count(w) for w in E.CASES[k]["keys"]) for k in E.CASES}
    best = max(scores, key=lambda k: scores[k])
    return best if scores[best] > 0 else None


def _tokens(text: str) -> set[str]:
    return {w for w in re.findall(r"[a-z0-9.+%]+", text.lower()) if w not in _STOP and len(w) > 1}


def best_answer(question: str, answers: list[dict[str, Any]]) -> int:
    """The number of the letter's question that the user's question asks: most shared words, ties to the first.
    "question 3" / "Q1.3" names it outright."""
    m = re.search(r"\bquestion\s*(\d{1,2})\b|\bq[1-5]\.(\d{1,2})\b", question.lower())
    if m:
        n = int(m.group(1) or m.group(2))
        if any(a["n"] == n for a in answers):
            return n
    mine = _tokens(question)
    scored = [(len(mine & _tokens(a["question"])), -a["n"], a["n"]) for a in answers]
    return max(scored)[2] if scored else 1


@app.post(f"{Q}/ask")
def ask(body: dict[str, Any] = Body(...)) -> dict[str, Any]:
    """A question typed in the Cockpit (R10 Q1). Not an auditor or regulator query -> 422 and the Cockpit's agent
    answers it as before. Otherwise the matching letter is answered on the books (reusing the one already answered
    in this Cockpit conversation, `item_id`, or the latest open one for that letter) and `focus` names the answer
    that answers the question, so the Cockpit can lead with it."""
    question = str(body.get("question") or "").strip()[:2000]
    if not question or not is_query(question):
        raise HTTPException(422, "Not an auditor or regulator query.")
    case = case_for(question)
    if case is None:
        raise HTTPException(422, "The question does not match a P7 response template. Prepared letters cover staging "
                                 "and rebuttals, Stage 3 floors, construction ECL movement, the 1.5% CRWA floor and "
                                 "restructuring; attach the letter to answer another.")
    items = _load()["items"]
    item = items.get(str(body.get("item_id") or ""))
    if not item or item.get("case") != case:
        item = next((i for i in sorted(items.values(), key=lambda x: x.get("created_at") or "", reverse=True)
                     if i.get("case") == case and i.get("status") != "Sent"), None)
    if item is None:
        f = next((p for p in LETTERS.glob(f"{case}_*.docx")), None)
        if f is None:
            raise HTTPException(404, "That prepared letter is not on file.")
        item = _create(f.name, f.read_bytes())
    n = best_answer(question, item["answers"])
    return {"question": question, "case": case, "focus": n, "item": item}


@app.get(f"{Q}/items/{{qid}}")
def get_item(qid: str) -> dict[str, Any]:
    return _item(qid)


@app.post(f"{Q}/items/{{qid}}/answers/{{n}}")
def edit_answer(qid: str, n: int, body: dict[str, Any] = Body(...)) -> dict[str, Any]:
    with _LOCK:
        s = _load()
        i = s["items"].get(qid)
        if not i:
            raise HTTPException(404, "That query does not exist.")
        a = next((x for x in i["answers"] if x["n"] == n), None)
        if a is None:
            raise HTTPException(404, "That question does not exist.")
        if "draft" in body:
            a["draft"] = str(body["draft"])[:20000]
            a["edited"] = True
        if body.get("status") in STATUSES:
            a["status"] = body["status"]
        i["log"].append({"at": _now(), "text": f"Answer {n} edited."})
        _save(s)
        return i


@app.post(f"{Q}/items/{{qid}}/status")
def set_status(qid: str, body: dict[str, Any] = Body(...)) -> dict[str, Any]:
    with _LOCK:
        s = _load()
        i = s["items"].get(qid)
        if not i:
            raise HTTPException(404, "That query does not exist.")
        if "status" in body:
            if body["status"] not in STATUSES:
                raise HTTPException(422, f"Status is one of {', '.join(STATUSES)}.")
            i["status"] = body["status"]
        for k in ("owner", "due"):
            if k in body:
                i[k] = str(body[k])[:80]
        i["log"].append({"at": _now(), "text": "Tracker updated: " + ", ".join(f"{k} {body[k]}" for k in ("status", "owner", "due") if k in body)})
        _save(s)
        return i


def _plain(md: str) -> str:
    return md.replace("**", "")


def _meta(i: dict[str, Any], title: str):
    from backend.house_style import Meta
    return Meta(title, book="corporate", period=i.get("period", ""), subject=i.get("subject") or i["title"], reference=i.get("reference", ""))


@app.get(f"{Q}/items/{{qid}}/export.docx")
def export_docx(qid: str) -> Response:
    from backend.house_style import docx as hd
    i = _item(qid)
    doc = hd.document(_meta(i, f"Response: {i['title']}"), kind="Regulatory response", subtitle=i.get("sender", ""),
                      details=[("Reference", i.get("reference", "")), ("From", i["from"]), ("Article", i["article"]), ("Status", i["status"])])
    for a in i["answers"]:
        hd.heading(doc, f"{a['n']}. {a['question']}", 2)
        doc.add_paragraph(_plain(a["draft"]))
        t = a.get("table")
        if t and t.get("rows"):
            cols = t["columns"]
            rows = [[r.get(f"{c['key']}__d", "") for c in cols] for r in t["rows"][:40]]
            if t.get("total"):
                rows.append([t["total"].get(f"{c['key']}__d", "") for c in cols])
            hd.table(doc, [c["label"] for c in cols], rows)
        tr = a["trace"]
        doc.add_paragraph(f"Basis: {tr['article']} · sources {', '.join(tr['sources'])} · period {tr['period']} · {a['powered_by']['pack']}").runs[0].italic = True
    buf = io.BytesIO()
    doc.save(buf)
    name = f"Response_{i['case']}_{i['id']}.docx"
    return Response(buf.getvalue(), media_type=DOCX, headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.get(f"{Q}/items/{{qid}}/export.xlsx")
def export_xlsx(qid: str) -> Response:
    from backend.house_style import xlsx as hx
    i = _item(qid)
    m = _meta(i, f"Response: {i['title']}")
    wb = hx.workbook(m)
    hx.table_sheet(wb, "Responses", m, [hx.Col("n", "No."), hx.Col("question", "Question"), hx.Col("answer", "Answer"), hx.Col("article", "Basis"), hx.Col("status", "Status")],
                   [{"n": a["n"], "question": a["question"], "answer": _plain(a["draft"]), "article": a["trace"]["article"], "status": a.get("status", "")} for a in i["answers"]])
    for a in i["answers"]:
        t = a.get("table")
        if t and t.get("rows"):
            cols = [hx.Col(c["key"], c["label"], "sar_mn" if c["kind"] == "sar" else ("pct" if c["kind"] == "pct" else ("count" if c["kind"] == "count" else "text"))) for c in t["columns"]]
            rows = [{c["key"]: r.get(c["key"]) for c in t["columns"]} for r in t["rows"]]
            hx.table_sheet(wb, f"Q{a['n']} schedule", m, cols, rows, title=a["question"][:120])
    hx.about_sheet(wb, m, [("From", i["from"]), ("Basis", i["article"]), ("Status", i["status"]), ("Powered by", E.PACK)])
    buf = io.BytesIO()
    wb.save(buf)
    name = f"Response_{i['case']}_{i['id']}.xlsx"
    return Response(buf.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ------------------------------------------------------------------ BRF routes
R = "/api/v1/brf"


def _brf_state(s: dict[str, Any], rid: str) -> dict[str, Any]:
    st = s["brf"].setdefault(rid, {"state": "Draft", "period": None, "template": None, "edits": {}, "schedule": None, "published_at": None, "log": []})
    return st


def _fields(rid: str, st: dict[str, Any]) -> list[dict[str, Any]]:
    out = []
    for f in B.field_map(rid):
        e = st["edits"].get(f["id"]) or {}
        out.append({**f, **{k: v for k, v in e.items() if k in ("filters", "calc", "label")}, "edited": bool(e)})
    return out


def _check_rid(rid: str) -> None:
    if rid not in B.RETURNS:
        raise HTTPException(404, "That return does not exist.")


@app.get(f"{R}/health")
def brf_health() -> dict[str, Any]:
    return {"status": "ok", "service": "brf", "as_of": E.as_of()}


@app.get(f"{R}/returns")
def returns() -> list[dict[str, Any]]:
    s = _load()
    return [{"id": rid, **r, "state": _brf_state(s, rid)["state"], "period": _brf_state(s, rid)["period"],
             "fields": len(B.field_map(rid)), "schedule": _brf_state(s, rid)["schedule"]} for rid, r in B.RETURNS.items()]


@app.get(f"{R}/returns/{{rid}}")
def get_return(rid: str) -> dict[str, Any]:
    _check_rid(rid)
    s = _load()
    st = _brf_state(s, rid)
    fl = B.flow(rid)
    for n in fl["nodes"]:
        if n["kind"] == "calc":
            f = next(x for x in _fields(rid, st) if f"calc:{x['id']}" == n["id"])
            n["detail"] = f"{f['calc']} where {f['filters'] or '—'}"
            n["label"] = f["label"]
    return {"id": rid, **B.RETURNS[rid], **st, "fields": _fields(rid, st), "flow": fl, "checks": B.checks(rid), "periods": E.periods()[-6:], "as_of": E.as_of()}


@app.post(f"{R}/returns/{{rid}}/upload")
async def brf_upload(rid: str, file: UploadFile = File(...)) -> dict[str, Any]:
    _check_rid(rid)
    data = await file.read()
    labels: list[str] = []
    if (file.filename or "").lower().endswith(".xlsx"):
        labels = [l for l in letter_text(file.filename, data).splitlines() if l.strip()][:60]
    with _LOCK:
        s = _load()
        st = _brf_state(s, rid)
        st["template"] = {"file": file.filename, "lines": len(labels), "at": _now()}
        st["log"].append({"at": _now(), "text": f"Template {file.filename} read: {len(B.field_map(rid))} fields mapped to existing domains."})
        _save(s)
    return get_return(rid)


@app.post(f"{R}/returns/{{rid}}/field")
def brf_field(rid: str, body: dict[str, Any] = Body(...)) -> dict[str, Any]:
    _check_rid(rid)
    fid = str(body.get("id", ""))
    if fid not in {f["id"] for f in B.field_map(rid)}:
        raise HTTPException(404, "That field is not in this return.")
    with _LOCK:
        s = _load()
        st = _brf_state(s, rid)
        if st["state"] in ("Locked", "Run", "Published"):
            st["state"] = "Draft"
        st["edits"][fid] = {k: str(body[k])[:300] for k in ("filters", "calc", "label") if k in body}
        st["log"].append({"at": _now(), "text": f"Field {fid} edited."})
        _save(s)
    return get_return(rid)


@app.post(f"{R}/returns/{{rid}}/lock")
def brf_lock(rid: str, body: dict[str, Any] = Body(default={})) -> dict[str, Any]:
    _check_rid(rid)
    period = str(body.get("period") or E.as_of())
    if period not in E.periods():
        raise HTTPException(422, f"No book at {period}.")
    with _LOCK:
        s = _load()
        st = _brf_state(s, rid)
        st.update(state="Locked", period=period)
        st["log"].append({"at": _now(), "text": f"Locked for {period}."})
        _save(s)
    return get_return(rid)


@app.post(f"{R}/returns/{{rid}}/run")
def brf_run(rid: str) -> dict[str, Any]:
    _check_rid(rid)
    s = _load()
    st = _brf_state(s, rid)
    if not st.get("period"):
        raise HTTPException(409, "Lock the return for a period first.")
    out = B.run(rid, st["period"])
    with _LOCK:
        s = _load()
        st = _brf_state(s, rid)
        st["state"] = "Run"
        st["last_run"] = {"at": _now(), "checks": [c["status"] for c in out["checks"]]}
        st["log"].append({"at": _now(), "text": f"Run for {st['period']}: {sum(c['status'] == 'PASS' for c in out['checks'])} of {len(out['checks'])} checks pass."})
        _save(s)
    return out


@app.post(f"{R}/returns/{{rid}}/publish")
def brf_publish(rid: str, body: dict[str, Any] = Body(default={})) -> dict[str, Any]:
    _check_rid(rid)
    with _LOCK:
        s = _load()
        st = _brf_state(s, rid)
        if st["state"] != "Run":
            raise HTTPException(409, "Run the return before publishing it.")
        st.update(state="Published", schedule=str(body.get("schedule") or "Quarterly, 15 working days after quarter-end"), published_at=_now())
        st["log"].append({"at": _now(), "text": f"Published; schedule {st['schedule']}."})
        _save(s)
    return get_return(rid)


@app.get(f"{R}/returns/{{rid}}/export.xlsx")
def brf_export(rid: str) -> Response:
    _check_rid(rid)
    from backend.house_style import Meta
    from backend.house_style import xlsx as hx
    st = _brf_state(_load(), rid)
    out = B.run(rid, st.get("period") or E.as_of())
    m = Meta(f"{rid} {out['title']}", book="corporate", period=out["period"])
    wb = hx.workbook(m)
    g = out["grid"]
    kind = {"sar": "sar_mn", "pct": "pct", "count": "count"}
    cols = [hx.Col(c["key"], c["label"], kind.get(c["kind"], "text")) for c in g["columns"]]
    rows = [{c["key"]: r.get(c["key"]) for c in g["columns"]}
            for r in g["rows"] + ([g["total"]] if g.get("total") else [])]
    hx.table_sheet(wb, "Return", m, cols, rows, note=out["note"].replace("**", ""))
    hx.table_sheet(wb, "Checks", m, [hx.Col("label", "Check"), hx.Col("rule", "Rule"), hx.Col("status", "Result"), hx.Col("detail", "Detail")], out["checks"])
    hx.about_sheet(wb, m, [("Return", out["title"]), ("Basis", B.RETURNS[rid]["article"]), ("Powered by", E.PACK)])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="BRF_{rid}_{out["period"]}.xlsx"'})


# ------------------------------------------------------------------ BRF studio: chat + flowchart (PLAN_R10 Session B)
S = "/api/v1/brf/studio"
BRF_TEMPLATES = V5 / "demo_inputs" / "brf"


def _graph_in(raw: Any) -> dict[str, Any]:
    g = json.loads(raw) if isinstance(raw, str) and raw.strip() else (raw or {})
    if not isinstance(g, dict):
        raise HTTPException(422, "The flow must be an object with nodes and edges.")
    nodes = [n for n in (g.get("nodes") or []) if isinstance(n, dict) and n.get("id") and n.get("kind") in G.KINDS][:400]
    for n in nodes:
        n.setdefault("label", n["id"])
        n["spec"] = n.get("spec") if isinstance(n.get("spec"), dict) else {}
    edges = [{"id": e.get("id") or f"e:{e['source']}>{e['target']}", "source": e["source"], "target": e["target"]}
             for e in (g.get("edges") or []) if isinstance(e, dict) and e.get("source") and e.get("target")][:1000]
    return {"nodes": nodes, "edges": edges}


def _period_ok(period: str | None) -> str:
    p = (period or "").strip() or E.as_of()
    if p not in E.periods():
        raise HTTPException(422, f"No book at {p}.")
    return p


@app.get(f"{S}/quick")
def studio_quick() -> dict[str, Any]:
    files = sorted(p.name for p in BRF_TEMPLATES.glob("*.xlsx")) if BRF_TEMPLATES.exists() else []
    return {"quick": [{"id": rid, "title": B.RETURNS[rid]["title"], "lines": len(G.QUICK[rid])} for rid in G.QUICK],
            "templates": files, "periods": E.periods()[-6:], "as_of": E.as_of(), "help": G.HELP}


@app.get(f"{S}/templates/{{name}}")
def studio_template_file(name: str) -> Response:
    p = BRF_TEMPLATES / pathlib.Path(name).name
    if not p.exists() or p.suffix.lower() != ".xlsx":
        raise HTTPException(404, "That template does not exist.")
    return Response(p.read_bytes(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{p.name}"'})


@app.post(f"{S}/template")
async def studio_template(file: UploadFile = File(...)) -> dict[str, Any]:
    data = await file.read()
    lines = G.template_lines(file.filename or "", data)
    if not lines:
        raise HTTPException(422, "No return lines found in that file. Put one line label per row in the first column.")
    return {"file": file.filename, "lines": lines}


@app.post(f"{S}/chat")
async def studio_chat(message: str = Form(""), graph: str = Form(""), template: str = Form(""), period: str = Form(""),
                      file: UploadFile | None = File(None)) -> dict[str, Any]:
    g = _graph_in(graph)
    tpl = json.loads(template) if template.strip() else []
    tpl = [str(x)[:120] for x in tpl if str(x).strip()][:80] if isinstance(tpl, list) else []
    attached = None
    if file is not None and file.filename:
        lines = G.template_lines(file.filename, await file.read())
        if not lines:
            raise HTTPException(422, "No return lines found in that file. Put one line label per row in the first column.")
        tpl, attached = lines, file.filename
    p = _period_ok(period)
    out = G.chat(message, g, tpl, p)
    if attached and out["action"] == "none" and not re.search(r"\w", message or ""):
        out["reply"] = f"Read **{attached}**: {len(tpl)} lines. Say “complete all data for this template” to map and fill them."
    elif attached and out["action"] == "none":
        out = G.chat("complete all data", g, tpl, p) if not re.search(r"remove|delete|lock|publish|reconcil|take|check|filter", message.lower()) else out
    out["template"] = out.get("template") or tpl
    out["file"] = attached
    p = out.get("period") or p
    out["period"] = p
    out["evaluation"] = G.evaluate(out["graph"], p) if out["graph"]["nodes"] else None
    return out


@app.post(f"{S}/evaluate")
def studio_evaluate(body: dict[str, Any] = Body(...)) -> dict[str, Any]:
    return G.evaluate(_graph_in(body.get("graph")), _period_ok(body.get("period")))


@app.post(f"{S}/publish")
def studio_publish(body: dict[str, Any] = Body(...)) -> dict[str, Any]:
    g = _graph_in(body.get("graph"))
    p = _period_ok(body.get("period"))
    ev = G.evaluate(g, p)
    rec = {"id": uuid.uuid4().hex[:8], "title": str(body.get("title") or "BRF return")[:120], "period": p, "at": _now(),
           "lines": len(ev["outputs"]), "checks_pass": sum(c.get("status") == "PASS" for c in ev["checks"]), "checks": len(ev["checks"]),
           "graph": g}
    with _LOCK:
        s = _load()
        s.setdefault("studio", []).append(rec)
        s["studio"] = s["studio"][-50:]
        _save(s)
    return {k: v for k, v in rec.items() if k != "graph"} | {"summary": ev["summary"]}


@app.post(f"{S}/export.xlsx")
def studio_export(body: dict[str, Any] = Body(...)) -> Response:
    """The populated return (lines, checks) as a house-style workbook."""
    from backend.house_style import Meta
    from backend.house_style import xlsx as hx
    g = _graph_in(body.get("graph"))
    p = _period_ok(body.get("period"))
    ev = G.evaluate(g, p)
    title = str(body.get("title") or "BRF return")[:80]
    m = Meta(title, book="corporate", period=p)
    wb = hx.workbook(m)
    rows = [{"code": o["code"], "line": o["label"], "value": o.get("value"), "shown": o.get("display", "")} for o in ev["outputs"]]
    hx.table_sheet(wb, "Return", m, [hx.Col("code", "Code"), hx.Col("line", "Line"), hx.Col("value", "Value", "num"), hx.Col("shown", "As shown")], rows,
                   note=ev["summary"].replace("**", ""))
    hx.table_sheet(wb, "Checks", m, [hx.Col("label", "Check"), hx.Col("rule", "Rule"), hx.Col("status", "Result"), hx.Col("detail", "Detail")],
                   [{k: c.get(k, c.get("error", "")) for k in ("label", "rule", "status", "detail")} for c in ev["checks"]])
    hx.about_sheet(wb, m, [("Return", title), ("Basis", "Template lines mapped to the Data Domains"), ("Powered by", E.PACK)])
    buf = io.BytesIO()
    wb.save(buf)
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", title).strip("_") or "BRF"
    return Response(buf.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{safe}_{p}.xlsx"'})
