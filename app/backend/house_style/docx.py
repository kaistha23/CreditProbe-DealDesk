"""House-style Word documents with python-docx.

    from backend.house_style import Meta, docx as hd
    doc = hd.document(meta, kind="Credit committee pack", subtitle="Q2 2026")   # cover page included
    hd.heading(doc, "1. Executive summary")
    hd.para(doc, "ECL rose **SAR 965m** in the quarter.")
    hd.bullets(doc, ["one", "two"])
    hd.table(doc, ["Sector", "ECL (SAR m)"], [["Construction", 2755.0]], kinds=["", "sar_mn"])
    hd.figure(doc, png_bytes, caption="Figure 1. ECL by sector")
    hd.page_break(doc)
    data = hd.to_bytes(doc)

The cover page carries the brand line, the document kind, the title, the subtitle and a details table
(book, period, prepared by, generated, reference, classification). From page 2 every page has a running
header (title · book · period, classification) and a footer ("CreditProbe · <generated line> · Page X of
Y"). Headings, body text and tables use the house styles: Heading 1-3 in the brand colour, the "CP
Table" table style with a brand header row, banded rows and right-aligned numbers.
"""
from __future__ import annotations

import io
from collections.abc import Sequence
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from backend.house_style import ACCENT, BAND, FONT, INK, MUTED, RULE, STATUS, STYLE_ID, WHITE, Meta

TABLE_STYLE = "CP Table"
COVER_STYLE = "CP Cover Title"


def _rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


def _field(run, code: str) -> None:
    for kind, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif text == "1":
            t = OxmlElement("w:t")
            t.text = "1"
            run._r.append(t)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {code} "
            run._r.append(el)


def shade(cell, hex6: str) -> None:
    tc = cell._tc.get_or_add_tcPr()
    for old in tc.findall(qn("w:shd")):
        tc.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6)
    tc.append(shd)


def _bottom_rule(paragraph, colour: str = ACCENT, size: int = 12) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), colour)
    bdr.append(b)
    ppr.append(bdr)


def _styles(doc) -> None:
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = _rgb(INK)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, before, after in (("Heading 1", 16, 16, 6), ("Heading 2", 13, 12, 4), ("Heading 3", 11, 10, 3)):
        h = st[name]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = _rgb(ACCENT)
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
        hr = h.element.get_or_add_rPr()
        hf = hr.find(qn("w:rFonts"))
        if hf is None:
            hf = OxmlElement("w:rFonts")
            hr.append(hf)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            hf.set(qn(a), FONT)
    for name in ("Title", "Subtitle"):
        st[name].font.name = FONT
    if COVER_STYLE not in [s.name for s in st]:
        cs = st.add_style(COVER_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        cs.base_style = st["Normal"]
        cs.font.size = Pt(28)
        cs.font.bold = True
        cs.font.color.rgb = _rgb(INK)
        cs.paragraph_format.space_after = Pt(8)
    if TABLE_STYLE not in [s.name for s in st]:
        ts = st.add_style(TABLE_STYLE, WD_STYLE_TYPE.TABLE)
        ts.base_style = st["Table Grid"]
        ts.font.name = FONT
        ts.font.size = Pt(9)
    caption = st["Caption"] if "Caption" in [s.name for s in st] else st.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = _rgb(MUTED)


def document(meta: Meta, kind: str = "", subtitle: str = "", cover: bool | str = True, landscape: bool = False,
             details: Sequence[tuple[str, str]] = ()) -> Any:
    """A new house-style document.

    cover=True      page 1 is the cover page; the running header starts on page 2.
    cover="block"   a compact cover block opens page 1 and the text continues below it (a fixed-length
                    report whose page 1 is "cover and executive summary").
    cover=False     no cover; the running header is on every page.
    """
    doc = Document()
    _styles(doc)
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.header_distance = sec.footer_distance = Cm(1.0)
    cp = doc.core_properties
    cp.title = meta.title
    cp.subject = meta.scope_line()
    cp.author = meta.generated_by or "CreditProbe"
    cp.keywords = STYLE_ID
    cp.category = kind or "CreditProbe document"
    cp.comments = meta.generated_line()

    # running header and footer (from page 2 when there is a cover)
    hp = sec.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run("CREDITPROBE  ")
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = _rgb(ACCENT)
    r = hp.add_run(" · ".join(x for x in (meta.title, meta.book_label, meta.period) if x))
    r.font.size = Pt(8)
    r.font.color.rgb = _rgb(MUTED)
    r = hp.add_run(f"\t\t{meta.classification}")
    r.font.size = Pt(8)
    r.bold = True
    r.font.color.rgb = _rgb(MUTED)
    _bottom_rule(hp, RULE, 6)
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    r = fp.add_run(f"CreditProbe · {meta.generated_line()}\t\tPage ")
    r.font.size = Pt(8)
    r.font.color.rgb = _rgb(MUTED)
    for code, tail in (("PAGE", " of "), ("NUMPAGES", "")):
        f = fp.add_run()
        f.font.size = Pt(8)
        _field(f, code)
        if tail:
            t = fp.add_run(tail)
            t.font.size = Pt(8)
            t.font.color.rgb = _rgb(MUTED)
    if cover:
        sec.different_first_page_header_footer = True
        _cover(doc, meta, kind, subtitle, details, compact=cover == "block")
    return doc


def _cover(doc, meta: Meta, kind: str, subtitle: str, details: Sequence[tuple[str, str]], compact: bool = False) -> None:
    bar = doc.add_table(rows=1, cols=1)
    bar.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = bar.rows[0].cells[0]
    shade(c, ACCENT)
    p = c.paragraphs[0]
    r = p.add_run("CREDITPROBE")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = _rgb(WHITE)
    r = p.add_run("    Credit risk intelligence")
    r.font.size = Pt(10)
    r.font.color.rgb = _rgb(WHITE)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(14 if compact else 90)
    if kind:
        k = sp.add_run(kind.upper())
        k.bold = True
        k.font.size = Pt(11)
        k.font.color.rgb = _rgb(ACCENT)
    t = doc.add_paragraph(meta.title, style=COVER_STYLE)
    _bottom_rule(t, ACCENT, 18)
    if subtitle:
        s = doc.add_paragraph()
        r = s.add_run(subtitle)
        r.font.size = Pt(14)
        r.font.color.rgb = _rgb(MUTED)
    if not compact:
        doc.add_paragraph().paragraph_format.space_before = Pt(60)
    rows = [("Book", meta.book_label), ("Period", meta.period), ("Scope", meta.subject),
            ("Prepared by", meta.generated_by or "CreditProbe"), ("Generated", meta.when),
            ("Reference", meta.reference), *details, ("Classification", meta.classification)]
    rows = [(k, v) for k, v in rows if v]
    tbl = doc.add_table(rows=0, cols=2)
    for k, v in rows:
        cells = tbl.add_row().cells
        cells[0].text = ""
        kr = cells[0].paragraphs[0].add_run(k)
        kr.bold = True
        kr.font.size = Pt(9.5)
        kr.font.color.rgb = _rgb(MUTED)
        cells[1].text = ""
        vr = cells[1].paragraphs[0].add_run(str(v))
        vr.font.size = Pt(9.5)
        for cc in cells:
            _cell_border(cc, bottom=RULE)
    for row in tbl.rows:
        row.cells[0].width = Cm(4.5)
        row.cells[1].width = Cm(11.5)
    if compact:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    else:
        page_break(doc)


def _cell_border(cell, bottom: str | None = None) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "right", "bottom"):
        e = OxmlElement(f"w:{side}")
        if side == "bottom" and bottom:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), bottom)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tcpr.append(b)


def heading(doc, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def para(doc, text: str, size: float | None = None, bold: bool = False, italic: bool = False,
         colour: str | None = None, align: str | None = None):
    """A paragraph; **double asterisks** mark bold runs."""
    p = doc.add_paragraph()
    for i, chunk in enumerate(str(text).split("**")):
        if not chunk:
            continue
        r = p.add_run(chunk)
        if size:
            r.font.size = Pt(size)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        if colour:
            r.font.color.rgb = _rgb(colour.lstrip("#"))
    if align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def note(doc, text: str):
    return para(doc, text, size=8.5, italic=True, colour=MUTED)


def bullets(doc, items: Sequence[str], numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    for it in items:
        p = doc.add_paragraph(style=style)
        for i, chunk in enumerate(str(it).split("**")):
            if chunk:
                p.add_run(chunk).bold = i % 2 == 1


def fmt(value: Any, kind: str = "") -> str:
    """A number as the screen shows it, by column kind."""
    if value is None or value == "":
        return "—"
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return str(value)
    if kind == "sar_mn":
        return f"{value:,.1f}"
    if kind in ("sar", "int", "days", "sar_k"):
        return f"{value:,.0f}"
    if kind in ("pct", "pct1"):
        return f"{value:,.{1 if kind == 'pct1' else 2}f}%"
    if kind in ("share", "prob"):
        return f"{value * 100:,.2f}%"
    if kind == "bps":
        return f"{value:,.0f} bps"
    if kind == "x":
        return f"{value:,.2f}x"
    if kind == "num4":
        return f"{value:.4f}"
    return f"{value:,.0f}" if isinstance(value, int) else f"{value:,.2f}"


def table(doc, header: Sequence[str], rows: Sequence[Sequence[Any]], kinds: Sequence[str] | None = None,
          widths: Sequence[float] | None = None, right_from: int | None = None, size: float = 8.5,
          total_row: bool = False, gap: bool = True):
    """A house-style table: brand header row (repeats on each page), banded rows, numbers right-aligned.

    Cells may be numbers (formatted by `kinds`) or text. A text cell reading RED/AMBER/GREEN (or PASS/FAIL)
    is coloured and keeps its word.
    """
    t = doc.add_table(rows=1, cols=len(header))
    t.style = doc.styles[TABLE_STYLE]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    kinds = list(kinds or [""] * len(header))
    hdr = t.rows[0]
    trpr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    for i, h in enumerate(header):
        c = hdr.cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = _rgb(WHITE)
        shade(c, ACCENT)
        if _numeric_col(rows, i, kinds, right_from):
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        last = total_row and k == len(rows) - 1
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            text = fmt(v, kinds[i] if i < len(kinds) else "")
            run = p.add_run(text)
            run.font.size = Pt(size)
            run.bold = last
            if _numeric_col(rows, i, kinds, right_from):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            st = STATUS.get(str(v).upper()) if isinstance(v, str) else None
            if st:
                shade(cells[i], st[0])
                run.bold = True
                run.font.color.rgb = _rgb(st[1])
            elif k % 2:
                shade(cells[i], BAND)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Cm(w)
    if gap:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def _numeric_col(rows, i: int, kinds: Sequence[str], right_from: int | None) -> bool:
    if right_from is not None:
        return i >= right_from
    if i < len(kinds) and kinds[i] and kinds[i] not in ("text", "rating"):
        return True
    vals = [r[i] for r in rows if i < len(r) and r[i] not in (None, "")]
    return bool(vals) and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals)


def key_figures(doc, items: Sequence[tuple[str, str]], per_row: int = 3) -> None:
    """Key-figure tiles: a row of shaded cells, label above value."""
    rows = [items[i:i + per_row] for i in range(0, len(items), per_row)]
    t = doc.add_table(rows=len(rows), cols=per_row)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, chunk in enumerate(rows):
        for ci in range(per_row):
            c = t.rows[ri].cells[ci]
            c.text = ""
            shade(c, BAND)
            _cell_border(c)
            if ci >= len(chunk):
                continue
            label, value = chunk[ci]
            p = c.paragraphs[0]
            lr = p.add_run(label.upper())
            lr.font.size = Pt(7.5)
            lr.bold = True
            lr.font.color.rgb = _rgb(MUTED)
            p2 = c.add_paragraph()
            vr = p2.add_run(str(value))
            vr.font.size = Pt(14)
            vr.bold = True
            vr.font.color.rgb = _rgb(ACCENT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def figure(doc, image: bytes | io.BytesIO, width_cm: float = 16.0, caption: str = "") -> None:
    buf = io.BytesIO(image) if isinstance(image, (bytes, bytearray)) else image
    doc.add_picture(buf, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        doc.add_paragraph(caption, style="Caption")


def page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def to_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def chart_style() -> None:
    """Matplotlib defaults for charts embedded in Word: the house palette and font, light grid."""
    import matplotlib
    matplotlib.use("Agg")
    from cycler import cycler
    import matplotlib.pyplot as plt
    from backend.house_style import CHART
    plt.rcParams.update({
        "axes.prop_cycle": cycler(color=CHART), "font.family": "sans-serif", "font.size": 8,
        "axes.edgecolor": "#" + RULE, "axes.labelcolor": "#" + INK, "xtick.color": "#" + MUTED,
        "ytick.color": "#" + MUTED, "axes.spines.top": False, "axes.spines.right": False,
        "axes.grid": True, "grid.color": "#" + BAND, "grid.linewidth": 0.8, "axes.axisbelow": True,
        "axes.titleweight": "bold", "axes.titlesize": 9, "axes.titlecolor": "#" + INK, "legend.frameon": False,
    })


__all__ = ["document", "heading", "para", "note", "bullets", "table", "key_figures", "figure", "page_break",
           "to_bytes", "fmt", "shade", "chart_style", "TABLE_STYLE", "COVER_STYLE"]
