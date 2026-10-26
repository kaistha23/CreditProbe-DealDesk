"""Routes for P3 / P4 / P6, mounted on the Early Warning API (port 8501) under /api/v1/ews/intel."""
from __future__ import annotations

import io
from typing import Any

from fastapi import APIRouter, Body, HTTPException, Query, Request
from fastapi.responses import Response

from backend.risk_intel import diagnostics, ecl_impact, store, watch

router = APIRouter(prefix="/api/v1/ews/intel")


def _who(request: Request) -> str:
    try:
        from backend.access.middleware import _user
        u = _user(request) or {}
        return str(u.get("username") or u.get("name") or "admin")
    except Exception:  # noqa: BLE001
        return "admin"


@router.get("/pd-by-bucket")
def get_pd() -> dict[str, Any]:
    return store.pd_by_bucket()


@router.put("/pd-by-bucket")
def put_pd(request: Request, body: dict[str, Any] = Body(...)) -> dict[str, Any]:
    try:
        return store.set_pd_by_bucket(body.get("values") or body, _who(request))
    except ValueError as e:
        raise HTTPException(422, {"error_code": "PD_INVALID", "message": str(e)}) from e


@router.get("/ecl-impact")
def get_impact(domain: str = Query("corporate"), period: str = Query(""), segment: str = Query("")) -> dict[str, Any]:
    return ecl_impact.impact(domain, period, segment)


@router.get("/watch")
def get_watch(tier: int | None = Query(None)) -> dict[str, Any]:
    return watch.view(tier)


@router.get("/actions")
def get_actions() -> list[dict[str, Any]]:
    return store.actions()


@router.post("/actions", status_code=201)
def post_action(request: Request, body: dict[str, Any] = Body(...)) -> dict[str, Any]:
    kind = str(body.get("kind") or "")
    if kind not in ("bank_watchlist", "accepted_risk_watch", "sicr_assessment"):
        raise HTTPException(422, {"error_code": "ACTION_UNKNOWN", "message": "Unknown action."})
    domain = "retail" if body.get("domain") == "retail" else "corporate"
    buckets = [b for b in (body.get("buckets") or ["critical", "high"]) if b in ecl_impact.BUCKETS]
    ids = body.get("obligor_ids") or ecl_impact.owners_in_scope(domain, buckets, body.get("period") or "", body.get("segment") or "")
    names = {r["owner"]: r for r in ecl_impact.impact(domain, body.get("period") or "", body.get("segment") or "")["top_contributors"]}
    out = []
    for oid in ids[:200]:
        r = names.get(oid, {})
        out.append(store.add_action(kind, {"obligor_id": oid, "domain": domain, "name": r.get("owner_name") or body.get("name") or oid,
                                           "sector": r.get("segment", ""), "exposure": r.get("ead", 0.0), "ecl": r.get("current_ecl", 0.0),
                                           "tier": body.get("tier", 2), "note": body.get("note", "")}, _who(request)))
    return {"kind": kind, "added": len(out), "obligor_ids": ids[:200]}


@router.get("/diagnostics/waterfall")
def get_waterfall(book: str = Query("corporate")) -> dict[str, Any]:
    return diagnostics.waterfall(book)


@router.get("/diagnostics/anomalies")
def get_anomalies(book: str = Query("corporate")) -> dict[str, Any]:
    return diagnostics.anomalies(book)


@router.get("/diagnostics/coverage")
def get_coverage(book: str = Query("corporate")) -> dict[str, Any]:
    return diagnostics.coverage(book)


@router.get("/diagnostics/shocks")
def get_shocks(book: str = Query("corporate"), pd: float = Query(20.0), lgd: float = Query(5.0),
               macro: float = Query(0.0)) -> dict[str, Any]:
    return diagnostics.shocks(book, pd, lgd, macro)


def _fmt(v: float) -> str:
    return f"SAR {v:,.1f} million"


@router.get("/memo.docx")
def memo(domain: str = Query("corporate"), period: str = Query("")) -> Response:
    """Draft memo to the CRO: the ECL impact of current warnings, the severe case and the top contributors (Word)."""
    from docx import Document
    from docx.shared import Pt

    d = ecl_impact.impact(domain, period)
    w = watch.view()
    t = d["total"]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading("Memo to the Chief Risk Officer", level=1)
    doc.add_paragraph(f"Subject: ECL impact of current early warnings ({domain.title()} book, {d['period']})")
    doc.add_paragraph("From: Credit Risk · Status: draft for review")
    doc.add_heading("Summary", level=2)
    doc.add_paragraph(
        f"{t['names']:,} {d['labels']['names']} carry an early warning, with exposure of {_fmt(t['ead'])} and current ECL of "
        f"{_fmt(t['current_ecl'])}. If the warnings come true at the bucket probabilities below, expected ECL rises to "
        f"{_fmt(t['expected_ecl'])} (+{_fmt(t['uplift'])}). In the severe case, every Critical and High name at 90+ days "
        f"past due, ECL reaches {_fmt(t['severe_ecl'])} (+{_fmt(t['severe_uplift'])}).")
    doc.add_heading("By warning bucket", level=2)
    tb = doc.add_table(rows=1, cols=5)
    tb.style = "Light Grid Accent 1"
    for c, h in zip(tb.rows[0].cells, ("Bucket", "PD (management estimate)", "Names", "Current ECL", "Expected ECL")):
        c.text = h
    for b in d["by_bucket"]:
        r = tb.add_row().cells
        r[0].text, r[1].text, r[2].text = b["label"], f"{b['pd_pct']:.0f}%", f"{b['names']:,}"
        r[3].text, r[4].text = _fmt(b["current_ecl"]), _fmt(b["expected_ecl"])
    if d["top_contributors"]:
        doc.add_heading("Top 10 contributors", level=2)
        tb = doc.add_table(rows=1, cols=5)
        tb.style = "Light Grid Accent 1"
        for c, h in zip(tb.rows[0].cells, ("Borrower", "Sector", "Bucket", "Exposure", "ECL uplift")):
            c.text = h
        for x in d["top_contributors"]:
            r = tb.add_row().cells
            r[0].text, r[1].text, r[2].text = f"{x['owner_name']} ({x['owner']})", x["segment"], store.BUCKET_LABELS[x["bucket"]]
            r[3].text, r[4].text = _fmt(x["ead"]), _fmt(x["uplift"])
    doc.add_heading("Accepted Risk Watch", level=2)
    doc.add_paragraph(
        f"{w['names']} names on the list, exposure {_fmt(w['exposure'])}; {w['escalations']} meet the escalation rule "
        f"({w['rule']}). If Tier 1 moved to Stage 2, ECL would rise by {_fmt(w['tier1_stage2']['uplift'])}.")
    doc.add_heading("Requested decisions", level=2)
    for s in ("Note the expected ECL of current warnings and the severe case.",
              "Approve SICR assessments for the Critical names and the escalated Accepted Risk Watch names.",
              "Confirm the bucket probabilities (management estimate) for the next ECL committee."):
        doc.add_paragraph(s, style="List Bullet")
    doc.add_paragraph(f"Powered by {d['powered_by']['pack']} · {d['powered_by']['rule_id']} · {d['powered_by']['article']}")
    buf = io.BytesIO()
    doc.save(buf)
    name = f"CRO_memo_ECL_of_warnings_{domain}_{d['period']}.docx"
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})
