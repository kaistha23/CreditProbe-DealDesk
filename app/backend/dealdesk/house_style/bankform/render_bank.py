"""R14 session p: the bank-form proposal laid out like the bank's own paper (SIB_GAP_AUDIT.md A1–A5, B1–B3).

    data, name = docx(paper)            # Word: no cover; header block on every page (Word page header); part banners;
                                        # action requested and recommendation in ruled boxes; facilities on a landscape page;
                                        # __key concerns__ bold and underlined; signature grid and blank boxes
    data, name, pages = pdf(paper)      # the same layout in PDF (reportlab); pages = {section_key: first page}, "total"
    parts(paper)                        # the sections grouped by voice: Action requested (business), Credit division comments,
                                        # CreditProbe analysis, Credit division recommendation, Approvals

Content model as render.py (blocks.py); text markup: **bold**, __bold underlined__. House style kept for the check
(`backend.house_style.check.docx(data, cover=False) == []`).
"""
from __future__ import annotations

import io
import re
from typing import Any

from backend.house_style import ACCENT, BAND, INK, MUTED, RULE, Meta
from backend.house_style import docx as hd

BODY_PT = 8.8
NUMERIC = re.compile(r"^[\s(+\-–]*[\d,.]+%?x?\)?$|^–$")
MARK = re.compile(r"(\*\*.+?\*\*|__.+?__)")

PARTS = [("action", "ACTION REQUESTED", "Business · Corporate Banking Group", ("action", "facilities", "security", "terms")),
         ("comments", "CREDIT DIVISION COMMENTS", "Credit division · independent review", ("group_exposure", "background", "relationship", "projects", "syndication", "escrow",
                                                                                          "sme_decision", "cbrbr", "bureau", "other_banks", "new_project", "financials",
                                                                                          "synopsis", "geopolitical", "exceptions")),
         ("creditprobe", "CREDITPROBE ANALYSIS", "Independent re-performance by CreditProbe, in the credit voice", ("cp_group", "cp_counterparty", "cp_selected", "cp_findings", "cp_raroc", "cp_forward", "cp_stress",
                                                                                                                    "cp_precedent", "cp_swot")),
         ("recommendation", "CREDIT DIVISION RECOMMENDATION", "Credit division", ("strengths", "recommendation")),
         ("approvals", "APPROVALS", "Analyst · VP Credit · CCO · CRO · Credit committee", ("signatures",))]
BOXED = ("action", "recommendation")
LANDSCAPE = ("facilities",)
#: parts that open on a new page, as in the bank's paper (comments from page 2, the recommendation and approvals at the end)
NEW_PAGE = ("comments", "creditprobe", "recommendation", "approvals")


def parts(paper: dict[str, Any]) -> list[dict[str, Any]]:
    by = {s["key"]: s for s in paper["sections"]}
    out, used = [], {"header"}
    for key, title, sub, keys in PARTS:
        secs = [s for s in paper["sections"] if s["key"] in keys and s["key"] in by]  # Round K: the paper's order (a section-order rule)
        used |= {s["key"] for s in secs}
        if secs:
            out.append({"key": key, "title": title, "subtitle": sub, "sections": secs})
    rest = [s for s in paper["sections"] if s["key"] not in used]
    if rest:
        out[1 if len(out) > 1 else 0]["sections"] += rest
    return out


def _num(s: dict[str, Any], i: int) -> str:
    return f"{s['n']}." if s.get("n") else f"CP{i}."


def _meta(paper: dict[str, Any], generated_by: str) -> Meta:
    return Meta(f"Credit proposal · {paper['customer']}", book="corporate", period=str(paper.get("as_of_label") or paper.get("as_of") or ""),
                subject=" · ".join(x for x in (paper.get("sector"), paper.get("rating") and f"rating {paper['rating']}", paper.get("group")) if x),
                generated_by=generated_by, reference=f"{paper['proposal_id']}-BFP")


def _plain(t: Any) -> str:
    return str(t).replace("**", "").replace("__", "")


# ================================================================================================ Word

def _runs(par, text: str, size: float = BODY_PT, colour: str | None = None, bold: bool = False) -> None:
    from docx.shared import Pt, RGBColor
    for chunk in MARK.split(str(text)):
        if not chunk:
            continue
        b, u = bold, False
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            chunk, b = chunk[2:-2], True
        elif chunk.startswith("__") and chunk.endswith("__") and len(chunk) > 4:
            chunk, b, u = chunk[2:-2].replace("**", ""), True, True
        r = par.add_run(chunk)
        r.bold, r.underline = b, u
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(colour or INK)


def _borders(tbl, colour: str = RULE, size: int = 6) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:color"), colour)
        b.append(e)
    pr.append(b)


def _widths(tb, ws: list[float]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm
    tb.autofit = False
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tb._tbl.tblPr.append(lay)
    for i, gc in enumerate(tb._tbl.tblGrid.findall(qn("w:gridCol"))):
        if i < len(ws):
            gc.set(qn("w:w"), str(int(Cm(ws[i]).twips)))
    for row in tb.rows:
        for i, c in enumerate(row.cells[:len(ws)]):
            c.width = Cm(ws[i])


def _scale(ws: list[float] | None, n: int, total: float) -> list[float]:
    ws = list(ws or [])
    if len(ws) != n:
        ws = [total / n] * n
    s = sum(ws)
    return [w * total / s for w in ws]


def _table(host, doc, header: list[str], rows: list[list[Any]], widths: list[float] | None, size: float, total_cm: float, source: str = "") -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    n = len(header)
    tb = host.add_table(rows=1 + len(rows), cols=n)
    tb.style = doc.styles["CP Table"]
    _widths(tb, _scale(widths, n, total_cm))
    for j, h in enumerate(header):
        c = tb.rows[0].cells[j]
        c.text = ""
        hd.shade(c, ACCENT)
        _runs(c.paragraphs[0], _plain(h), size=size, colour="FFFFFF", bold=True)
    total_row = bool(rows) and str(rows[-1][0]).startswith("**")
    for i, r in enumerate(rows, start=1):
        for j in range(n):
            v = r[j] if j < len(r) else ""
            c = tb.rows[i].cells[j]
            c.text = ""
            _runs(c.paragraphs[0], str(v).replace("    ", " "), size=size, bold=total_row and i == len(rows))
            if i % 2 == 0:
                hd.shade(c, BAND)
    for j in range(1, n):
        vals = [_plain(r[j]).strip() for r in rows if j < len(r) and _plain(r[j]).strip() not in ("", "–", "-")]
        if vals and all(NUMERIC.match(v) for v in vals):
            for row in tb.rows:
                row.cells[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if source:
        p = host.add_paragraph()
        _runs(p, f"Source: {source}", size=6.8, colour=MUTED)
        p.paragraph_format.space_after = Pt(3)


def _kv(host, doc, pairs: list[tuple[str, str]], cols: int, total_cm: float, size: float = 7.6) -> None:
    per = max(1, cols)
    rows = [pairs[i:i + per] for i in range(0, len(pairs), per)]
    tb = host.add_table(rows=len(rows), cols=per * 2)
    tb.style = doc.styles["Table Grid"]
    _widths(tb, _scale([3.0, total_cm / per - 3.0] * per, per * 2, total_cm))
    for ri, chunk in enumerate(rows):
        for ci, (k, v) in enumerate(chunk):
            kc, vc = tb.rows[ri].cells[ci * 2], tb.rows[ri].cells[ci * 2 + 1]
            kc.text, vc.text = "", ""
            hd.shade(kc, BAND)
            _runs(kc.paragraphs[0], k, size=size - 0.4, colour=MUTED, bold=True)
            _runs(vc.paragraphs[0], v, size=size)


def _blocks(host, doc, blocks: list[dict[str, Any]], total_cm: float) -> None:
    from docx.shared import Pt
    for b in blocks:
        t = b["t"]
        if t == "p":
            p = host.add_paragraph()
            _runs(p, b["text"], size=BODY_PT - (1 if b.get("exempt") else 0), colour=MUTED if b.get("exempt") else None)
            p.paragraph_format.space_after = Pt(3)
        elif t == "h":
            p = host.add_paragraph()
            _runs(p, b["text"], size=BODY_PT + 0.4, colour=ACCENT, bold=True)
            p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(4), Pt(2)
        elif t in ("bullets", "numbered"):
            for k, it in enumerate(b["items"], start=1):
                p = host.add_paragraph()
                p.paragraph_format.left_indent = Pt(14)
                p.paragraph_format.first_line_indent = Pt(-11)
                p.paragraph_format.space_after = Pt(1.5)
                _runs(p, (f"{k}.  " if t == "numbered" else "•  ") + str(it))
        elif t == "table":
            _table(host, doc, b["header"], b["rows"], b.get("widths"), b.get("size") or 7.4, total_cm, b.get("source", ""))
        elif t == "kv":
            _kv(host, doc, b["pairs"], b.get("cols", 2), total_cm)
            host.add_paragraph().paragraph_format.space_after = Pt(1)
        elif t == "sign":
            rows = [[r, n or "", "", ""] for r, n in b["rows"]]
            tb = host.add_table(rows=1 + len(rows), cols=4)
            tb.style = doc.styles["Table Grid"]
            _widths(tb, _scale([5.0, 4.6, 4.8, 3.0], 4, total_cm))
            for j, h in enumerate(("Role", "Name", "Signature", "Date")):
                c = tb.rows[0].cells[j]
                c.text = ""
                hd.shade(c, BAND)
                _runs(c.paragraphs[0], h, size=8, bold=True)
            for i, r in enumerate(rows, start=1):
                for j, v in enumerate(r):
                    c = tb.rows[i].cells[j]
                    c.text = ""
                    _runs(c.paragraphs[0], v, size=8.5)
                    c.add_paragraph("")
            host.add_paragraph().paragraph_format.space_after = Pt(2)
        elif t == "box":
            tb = host.add_table(rows=1, cols=1)
            _borders(tb, INK, 8)
            _widths(tb, [total_cm])
            c = tb.rows[0].cells[0]
            c.text = ""
            _runs(c.paragraphs[0], b["title"], size=8.6, bold=True)
            for _ in range(b.get("lines", 4)):
                c.add_paragraph("")
            _runs(c.add_paragraph(), "Signature: ______________________    Name: ______________________    Date: ____________", size=8, colour=MUTED)
            host.add_paragraph().paragraph_format.space_after = Pt(2)


def _banner(doc, title: str, sub: str, total_cm: float) -> None:
    from docx.shared import Pt
    tb = doc.add_table(rows=1, cols=1)
    _widths(tb, [total_cm])
    c = tb.rows[0].cells[0]
    c.text = ""
    hd.shade(c, ACCENT)
    p = c.paragraphs[0]
    _runs(p, title, size=10.5, colour="FFFFFF", bold=True)
    _runs(p, f"   {sub}", size=8, colour="DCE6F2")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _header_block(doc, sec, line: list[tuple[str, str]], width_cm: float, form_title: str = "CREDIT PROPOSAL") -> None:
    """The compact header block in the Word page header of a section (every page), keeping the house-style marker."""
    from docx.shared import Cm, Pt
    hdr = sec.header
    hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0]
    hp.text = ""
    _runs(hp, "CREDITPROBE  ", size=7, colour=ACCENT, bold=True)
    _runs(hp, f"{form_title} (BANK FORM) · INTERNAL", size=7, colour=MUTED, bold=True)
    hp.paragraph_format.space_after = Pt(1)
    per = 4
    rows = [line[i:i + per] for i in range(0, len(line), per)]
    tb = hdr.add_table(rows=len(rows), cols=per * 2, width=Cm(width_cm))
    tb.style = doc.styles["Table Grid"]
    _borders(tb, RULE, 4)
    _widths(tb, _scale([1.5, 2.9] * per, per * 2, width_cm))
    for ri, chunk in enumerate(rows):
        for ci in range(per):
            k, v = chunk[ci] if ci < len(chunk) else ("", "")
            kc, vc = tb.rows[ri].cells[ci * 2], tb.rows[ri].cells[ci * 2 + 1]
            kc.text, vc.text = "", ""
            if k:
                hd.shade(kc, BAND)
            _runs(kc.paragraphs[0], k, size=6.4, colour=MUTED, bold=True)
            _runs(vc.paragraphs[0], v, size=6.6)


def _new_section(doc, landscape: bool):
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.shared import Cm
    prev = doc.sections[-1]
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    w, h = sorted((prev.page_width, prev.page_height))
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = (h, w) if landscape else (w, h)
    for a in ("left_margin", "right_margin", "top_margin", "bottom_margin", "header_distance", "footer_distance"):
        setattr(sec, a, getattr(prev, a))
    sec.top_margin = Cm(3.3)
    return sec


def docx(paper: dict[str, Any], generated_by: str = "") -> tuple[bytes, str]:
    from docx.shared import Cm, Pt
    meta = _meta(paper, generated_by)
    doc = hd.document(meta, kind="Credit proposal", cover=False)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin, sec.bottom_margin = Cm(3.3), Cm(1.7)
    sec.header_distance, sec.footer_distance = Cm(0.7), Cm(0.8)
    line = [tuple(x) for x in (paper.get("header_line") or [("Customer", paper["customer"]), ("Ref", paper["proposal_id"])])]
    pw = 21.0 - 3.2
    lw = 29.7 - 3.2
    _header_block(doc, sec, line, pw, paper.get("form_title") or "CREDIT PROPOSAL")
    for r in sec.footer.paragraphs[0].runs:  # one line on a portrait page: "… · Page X of Y"
        if "\t\t" in r.text:  # only the text run (the PAGE / NUMPAGES field runs keep their field codes)
            r.text = r.text.replace("\t\t", "  ·  ")
    by = {s["key"]: s for s in paper["sections"]}
    # page 1: title line and the full header block (section 1)
    p = doc.add_paragraph()
    _runs(p, paper.get("form_title") or "CREDIT PROPOSAL", size=13, colour=ACCENT, bold=True)  # R16 c: "CREDIT DIVISION RESPONSE"
    _runs(p, f"   {paper['customer']} · {paper['proposal_id']} · {paper.get('as_of_label') or ''}", size=9, colour=MUTED)
    hd.heading(doc, "1. Header block", 2)
    if "header" in by:
        _blocks(doc, doc, by["header"]["blocks"], pw)
    cp_i = 0
    for pi, part in enumerate(parts(paper)):
        if pi and part["key"] in NEW_PAGE:
            hd.page_break(doc)
        _banner(doc, part["title"], part["subtitle"], pw)
        for s in part["sections"]:
            if s["key"] in LANDSCAPE:
                ls = _new_section(doc, True)
                _header_block(doc, ls, line, lw, paper.get("form_title") or "CREDIT PROPOSAL")
            if not s.get("n"):
                cp_i += 1
            hd.heading(doc, f"{_num(s, cp_i)} {s['heading']}", 2)
            if s["key"] in BOXED:
                tb = doc.add_table(rows=1, cols=1)
                _borders(tb, INK, 10)
                _widths(tb, [pw - 0.2])
                cell = tb.rows[0].cells[0]
                cell.paragraphs[0].text = ""
                _blocks(cell, doc, s["blocks"], pw - 0.6)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            else:
                _blocks(doc, doc, s["blocks"], lw if s["key"] in LANDSCAPE else pw)
            if s["key"] in LANDSCAPE:
                ps = _new_section(doc, False)
                _header_block(doc, ps, line, pw, paper.get("form_title") or "CREDIT PROPOSAL")
    safe = re.sub(r"[^\w-]+", "-", f"Credit-proposal-{paper['proposal_id']}-{paper['customer']}")[:90].strip("-")
    return hd.to_bytes(doc), f"{safe}.docx"


# ================================================================================================ PDF

def _esc(t: Any) -> str:
    s = str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    s = re.sub(r"__(.+?)__", lambda m: "<u><b>" + m.group(1).replace("**", "") + "</b></u>", s)
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)


def pdf(paper: dict[str, Any], generated_by: str = "") -> tuple[bytes, str, dict[str, Any]]:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (BaseDocTemplate, Frame, KeepTogether, NextPageTemplate, PageBreak, PageTemplate, Paragraph, Spacer, Table,
                                    TableStyle)

    meta = _meta(paper, generated_by)
    acc, band, muted, rule, ink = (colors.HexColor("#" + x) for x in (ACCENT, BAND, MUTED, RULE, INK))
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9.2, leading=9.2 * 1.3, spaceAfter=3.5, textColor=ink)
    small = ParagraphStyle("s", parent=body, fontSize=7, leading=8.6, spaceAfter=0)
    note = ParagraphStyle("n", parent=body, fontSize=6.8, textColor=muted, fontName="Helvetica-Oblique", spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=11, textColor=acc, spaceBefore=8, spaceAfter=3)
    h3 = ParagraphStyle("h3", parent=body, fontName="Helvetica-Bold", fontSize=9.2, textColor=acc, spaceBefore=4, spaceAfter=2)
    head = ParagraphStyle("th", parent=small, textColor=colors.white, fontName="Helvetica-Bold")
    LM = 1.6 * cm
    PW, LW = A4[0] - 2 * LM, A4[1] - 2 * LM
    line = [tuple(x) for x in (paper.get("header_line") or [("Customer", paper["customer"]), ("Ref", paper["proposal_id"])])]
    width = {"w": PW}

    def tbl(header, rows, widths=None, size=7.4, W=None):
        W = W or width["w"]
        size = size * 1.08  # Helvetica in the PDF, read at the Word table size
        st = ParagraphStyle("c", parent=small, fontSize=size, leading=size * 1.22)
        sth = ParagraphStyle("ch", parent=head, fontSize=size, leading=size * 1.22)
        data = [[Paragraph(_esc(h), sth) for h in header]] + [[Paragraph(_esc(c).replace("    ", "&nbsp;&nbsp;&nbsp;&nbsp;"), st) for c in r] for r in rows]
        cw = [W * w / sum(widths) for w in widths] if widths and len(widths) == len(header) else None
        t = Table(data, colWidths=cw, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), acc), ("GRID", (0, 0), (-1, -1), 0.25, rule), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                 ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
                 ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]
        for i in range(2, len(data), 2):
            style.append(("BACKGROUND", (0, i), (-1, i), band))
        t.setStyle(TableStyle(style))
        return t

    def flow(bs, W):
        out = []
        for b in bs:
            t = b["t"]
            if t == "p":
                out.append(Paragraph(_esc(b["text"]), note if b.get("exempt") else body))
            elif t == "h":
                out.append(Paragraph(_esc(b["text"]), h3))
            elif t in ("bullets", "numbered"):
                for k, x in enumerate(b["items"], start=1):
                    out.append(Paragraph(("%d.&nbsp;&nbsp;" % k if t == "numbered" else "•&nbsp;&nbsp;") + _esc(x),
                                         ParagraphStyle("li", parent=body, leftIndent=12, firstLineIndent=-10, spaceAfter=1.5)))
            elif t == "table":
                out.append(tbl(b["header"], b["rows"], b.get("widths"), b.get("size") or 7.4, W))
                out.append(Paragraph(_esc("Source: " + b["source"]) if b.get("source") else "", note))
            elif t == "kv":
                per = b.get("cols", 2)
                rows = []
                for i in range(0, len(b["pairs"]), per):
                    r = []
                    for k, v in b["pairs"][i:i + per]:
                        r += [f"**{k}**", v]
                    rows.append(r + [""] * (per * 2 - len(r)))
                kt = Table([[Paragraph(_esc(c), small) for c in r] for r in rows], colWidths=[W * w / (8.5 * per) for w in [3.0, 5.5] * per])
                kst = [("GRID", (0, 0), (-1, -1), 0.25, rule), ("VALIGN", (0, 0), (-1, -1), "TOP")]
                for j in range(0, per * 2, 2):
                    kst.append(("BACKGROUND", (j, 0), (j, -1), band))
                kt.setStyle(TableStyle(kst))
                out += [kt, Spacer(1, 4)]
            elif t == "sign":
                out.append(tbl(["Role", "Name", "Signature", "Date"], [[r, n or "", "<br/><br/>", ""] for r, n in b["rows"]], [5.0, 4.6, 4.8, 3.0], 8.4, W))
                out.append(Spacer(1, 6))
            elif t == "box":
                bx = Table([[Paragraph(f"<b>{_esc(b['title'])}</b>", body)]] + [[""] for _ in range(b.get("lines", 4))]
                           + [[Paragraph("Signature: ____________________   Name: ____________________   Date: __________", small)]], colWidths=[W])
                bx.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, ink)]))
                out.append(KeepTogether([bx, Spacer(1, 6)]))
        return out

    pages: dict[str, int] = {}

    class Doc(BaseDocTemplate):
        def afterFlowable(self, f):  # noqa: N802 - reportlab hook
            k = getattr(f, "_section_key", None)
            if k and k not in pages:
                pages[k] = self.page

    def frame_page(c, d, size):
        w, h = size
        c.saveState()
        c.setFont("Helvetica-Bold", 7)
        c.setFillColor(acc)
        c.drawString(LM, h - 0.9 * cm, "CREDITPROBE")
        c.setFillColor(muted)
        c.drawString(LM + 2.0 * cm, h - 0.9 * cm, f"{paper.get('form_title') or 'CREDIT PROPOSAL'} (BANK FORM) · INTERNAL")
        per = 4
        cw = (w - 2 * LM) / per
        y0 = h - 1.1 * cm
        rows = [line[i:i + per] for i in range(0, len(line), per)]
        rh = 0.42 * cm
        for ri, chunk in enumerate(rows):
            for ci, (k, v) in enumerate(chunk):
                x, y = LM + ci * cw, y0 - (ri + 1) * rh
                c.setFillColor(band)
                c.rect(x, y, cw * 0.32, rh, stroke=0, fill=1)
                c.setStrokeColor(rule)
                c.setLineWidth(0.3)
                c.rect(x, y, cw, rh, stroke=1, fill=0)
                c.setFillColor(muted)
                c.setFont("Helvetica-Bold", 5.8)
                c.drawString(x + 2, y + 3.2, str(k)[:14])
                c.setFillColor(ink)
                c.setFont("Helvetica", 6.2)
                c.drawString(x + cw * 0.32 + 2, y + 3.2, str(v)[:int(cw / 3.1)])
        c.setFont("Helvetica", 7)
        c.setFillColor(muted)
        c.drawString(LM, 0.8 * cm, f"CreditProbe · {meta.generated_line()}"[:110])
        c.drawRightString(w - LM, 0.8 * cm, f"Page {d.page}")
        c.restoreState()

    buf = io.BytesIO()
    doc = Doc(buf, pagesize=A4, leftMargin=LM, rightMargin=LM, topMargin=2.3 * cm, bottomMargin=1.5 * cm, title=meta.title,
              author=meta.generated_by or "CreditProbe", subject=meta.scope_line(), keywords="creditprobe-house-style-1")
    fp = Frame(LM, 1.5 * cm, PW, A4[1] - 4.6 * cm, id="p")
    fl = Frame(LM, 1.5 * cm, LW, A4[0] - 4.6 * cm, id="l")
    doc.addPageTemplates([PageTemplate(id="portrait", frames=[fp], pagesize=A4, onPage=lambda c, d: frame_page(c, d, A4)),
                          PageTemplate(id="landscape", frames=[fl], pagesize=landscape(A4), onPage=lambda c, d: frame_page(c, d, landscape(A4)))])
    story: list[Any] = []
    by = {s["key"]: s for s in paper["sections"]}
    t0 = Paragraph(f"<font color='#{ACCENT}' size=13><b>{_esc(paper.get('form_title') or 'CREDIT PROPOSAL')}</b></font>&nbsp;&nbsp;<font color='#{MUTED}'>{_esc(paper['customer'])} · "
                   f"{_esc(paper['proposal_id'])} · {_esc(paper.get('as_of_label') or '')}</font>", body)
    story.append(t0)
    hh = Paragraph("1. Header block", h2)
    hh._section_key = "header"
    story.append(hh)
    if "header" in by:
        story += flow(by["header"]["blocks"], PW)

    def banner(title, sub):
        b = Table([[Paragraph(f"<font color='white'><b>{title}</b></font>&nbsp;&nbsp;&nbsp;<font color='#DCE6F2' size=7.5>{_esc(sub)}</font>", body)]], colWidths=[width["w"]])
        b.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), acc), ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        return b

    cp_i = 0
    for pi, part in enumerate(parts(paper)):
        if pi and part["key"] in NEW_PAGE:
            story.append(PageBreak())
        story += [Spacer(1, 4), banner(part["title"], part["subtitle"]), Spacer(1, 2)]
        for s in part["sections"]:
            land = s["key"] in LANDSCAPE
            if land:
                story += [NextPageTemplate("landscape"), PageBreak()]
                width["w"] = LW
            if not s.get("n"):
                cp_i += 1
            h = Paragraph(_esc(f"{_num(s, cp_i)} {s['heading']}"), h2)
            h._section_key = s["key"]
            fl_ = flow(s["blocks"], width["w"] - (14 if s["key"] in BOXED else 0))
            if s["key"] in BOXED:
                bx = Table([[x] for x in fl_], colWidths=[width["w"]])
                bx.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 1.0, ink), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]))
                story += [h, bx]
            else:
                story += [h] + fl_
            if land:
                story += [NextPageTemplate("portrait"), PageBreak()]
                width["w"] = PW
    doc.build(story)
    data = buf.getvalue()
    from pypdf import PdfReader
    total = len(PdfReader(io.BytesIO(data)).pages)
    safe = re.sub(r"[^\w-]+", "-", f"Credit-proposal-{paper['proposal_id']}-{paper['customer']}")[:90].strip("-")
    return data, f"{safe}.pdf", {"sections": pages, "total": total}


def page_map(paper: dict[str, Any]) -> dict[str, Any]:
    """{section_key: first page} and the page count, from the PDF layout (the on-screen preview paginates by it)."""
    return pdf(paper)[2]
