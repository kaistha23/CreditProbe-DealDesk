"""A paper in Word (house style, `backend.house_style.docx`) and PDF (reportlab), from one content model.

    data, name = docx(paper, kind="Credit proposal")                 # check.docx(data) == []
    data, name = pdf(paper, kind="Credit proposal")
    parts=[{"title": "Decision summary", "sections": [...]}, {"title": "Full paper", "sections": paper["sections"]}]
    lets the committee paper put its 2-page summary before the full paper (a page break between parts).
"""
from __future__ import annotations

import io
import re
from typing import Any

from backend.house_style import ACCENT, BAND, INK, MUTED, RULE, Meta
from backend.house_style import docx as hd

BODY_PT = 9.0
NUMERIC = re.compile(r"^[\s(+\-–]*[\d,.]+%?x?\)?$|^–$")


def _plain(t: Any) -> str:
    return str(t).replace("**", "")


def _right_from(rows: list[list[Any]]) -> int | None:
    if not rows:
        return None
    n = len(rows[0])
    for i in range(1, n):
        if all(NUMERIC.match(_plain(r[j]).strip() or "–") for r in rows for j in range(i, min(n, len(r)))):
            return i
    return None


def _meta(paper: dict[str, Any], kind: str, generated_by: str) -> Meta:
    return Meta(f"{kind} · {paper['customer']}", book="corporate", period=str(paper.get("as_of_label") or paper.get("as_of") or ""),
                subject=" · ".join(x for x in (paper.get("sector"), paper.get("rating") and f"rating {paper['rating']}", paper.get("group")) if x),
                generated_by=generated_by, reference=f"{paper['proposal_id']}-{'CP' if 'ommittee' in kind else 'BFP'}")


def _parts(paper: dict[str, Any], parts: list[dict[str, Any]] | None) -> list[dict[str, Any]]:
    return parts or [{"title": "", "sections": paper["sections"], "numbered": True}]


# ------------------------------------------------------------------------------------------------ Word

def _kv(doc, pairs: list[tuple[str, str]], cols: int = 2) -> None:
    from docx.shared import Pt, RGBColor
    per = max(1, cols)
    rows = [pairs[i:i + per] for i in range(0, len(pairs), per)]
    t = doc.add_table(rows=len(rows), cols=per * 2)
    t.style = doc.styles["Table Grid"]
    for ri, chunk in enumerate(rows):
        for ci, (k, v) in enumerate(chunk):
            kc, vc = t.rows[ri].cells[ci * 2], t.rows[ri].cells[ci * 2 + 1]
            kc.text, vc.text = "", ""
            hd.shade(kc, BAND)
            r = kc.paragraphs[0].add_run(k)
            r.bold, r.font.size = True, Pt(7.5)
            r.font.color.rgb = RGBColor.from_string(MUTED)
            r2 = vc.paragraphs[0].add_run(v)
            r2.font.size = Pt(8)
            r2.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _box(doc, title: str, lines: int) -> None:
    from docx.shared import Pt
    t = doc.add_table(rows=1, cols=1)
    t.style = doc.styles["Table Grid"]
    c = t.rows[0].cells[0]
    c.text = ""
    r = c.paragraphs[0].add_run(title)
    r.bold, r.font.size = True, Pt(8.5)
    for _ in range(lines):
        c.add_paragraph("")
    c.add_paragraph("Signature: ______________________    Name: ______________________    Date: ____________").runs[0].font.size = Pt(8)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def _fit(tb, rows: list[list[str]], widths: list[float] | None, total_cm: float = 17.4) -> None:
    """Fixed column widths (Word and LibreOffice honour the grid, not only the cells) and numeric columns right-aligned."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm
    n = len(tb.columns)
    ws = list(widths or [])
    if len(ws) != n:
        ws = [total_cm / n] * n
    scale = total_cm / sum(ws) if sum(ws) > total_cm else 1.0
    ws = [w * scale for w in ws]
    tb.autofit = False
    tblpr = tb._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    grid = tb._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < n:
            gc.set(qn("w:w"), str(int(Cm(ws[i]).twips)))
    for row in tb.rows:
        for i, c in enumerate(row.cells[:n]):
            c.width = Cm(ws[i])
    for i in range(1, n):
        vals = [r[i] for r in rows if i < len(r) and str(r[i]).strip() not in ("", "–", "-")]
        if vals and all(NUMERIC.match(str(v).strip()) for v in vals):
            for row in tb.rows:
                row.cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _docx_blocks(doc, blocks: list[dict[str, Any]]) -> None:
    from docx.shared import Pt
    for b in blocks:
        t = b["t"]
        if t == "p":
            hd.para(doc, b["text"], size=BODY_PT)
        elif t == "h":
            hd.heading(doc, b["text"], 3)
        elif t in ("bullets", "numbered"):
            hd.bullets(doc, b["items"], numbered=t == "numbered")
            for p in doc.paragraphs[-len(b["items"]):]:
                for r in p.runs:
                    r.font.size = Pt(BODY_PT)
        elif t == "table":
            rows = b["rows"]
            total = bool(rows) and str(rows[-1][0]).startswith("**")
            clean = [[_plain(c) for c in r] for r in rows]
            tb = hd.table(doc, b["header"], clean, widths=b.get("widths"), size=b.get("size") or 7.5, total_row=total, gap=False)
            _fit(tb, clean, b.get("widths"))
            if b.get("source"):
                hd.note(doc, f"Source: {b['source']}")
            else:
                doc.add_paragraph().paragraph_format.space_after = Pt(1)
        elif t == "kv":
            _kv(doc, b["pairs"], b.get("cols", 2))
        elif t == "sign":
            tb = hd.table(doc, ["Role", "Name", "Signature", "Date"], [[r, n or "", "\n\n", ""] for r, n in b["rows"]], size=8.5, gap=True)
            _fit(tb, [], [5.0, 4.5, 4.9, 3.0])
        elif t == "box":
            _box(doc, b["title"], b.get("lines", 4))


def docx(paper: dict[str, Any], kind: str = "Credit proposal", generated_by: str = "", parts: list[dict[str, Any]] | None = None,
         subtitle: str = "") -> tuple[bytes, str]:
    from docx.shared import Cm
    meta = _meta(paper, kind, generated_by)
    details = [("Request", str(paper.get("request_kind") or "").capitalize() or "–"),
               ("Total limits", f"{paper['currency']} {paper['totals']['existing']:,.1f}m → {paper['currency']} {paper['totals']['proposed']:,.1f}m")]
    doc = hd.document(meta, kind=kind, subtitle=subtitle or f"{paper['customer']} · {paper['proposal_id']}", details=details)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.8)
    for pi, part in enumerate(_parts(paper, parts)):
        if pi:
            hd.page_break(doc)
        if part.get("title"):
            hd.heading(doc, part["title"], 1)
        for i, s in enumerate(part["sections"], start=1):
            hd.heading(doc, (f"{i}. " if part.get("numbered", True) else "") + s["heading"], 2)
            _docx_blocks(doc, s["blocks"])
    safe = re.sub(r"[^\w-]+", "-", f"{kind}-{paper['proposal_id']}-{paper['customer']}")[:90].strip("-")
    return hd.to_bytes(doc), f"{safe}.docx"


# ------------------------------------------------------------------------------------------------ PDF

def _esc(t: Any) -> str:
    s = str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)


def pdf(paper: dict[str, Any], kind: str = "Credit proposal", generated_by: str = "", parts: list[dict[str, Any]] | None = None) -> tuple[bytes, str]:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    meta = _meta(paper, kind, generated_by)
    acc, band, muted, rule = (colors.HexColor("#" + x) for x in (ACCENT, BAND, MUTED, RULE))
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=BODY_PT, leading=BODY_PT * 1.25, spaceAfter=3)
    small = ParagraphStyle("s", parent=body, fontSize=7, leading=8.6, spaceAfter=0)
    note = ParagraphStyle("n", parent=body, fontSize=7.5, textColor=muted, fontName="Helvetica-Oblique", spaceAfter=5)
    h1 = ParagraphStyle("h1", parent=body, fontName="Helvetica-Bold", fontSize=15, textColor=acc, spaceBefore=6, spaceAfter=8)
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=12, textColor=acc, spaceBefore=10, spaceAfter=4)
    h3 = ParagraphStyle("h3", parent=body, fontName="Helvetica-Bold", fontSize=10, textColor=acc, spaceBefore=6, spaceAfter=3)
    head = ParagraphStyle("th", parent=small, textColor=colors.white, fontName="Helvetica-Bold")
    W = A4[0] - 3.6 * cm

    def tbl(header, rows, widths=None, size=7.5):
        st = ParagraphStyle("c", parent=small, fontSize=size, leading=size * 1.22)
        sth = ParagraphStyle("ch", parent=head, fontSize=size, leading=size * 1.22)
        cell = lambda c: _esc(c).replace("    ", "&nbsp;&nbsp;&nbsp;&nbsp;")  # noqa: E731 - sub-limit indent
        data = [[Paragraph(_esc(h), sth) for h in header]] + [[Paragraph(cell(c), st) for c in r] for r in rows]
        cw = None
        if widths:
            tot = sum(widths)
            cw = [W * w / tot for w in widths]
        t = Table(data, colWidths=cw, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), acc), ("GRID", (0, 0), (-1, -1), 0.25, rule), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                 ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
                 ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5)]
        for i in range(2, len(data), 2):
            style.append(("BACKGROUND", (0, i), (-1, i), band))
        t.setStyle(TableStyle(style))
        return t

    story: list[Any] = []
    # cover block
    bar = Table([[Paragraph("<b>CREDITPROBE</b>&nbsp;&nbsp;&nbsp;Credit risk intelligence", ParagraphStyle("bar", parent=body, textColor=colors.white, fontSize=11))]],
                colWidths=[W])
    bar.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), acc), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    story += [bar, Spacer(1, 1.2 * cm), Paragraph(f"<b>{_esc(kind.upper())}</b>", ParagraphStyle("k", parent=body, textColor=acc, fontSize=11)),
              Paragraph(_esc(meta.title), ParagraphStyle("t", parent=body, fontSize=22, leading=26, fontName="Helvetica-Bold", spaceAfter=10)),
              Paragraph(_esc(f"{paper['customer']} · {paper['proposal_id']}"), ParagraphStyle("st", parent=body, fontSize=13, textColor=muted, spaceAfter=24))]
    cov = [("Book", meta.book_label), ("Period", meta.period), ("Scope", meta.subject), ("Prepared by", meta.generated_by or "CreditProbe"),
           ("Generated", meta.when), ("Reference", meta.reference), ("Classification", meta.classification)]
    story.append(tbl(["Item", "Value"], [[k, v] for k, v in cov if v], widths=[4.5, 12.5], size=9))
    story.append(PageBreak())

    def blocks_flow(bs):
        out = []
        for b in bs:
            t = b["t"]
            if t == "p":
                out.append(Paragraph(_esc(b["text"]), body))
            elif t == "h":
                out.append(Paragraph(_esc(b["text"]), h3))
            elif t in ("bullets", "numbered"):
                out.append(ListFlowable([ListItem(Paragraph(_esc(x), body), leftIndent=12) for x in b["items"]],
                                        bulletType="1" if t == "numbered" else "bullet", start=1 if t == "numbered" else None, leftIndent=12,
                                        bulletFontSize=BODY_PT if t == "numbered" else 6))
            elif t == "table":
                out.append(tbl(b["header"], b["rows"], b.get("widths"), b.get("size") or 7.5))
                out.append(Paragraph(_esc("Source: " + b["source"]) if b.get("source") else "", note))
            elif t == "kv":
                pairs = b["pairs"]
                per = b.get("cols", 2)
                rows = []
                for i in range(0, len(pairs), per):
                    r = []
                    for k, v in pairs[i:i + per]:
                        r += [f"**{k}**", v]
                    rows.append(r + [""] * (per * 2 - len(r)))
                t2 = tbl([""] * (per * 2), rows, [3.2, 5.3] * per, 7.5)
                out.append(t2)
                out.append(Spacer(1, 4))
            elif t == "sign":
                out.append(tbl(["Role", "Name", "Signature", "Date"], [[r, n or "", "", ""] for r, n in b["rows"]] , [5.0, 4.5, 4.5, 3.0], 8.5))
                out.append(Spacer(1, 6))
            elif t == "box":
                bx = Table([[Paragraph(f"<b>{_esc(b['title'])}</b>", body)]] + [[""] for _ in range(b.get("lines", 4))]
                           + [[Paragraph("Signature: ____________________   Name: ____________________   Date: __________", small)]], colWidths=[W])
                bx.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, rule)]))
                out.append(KeepTogether([bx, Spacer(1, 6)]))
        return out

    for pi, part in enumerate(_parts(paper, parts)):
        if pi:
            story.append(PageBreak())
        if part.get("title"):
            story.append(Paragraph(_esc(part["title"]), h1))
        for i, s in enumerate(part["sections"], start=1):
            story.append(Paragraph(_esc((f"{i}. " if part.get("numbered", True) else "") + s["heading"]), h2))
            story += blocks_flow(s["blocks"])

    class NumberedCanvas(rl_canvas.Canvas):
        def __init__(self, *a, **kw):
            super().__init__(*a, **kw)
            self._pages = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            n = len(self._pages)
            for i, st in enumerate(self._pages, start=1):
                self.__dict__.update(st)
                if i > 1:
                    self.setFont("Helvetica-Bold", 7.5)
                    self.setFillColor(acc)
                    self.drawString(1.8 * cm, A4[1] - 1.2 * cm, "CREDITPROBE")
                    self.setFont("Helvetica", 7.5)
                    self.setFillColor(muted)
                    self.drawString(3.9 * cm, A4[1] - 1.2 * cm, f"{meta.title} · {meta.book_label} · {meta.period}"[:120])
                    self.drawRightString(A4[0] - 1.8 * cm, A4[1] - 1.2 * cm, meta.classification)
                    self.drawString(1.8 * cm, 1.0 * cm, f"CreditProbe · {meta.generated_line()}"[:110])
                    self.drawRightString(A4[0] - 1.8 * cm, 1.0 * cm, f"Page {i} of {n}")
                super().showPage()
            super().save()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.8 * cm, bottomMargin=1.6 * cm,
                            title=meta.title, author=meta.generated_by or "CreditProbe", subject=meta.scope_line(), keywords="creditprobe-house-style-1")
    doc.build(story, canvasmaker=NumberedCanvas)
    safe = re.sub(r"[^\w-]+", "-", f"{kind}-{paper['proposal_id']}-{paper['customer']}")[:90].strip("-")
    return buf.getvalue(), f"{safe}.pdf"


def pdf_pages(data: bytes) -> int:
    from pypdf import PdfReader
    return len(PdfReader(io.BytesIO(data)).pages)


# ------------------------------------------------------------------------------------------------ Markdown (exemplars, previews)

def markdown(paper: dict[str, Any], keys: list[str] | None = None, title: str = "") -> str:
    """The paper as Markdown with '#' headings: the format the house-style corpus reads (exemplars) and the dialog preview."""
    out = [f"# {title or paper['title']}", ""]
    for i, s in enumerate([s for s in paper["sections"] if keys is None or s["key"] in keys], start=1):
        out += [f"## {i}. {s['heading']}", ""]
        for b in s["blocks"]:
            t = b["t"]
            if t == "p":
                out += [b["text"], ""]
            elif t == "h":
                out += [f"**{b['text']}**", ""]
            elif t == "bullets":
                out += [f"- {x}" for x in b["items"]] + [""]
            elif t == "numbered":
                out += [f"{k}) {x}" for k, x in enumerate(b["items"], start=1)] + [""]
            elif t == "table":
                out += ["| " + " | ".join(map(str, b["header"])) + " |", "|" + "---|" * len(b["header"])]
                out += ["| " + " | ".join(_plain(c).strip() for c in r) + " |" for r in b["rows"]]
                out += ([f"Source: {b['source']}"] if b.get("source") else []) + [""]
            elif t == "kv":
                out += ["| Item | Value |", "|---|---|"] + [f"| {k} | {v} |" for k, v in b["pairs"]] + [""]
            elif t == "sign":
                out += ["| Role | Name | Signature | Date |", "|---|---|---|---|"] + [f"| {r} | {n} |  |  |" for r, n in b["rows"]] + [""]
            elif t == "box":
                out += [f"{b['title']}: ____________________", ""]
    return "\n".join(out).rstrip() + "\n"
