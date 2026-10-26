"""The proposal-check demo pack (PLAN_R14 p, K31 point 3): a fictional half-baked RM proposal for R14-RWQ in the bank form (Word)
with exactly 12 planted defects, the answer key, and the corrected version of the same proposal.

    .venv/bin/python -m backend.dealdesk.proposal_check.demo_pack      # writes data/dealdesk/demo_proposals/

Figures are read from Data Domains (R14-RWQ, published by ops/r14_p_publish_demo.py), so the planted deviations are exact:
the RM's turnover is 40% above the account's, the receivables include the related-party balance, and so on.
"""
from __future__ import annotations

import json
import pathlib
from typing import Any

from backend.dealdesk.proposal_check.demo_deal import ANALYST, NAME, PID, PRIOR_REF, RM

OUT = pathlib.Path(__file__).resolve().parents[3] / "data" / "dealdesk" / "demo_proposals"
HALF, FIXED = "R14-RWQ_RM_proposal_half_baked.docx", "R14-RWQ_RM_proposal_corrected.docx"

ANSWER_KEY = [
    (1, "undisclosed_returned_cheques", 8, "Relationship with the bank", "Says no cheques were returned in the last 12 months",
     "4 returned cheques in the 12 months to Sep 2026 (SAR 3.3m: 10 Nov 2025, 10 Feb, 10 May, 10 Aug 2026)", "dd_returned_cheque"),
    (2, "turnover_overstated", 8, "Relationship with the bank", "Claims account turnover of SAR {claim}m, 40% above the account",
     "Credit turnover SAR {turn}m over Oct 2025 – Sep 2026", "dd_account_conduct"),
    (3, "related_party_receivables_as_trade", 15, "Financial synopsis", "Trade receivables of SAR 274.0m include SAR 64.0m due from a related party",
     "FY2025: trade receivables SAR 210.0m + due from Ruwaq Real Estate Development Company SAR 64.0m", "dd_financials_detail"),
    (4, "ubo_bureau_missing", 11, "Credit bureau", "Says the shareholders' bureau reports are clear",
     "No bureau report of the UBO Sultan Hamad Al-Ruwaqi on file (requested 13 Oct 2026, CFR-2026-0611)", "bureau_entity_report"),
    (5, "unperfected_guarantee", 4, "Security and support", "Says all security is perfected",
     "Corporate guarantee of Ruwaq Holding Company (SAR 190.0m) pending; deadline 30 Jun 2026 passed", "dd_security"),
    (6, "stale_valuation", 4, "Security and support", "Calls the land valuation recent and gives no date",
     "Valuation dated 15 Apr 2025, 18 months old", "dd_security"),
    (7, "group_exposure_total_wrong", 6, "Group exposure and references", "Group exposure table totals SAR 250.0m",
     "Rows add up to SAR 226.0m (dd_group_exposure SAR 226.0m)", "dd_group_exposure"),
    (8, "covenant_without_numbers", 5, "Terms and conditions", "Covenants 'standard as per the bank's practice' with no thresholds",
     "Policy 4.2–4.4: DSCR ≥ 1.25x, debt / EBITDA ≤ 3.5x, current ratio ≥ 1.10x", "dd_policy_clause"),
    (9, "missing_project_status", 9, "Status of financed projects", "Section missing",
     "2 financed projects; Al Bariq Residences 72% done vs 88% planned, 243-day EOT not formalised, 2 certificates pending", "dd_project"),
    (10, "vague_sales_terms", 7, "Background and ownership", "Sales 'on usual market terms'",
     "Monthly certificates paid in 45 days; 10% retention; 10% advance", "dd_business_profile"),
    (11, "prior_condition_misstated", 5, "Terms and conditions", f"Says all conditions of {PRIOR_REF} are complied with",
     "2 of 3 conditions not complied: quarterly project reports; perfection of the parent guarantee", "dd_condition"),
    (12, "pricing_exception_undisclosed", 17, "Exceptions and other observations", "Says no policy exceptions",
     "Lifetime RAROC {raroc}% vs {hurdle}% hurdle at 150 bps on the new line", "dd_deal"),
]


def _facts() -> dict[str, Any]:
    from backend.dealdesk.house_style.bankform import ctx
    from backend.dealdesk.proposal_check import engine as E, rules
    D = ctx.load(PID)
    C = rules.Ctx(doc={"sections": []}, D=D, as_of=E._as_of(), deal=E._deal_row(PID))  # noqa: SLF001
    turn = rules._turnover_12m(C)[0]  # noqa: SLF001
    pr = rules._pricing(C)  # noqa: SLF001
    return {"turn": round(turn, 1), "claim": round(turn * 1.4, 1), "raroc": round(pr["raroc"], 1), "hurdle": round(pr["hurdle"], 1)}


# ------------------------------------------------------------------------------------------------ Word (bank form, RM's own paper)

def _doc():
    from docx import Document
    from docx.shared import Cm, Pt
    d = Document()
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Arial", Pt(9)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(1.8)
        s.top_margin = s.bottom_margin = Cm(1.5)
        s.header.paragraphs[0].text = f"CREDIT PROPOSAL · {NAME} · CIF-0074412 · Corporate Banking — Dammam Corporate Centre"
        s.header.paragraphs[0].runs[0].font.size = Pt(7)
    return d


def _h(d, text: str) -> None:
    from docx.shared import Pt
    p = d.add_heading(text, level=2)
    for r in p.runs:
        r.font.size, r.font.name = Pt(10.5), "Arial"


def _p(d, text: str, bold: bool = False) -> None:
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold = bold


def _t(d, rows: list[list[str]], widths: list[float] | None = None) -> None:
    from docx.shared import Cm, Pt
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = d.styles["Table Grid"]
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            cell = t.rows[i].cells[j]
            cell.text = ""
            run = cell.paragraphs[0].add_run(c)
            run.font.size = Pt(8)
            run.bold = i == 0 or c.startswith("Total")
            if widths:
                cell.width = Cm(widths[j])
    d.add_paragraph()


def write(fixed: bool, f: dict[str, Any]) -> bytes:
    d = _doc()
    _p(d, "CREDIT PROPOSAL — CORPORATE BANKING GROUP", bold=True)
    _t(d, [["Branch / unit", "Dammam Corporate Centre", "Customer", NAME],
           ["Date account opened", "6 Mar 2011", "CIF / base no.", "CIF-0074412"],
           ["Group no.", "GRP-1377", "Group", "Ruwaq Holding Group"],
           ["Related / connected party", "No", "Date", "12 Oct 2026"],
           ["AO / RO", f"{RM} / Bader Al-Qahtani", "Purpose code", "P34 Contracting guarantees"],
           ["Request", "[ ] New [ ] Existing [x] Renewal [x] Additional [ ] Amendment", "Review date", "31 Oct 2027"],
           ["Industry code", "F4100", "Collateral code", "CG / RE / ASG"],
           ["Rating existing / proposed", "BB / BB", "RM", RM]])
    _h(d, "2. Action requested")
    _p(d, f"Business requests approval, as at 12 Oct 2026, to renew the total limits of SAR 150.0m approved by {PRIOR_REF} on 22 Oct 2025 and to increase "
          "them to SAR 190.0m: renew the overdraft of SAR 20.0m and project limit 1 of SAR 60.0m, reduce project limit 2 from SAR 40.0m to SAR 30.0m, "
          "increase the general LG line from SAR 30.0m to SAR 40.0m and approve a new line of SAR 40.0m of finance against certified payment "
          "certificates; group exposure cap from SAR 200.0m to SAR 240.0m.")
    if fixed:
        _p(d, f"Of the 3 conditions of {PRIOR_REF}, 2 are not complied (see section 5); business asks for their completion as conditions prior to limit loading.")
    else:
        _p(d, f"All conditions of the last approval ({PRIOR_REF}) have been complied with.")
    _h(d, "3. Facilities")
    _t(d, [["Limit / sub-limit", "Existing", "Outstanding", "Proposed", "Cash margin", "Pricing / tenor"],
           ["Overdraft", "20.0", "17.5", "20.0", "–", "3M SAIBOR + 2.25%, floor 6.00%; 12 months"],
           ["Project limit 1: Al Bariq Residences", "60.0", "52.0", "60.0", "–", "To completion + 12 months"],
           ["  of which performance and advance-payment LGs", "60.0", "52.0", "60.0", "10%", "Commission 0.90% p.a."],
           ["Project limit 2: Eastern Schools Package 2", "40.0", "33.0", "30.0", "–", "To completion + 12 months"],
           ["  of which performance LGs", "40.0", "33.0", "30.0", "10%", "Commission 0.90% p.a."],
           ["General LG line", "30.0", "25.5", "40.0", "10%", "Commission 0.85% p.a.; up to 24 months"],
           ["Finance against certified payment certificates (new)", "0.0", "0.0", "40.0", "–", "3M SAIBOR + 1.50%, floor 5.50%; 120 days per drawing"],
           ["Total", "150.0", "128.0", "190.0", "", ""]], [5.6, 1.6, 1.8, 1.6, 1.6, 5.2])
    _p(d, "Study and documentation fee of SAR 0.08m (20 bps on the new line) payable on approval, as at 12 Oct 2026.")
    _h(d, "4. Security and support")
    _p(d, "Personal guarantee of Sultan Hamad Al-Ruwaqi for SAR 190.0m, net worth SAR 410.0m per statement dated 31 Mar 2026.")
    if fixed:
        _p(d, "Corporate guarantee of Ruwaq Holding Company for SAR 190.0m (net worth SAR 260.0m at 31 Dec 2025) is not yet perfected: the guarantor's board "
              "resolution is pending and the perfection deadline of 30 Jun 2026 has passed; perfection is a condition prior to limit loading.")
        _p(d, "First-degree mortgage over a 24,000 sqm plot in Dammam 2nd Industrial City, valued at SAR 95.0m by Coastline Valuers on 15 Apr 2025 "
              "(18 months old at 12 Oct 2026); a fresh valuation by a bank-approved valuer is a condition prior to limit loading.")
        _p(d, "Assignment of Al Bariq Residences proceeds (SAR 210.0m) acknowledged by the employer; undated security cheque for SAR 190.0m pending, due by 31 Dec 2026.")
    else:
        _p(d, "Corporate guarantee of Ruwaq Holding Company for SAR 190.0m, net worth SAR 260.0m at 31 Dec 2025.")
        _p(d, "Mortgage over land in Dammam 2nd Industrial City valued at SAR 95.0m per a recent valuation.")
        _p(d, "Assignment of Al Bariq Residences proceeds (SAR 210.0m) and an undated security cheque for SAR 190.0m.")
        _p(d, "All security is perfected and in place.")
    _h(d, "5. Terms and conditions")
    _p(d, "1. Finance against certified payment certificates at 80% of the certified amount; the last certificate of each project is not financed.")
    _p(d, "2. Employer payments received in the collection account with the bank and applied first to the related drawing.")
    _p(d, "3. Audited statements within 120 days of year end; interim statements within 60 days of the half year.")
    _p(d, "4. No LGs or funded drawings in favour of related parties.")
    if fixed:
        _p(d, "5. Covenants: DSCR not below 1.25x and total debt / EBITDA not above 3.5x, tested annually on the audited statements within 120 days of year end; "
              "current ratio not below 1.10x (Credit Policy 4.2–4.4).")
        _p(d, f"Status of existing conditions of {PRIOR_REF} (22 Oct 2025): 2 of the 3 are not complied.")
        _t(d, [["Item of " + PRIOR_REF, "Owner", "Status at 30 Sep 2026"],
               ["Quarterly project status reports with certified payment certificates", "RM", "Not complied: last report 30 Jun 2026"],
               ["Perfection of the parent corporate guarantee", "Credit administration", "Not complied: board resolution pending"],
               ["Audited statements within 120 days of year end", "RM", "Complied on 21 Apr 2026"]])
    else:
        _p(d, "5. Standard covenants as per the bank's practice.")
    _h(d, "6. Group exposure and references")
    total = "226.0" if fixed else "250.0"
    _t(d, [["Entity", "Total existing", "Total proposed", "Outstanding"],
           [NAME, "150.0", "190.0", "128.0"], ["Ruwaq Ready-Mix Company", "22.0", "22.0", "19.5"],
           ["Ruwaq Real Estate Development Company", "14.0", "14.0", "14.0"], ["Ruwaq Holding Company", "0.0", "0.0", "0.0"],
           ["Total", "186.0", total, "161.5"]])
    _p(d, f"References: {PRIOR_REF} of 22 Oct 2025 (renewal at SAR 150.0m, group cap SAR 200.0m); CC-2024-163 of 17 Oct 2024 (general LG line raised to SAR 30.0m).")
    _h(d, "7. Background and ownership")
    _p(d, "Trade licence CR-2050-7781: incorporated on 15 Nov 2009 in Dammam as an LLC; building and infrastructure contracting (grade 2); licence valid to 31 Mar 2027.")
    _t(d, [["Shareholder of Ruwaq Holding Company", "Born", "Nationality", "Share", "UBO"],
           ["Sultan Hamad Al-Ruwaqi", "1961", "Saudi", "70%", "Yes"], ["Huda Sultan Al-Ruwaqi", "1988", "Saudi", "30%", "Yes"]])
    _p(d, "Group entities: Ruwaq Holding Company (parent), Ruwaq Ready-Mix Company (concrete supply, 30% of purchases), Ruwaq Real Estate Development Company "
          "(residential development). Activity mix FY2025: residential buildings 61% (FY2024 58%), education and public buildings 39%; 1,180 staff.")
    if fixed:
        _p(d, "Sales terms: monthly certified payment certificates paid in 45 days; 10% retention released 12 months after handover; 10% advance recovered pro rata. "
              "Purchase terms: steel and cement on 60–90 days open account, as per business.")
    else:
        _p(d, "Sales are on usual market terms with the employers. The group is well established in the Eastern Province.")
    _h(d, "8. Relationship with the bank")
    _p(d, "The relationship dates from 6 Mar 2011 (CIF-0074412); limits of SAR 150.0m approved by " + PRIOR_REF + " on 22 Oct 2025, utilisation 78–92% of limits in 2026.")
    if fixed:
        _p(d, f"Credit turnover of SAR {f['turn']:.1f}m in the 12 months Oct 2025 to Sep 2026, 1.99x the existing limits; average balance SAR 8.3m.")
        _p(d, "4 cheques returned in the 12 months to Sep 2026 for SAR 3.3m (10 Nov 2025, 10 Feb 2026, 10 May 2026, 10 Aug 2026); business attributes them to "
              "delayed employer payments on Al Bariq Residences; maximum past due 12 days in Jul 2026.")
    else:
        _p(d, f"Account turnover through our accounts of SAR {f['claim']:.1f}m in the last 12 months, reflecting the group's growing business.")
        _p(d, "No cheques were returned in the last 12 months. The relationship has been satisfactory.")
    if fixed:
        _h(d, "9. Status of financed projects")
        _t(d, [["Project", "Employer", "Contract (SAR m)", "Done / planned", "Completion", "Pending certificates"],
               ["Al Bariq Residences", "Al Bariq Developments Company", "438.0", "72% / 88%", "31 Jul 2026 → 31 Mar 2027", "2 (SAR 21.0m)"],
               ["Eastern Schools Package 2", "Eastern Schools Projects Company", "260.0", "46% / 48%", "30 Jun 2027", "0"]])
        _p(d, "Al Bariq Residences is past its planned completion of 31 Jul 2026 at 72% work done; an extension of time of 243 days is requested and not formalised; "
              "the last payment was received on 20 Jun 2026. Eastern Schools Package 2 is within 2 pp of plan with receipts current.")
    _h(d, "10. Central bank risk bureau (CBRBR)")
    _t(d, [["Entity", "Funded limit / o/s", "Non-funded limit / o/s", "Banks", "Classification"],
           [NAME, "110.0 / 92.0", "210.0 / 168.0", "3", "Regular"], ["Ruwaq Ready-Mix Company", "40.0 / 31.0", "10.0 / 8.0", "2", "Regular"],
           ["Ruwaq Real Estate Development Company", "35.0 / 31.0", "0.0 / 0.0", "1", "Watch"]])
    _p(d, "CBRBR report dated 30 Sep 2026; the real estate sister company is on the watch list of its bank.")
    _h(d, "11. Credit bureau (SIMAH)")
    if fixed:
        _p(d, "SIMAH report of the company dated 6 Oct 2026: 9 facilities, SAR 286.0m outstanding, no overdue, 4 returned cheques (SAR 3.3m), score 602.")
        _p(d, "The bureau report of the UBO, Sultan Hamad Al-Ruwaqi, is not provided: requested on 13 Oct 2026 (CFR-2026-0611); obtaining it is a condition prior to limit loading.")
        _p(d, "Ruwaq Real Estate Development Company: SAR 2.4m overdue at 47 days and 1 court case (SAR 1.9m), report of 6 Oct 2026.")
    else:
        _p(d, "SIMAH reports of the company dated 6 Oct 2026 show 9 facilities and SAR 286.0m outstanding with no overdue.")
        _p(d, "The SIMAH reports of the shareholders are clear with no adverse remarks.")
    _h(d, "12. Other banks' facilities")
    _t(d, [["Bank", "Facilities", "Limit", "Outstanding (declared / bureau)"],
           ["Coastal Commerce Bank", "PB; APG; OD", "90.0", "71.0 / 72.0"], ["Red Sea Investment Bank", "LC; STL", "40.0", "28.0 / 29.5"]])
    _p(d, "Declared facilities agree with the bureau within SAR 1.5m at 30 Sep 2026.")
    _h(d, "13. New project details and cash flow")
    _p(d, "Not applicable: no new project is financed in this request (12 Oct 2026); the new line finances certified certificates of the 2 existing projects.")
    _h(d, "14. Financials")
    _t(d, [["SAR m", "FY2023", "FY2024", "FY2025", "6M2026"],
           ["Revenue", "300.0", "330.0", "360.0", "171.0"], ["Net profit", "9.8", "9.9", "9.2", "3.6"], ["EBITDA", "33.0", "36.0", "38.0", "18.2"],
           ["Bank debt", "120.0", "132.0", "150.0", "154.0"], ["Net worth", "102.0", "110.0", "118.0", "121.0"], ["Current ratio", "1.38x", "1.39x", "1.43x", "1.42x"]])
    _p(d, "Audited by Harbour Gate Auditors (bank panel), clean opinion for FY2023–FY2025; 6M2026 management accounts as per business.")
    _h(d, "15. Financial synopsis")
    _p(d, "Revenue rose to SAR 360.0m in FY2025 (PY SAR 330.0m, +9.1%); net profit fell to SAR 9.2m (PY SAR 9.9m).")
    if fixed:
        _p(d, "Trade receivables of SAR 210.0m at 31 Dec 2025 (PY SAR 178.0m), 58.0% of current assets; due from Ruwaq Real Estate Development Company "
              "(related party) SAR 64.0m, 210 days past due as per the ageing, shown separately.")
        _p(d, "Netting the related-party balance, adjusted working capital is SAR 44.0m (reported SAR 108.0m) and the adjusted current ratio 1.17x (reported 1.43x).")
    else:
        _p(d, "Trade receivables of SAR 274.0m at 31 Dec 2025 (PY SAR 230.0m), 75.7% of current assets, reflecting the growth in certified works.")
        _p(d, "Working capital of SAR 108.0m and a current ratio of 1.43x at 31 Dec 2025.")
    _h(d, "16. Effects of the current geopolitical and market situation")
    _p(d, "Business conveys steel prices rose 11% in 2026 and the contracts carry no escalation clause; with Brent high in Oct 2026 input and freight costs remain elevated.")
    _h(d, "17. Exceptions and other observations")
    if fixed:
        _p(d, f"1. Pricing exception: the new line at 3M SAIBOR + 150 bps gives a lifetime RAROC of {f['raroc']:.1f}% against the {f['hurdle']:.1f}% hurdle; "
              "business's reason: the line replaces another bank's facility and brings the employer collections of Al Bariq Residences to the bank.")
        _p(d, "2. Ruwaq Real Estate Development Company, a related party, owes SAR 64.0m to the obligor and is on its bank's watch list (30 Sep 2026).")
    else:
        _p(d, "No policy exceptions.")
    _h(d, "18. Key strengths")
    _p(d, "1. Relationship since 6 Mar 2011 with limits renewed in 2024 and 2025 (CC-2024-163, " + PRIOR_REF + ").")
    _p(d, "2. Revenue growth of 9.1% in FY2025 with a backlog of SAR 655.0m (1.8x FY2025 revenue).")
    _p(d, "3. Personal guarantee of the UBO with a net worth of SAR 410.0m (31 Mar 2026).")
    _h(d, "19. Recommendation")
    _p(d, f"{RM} recommends approval of total limits of SAR 190.0m and a group exposure cap of SAR 240.0m as set out in sections 2 to 5, as at 12 Oct 2026.")
    if fixed:
        _p(d, "Prior to limit loading: perfection of the parent guarantee; UBO bureau report; fresh valuation of the Dammam plot; the 2 outstanding project reports.")
    _h(d, "20. Signatures and approval")
    _t(d, [["Role", "Name", "Signature", "Date"], ["Relationship manager", RM, "", "12 Oct 2026"], ["Head of corporate banking, Dammam", "", "", ""],
           ["Credit analyst", ANALYST, "", ""]])
    import io
    b = io.BytesIO()
    d.save(b)
    return _stable(b.getvalue())


def _stable(data: bytes) -> bytes:
    """Fixed core properties and zip timestamps, so the files are byte-stable."""
    import io
    import zipfile
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in sorted(src.infolist(), key=lambda i: i.filename):
            body = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                import re
                body = re.sub(rb"<dcterms:(created|modified)([^>]*)>[^<]*</dcterms:\1>", rb"<dcterms:\1\2>2026-10-12T08:00:00Z</dcterms:\1>", body)
                body = re.sub(rb"<cp:lastModifiedBy>[^<]*</cp:lastModifiedBy>", b"<cp:lastModifiedBy>" + RM.encode() + b"</cp:lastModifiedBy>", body)
                body = re.sub(rb"<dc:creator>[^<]*</dc:creator>", b"<dc:creator>" + RM.encode() + b"</dc:creator>", body)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 10, 12, 8, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, body)
    return out.getvalue()


def answer_key(f: dict[str, Any]) -> dict[str, Any]:
    return {"proposal_id": PID, "customer": NAME, "rm": RM, "document": HALF, "corrected": FIXED, "planted_defects": 12,
            "note": "Fictional. Every defect is a claim or omission in the RM's half-baked proposal that the bank's data contradicts; "
                    "the proposal check must catch all 12, deterministically, in under 60 seconds (tests/v5/proposal_check).",
            "defects": [{"n": n, "defect_type": t, "section": s, "heading": h, "planted": a.format(**f), "truth": b.format(**f), "relation": r}
                        for n, t, s, h, a, b, r in ANSWER_KEY]}


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    f = _facts()
    (OUT / HALF).write_bytes(write(False, f))
    (OUT / FIXED).write_bytes(write(True, f))
    (OUT / "answer_key.json").write_text(json.dumps(answer_key(f), indent=1, ensure_ascii=False) + "\n")
    print("demo pack:", HALF, FIXED, "answer_key.json", f)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
