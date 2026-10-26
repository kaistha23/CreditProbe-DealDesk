"""Read an RM proposal into sections (PLAN_R14 p, K31).

    doc = read_docx(data, name)            # Word (bank form or memo): headings -> anatomy sections via the bankform coverage matcher
    doc = from_record(pid)                  # no document: the RM's submission record in dd_deal (purpose, ask, security, covenants)
    doc["sections"]  -> [{key, n, heading, paras: [{i, text}], tables: [[[cell]]], words}]
    doc["header"]    -> {label: value} from the first two-column table (the bank-form header block), when present

Deterministic: the same bytes give the same structure. Paragraph indices (`i`) are document-order indices into
`Document.paragraphs`, so the annotator can anchor comments to the exact paragraph.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
from typing import Any

from backend.dealdesk.house_style.bankform import coverage, registry

HEAD_RE = re.compile(r"^\s*(\d{1,2})[.)]\s+\S")


def _is_heading(p: Any) -> bool:
    st = (p.style.name if p.style is not None else "") or ""
    t = p.text.strip()
    if not t or len(t) > 90:
        return False
    if st.lower().startswith("heading") or st.lower() == "title":
        return True
    return bool(HEAD_RE.match(t)) and all(r.bold for r in p.runs if r.text.strip()) and bool(p.runs)


def _key_of(heading: str) -> str | None:
    k = coverage.match(heading)
    if k in registry.BY_KEY and registry.BY_KEY[k].part == "anatomy":
        return k
    return None


def _body_items(doc: Any) -> list[tuple[str, Any]]:
    """Paragraphs and tables in document order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    out = []
    for el in doc.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            out.append(("p", Paragraph(el, doc)))
        elif tag == "tbl":
            out.append(("t", Table(el, doc)))
    return out


def _table_rows(t: Any) -> list[list[str]]:
    rows = []
    for r in t.rows:
        cells, seen = [], set()
        for c in r.cells:
            if id(c._tc) in seen:  # merged cells repeat
                continue
            seen.add(id(c._tc))
            cells.append(c.text.strip())
        rows.append(cells)
    return rows


def read_docx(data: bytes, name: str = "proposal.docx") -> dict[str, Any]:
    from docx import Document
    doc = Document(io.BytesIO(data))
    pi = -1  # document-order index into Document.paragraphs (body paragraphs, tables excluded), as the annotator reads it
    sections: list[dict[str, Any]] = []
    cur: dict[str, Any] = {"key": "_preamble", "n": 0, "heading": "", "paras": [], "tables": [], "heading_i": None}
    header: dict[str, str] = {}
    title = ""
    for kind, item in _body_items(doc):
        if kind == "p":
            pi += 1
            t = item.text.strip()
            if not t:
                continue
            if _is_heading(item):
                k = _key_of(t)
                if k is not None:
                    sections.append(cur)
                    cur = {"key": k, "n": registry.BY_KEY[k].n, "heading": t, "paras": [], "tables": [], "heading_i": pi}
                    continue
                if not title and cur["key"] == "_preamble":
                    title = t
            cur["paras"].append({"i": pi, "text": t})
        else:
            rows = _table_rows(item)
            cur["tables"].append(rows)
            if not header and rows and all(len(r) in (2, 4, 6) for r in rows):
                for r in rows:
                    for j in range(0, len(r) - 1, 2):
                        if r[j]:
                            header[r[j].strip().rstrip(":")] = r[j + 1].strip()
    sections.append(cur)
    pre = sections[0]
    if pre["key"] == "_preamble" and header and not any(s["key"] == "header" for s in sections):
        pre["key"], pre["n"], pre["heading"] = "header", 1, "Header block (untitled)"
    secs = [s for s in sections if s["key"] != "_preamble"]
    merged: dict[str, dict[str, Any]] = {}
    for s in secs:  # a section split in two places counts once
        if s["key"] in merged:
            merged[s["key"]]["paras"] += s["paras"]
            merged[s["key"]]["tables"] += s["tables"]
        else:
            merged[s["key"]] = s
    for s in merged.values():
        s["words"] = sum(len(p["text"].split()) for p in s["paras"]) + sum(len(" ".join(c for r in t for c in r).split()) for t in s["tables"])
    return {"name": name, "kind": "docx", "sha": hashlib.sha256(data).hexdigest()[:16], "title": title, "header": header,
            "sections": [merged[k] for k in registry.ANATOMY if k in merged], "paragraphs": len(doc.paragraphs),
            "full_text": "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(" | ".join(c for r in t for c in r) for s in merged.values() for t in s["tables"])}


def from_record(pid: str, deal: dict[str, Any] | None = None) -> dict[str, Any]:
    """The RM's submission as recorded in dd_deal (no Word on file): purpose, ask lines, security, covenants, financials, memo text."""
    from backend.dealdesk.house_style.bankform import source as S
    rows = S.read("dd_deal", proposal_id=pid) if deal is None else [deal]
    if not rows:
        raise KeyError(pid)
    d = rows[-1]
    pj = d.get("proposal_json")
    if isinstance(pj, str) and pj:
        try:
            pj = json.loads(pj)
        except ValueError:
            pj = {}
    pj = pj if isinstance(pj, dict) else {}
    ask, sec, fin = pj.get("ask") or {}, pj.get("security") or {}, (pj.get("financials") or {}).get("years") or []
    secs = []

    def add(key: str, paras: list[str], tables: list[list[list[str]]] | None = None) -> None:
        paras = [p for p in paras if p]
        if not paras and not tables:
            return
        s = registry.BY_KEY[key]
        secs.append({"key": key, "n": s.n, "heading": s.heading, "paras": [{"i": None, "text": p} for p in paras], "tables": tables or [],
                     "heading_i": None, "words": sum(len(p.split()) for p in paras)})
    hdr = {k: str(d.get(c) or "") for k, c in (("Customer", "obligor_name"), ("CIF", "cif"), ("Group", "group_name"), ("RM", "rm"), ("Branch", "branch"))}
    add("header", [f"{k}: {v}" for k, v in hdr.items() if v])
    purpose = pj.get("purpose") or d.get("deal_name") or ""
    if ask:
        add("action", [f"Request: {purpose}." if purpose else "",
                       f"{str(ask.get('facility_type', '')).replace('_', ' ')} of {ask.get('currency', 'SAR')} {ask.get('amount')}m for {ask.get('tenor_years')} years "
                       f"at a margin of {ask.get('margin_bps')} bps."])
        add("facilities", [], [[["Facility", "Amount", "Tenor", "Margin"], [str(ask.get("facility_type", "")), str(ask.get("amount", "")),
                                                                             str(ask.get("tenor_years", "")), str(ask.get("margin_bps", ""))]]])
    col = [f"{c.get('type', '')} valued {c.get('value')}m on {c.get('valuation_date', 'date not given')}" for c in sec.get("collateral") or []]
    gua = [f"{g.get('type', '')} guarantee of {g.get('guarantor', '')}" for g in sec.get("guarantees") or []]
    add("security", col + gua)
    covs = sec.get("covenants") or []
    pkg = sec.get("covenant_package") or ask.get("covenant_package")
    add("terms", [("Covenant package: " + str(pkg) + ".") if pkg else ""] + [f"{c.get('type')} {c.get('threshold')}" for c in covs])
    if fin:
        add("financials", [], [[["Year", "Revenue", "EBITDA", "Net worth"]] + [[str(y.get("year")), str(y.get("revenue")), str(y.get("ebitda")), str(y.get("net_worth"))] for y in fin]])
    if pj.get("memo"):
        add("recommendation", [str(pj.get("memo"))])
    blob = json.dumps({k: d.get(k) for k in sorted(d) if k not in ("published_at",)}, sort_keys=True, default=str)
    return {"name": "submission record", "kind": "record", "sha": hashlib.sha256(blob.encode()).hexdigest()[:16], "title": d.get("deal_name") or pid,
            "header": hdr, "sections": [s for k in registry.ANATOMY for s in secs if s["key"] == k], "paragraphs": 0,
            "full_text": "\n".join(p["text"] for s in secs for p in s["paras"])}
