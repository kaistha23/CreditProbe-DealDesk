"""Round A: the Saqr Al Wadi Contracting sample deal (16 fictional documents) as real files, for UAT of the agentic intake.

The texts are the reference page's SAMPLE (docs/v5/reference/credit-request-agent.html), kept verbatim in
data/dealdesk/agent_sample/sample_texts.json. `build()` turns each into the file its name says: Word (.docx), PDF with a
text layer, a scanned PDF (page images only, no text layer), Excel (.xlsx) and photos / screenshots (.jpg, .png). The
files are committed; run `python -m backend.dealdesk.agent_intake.sample` to regenerate them.

    manifest()        {customer, market, rm, ask, notes, documents: [{name, kind, bytes}]}
    files()           [(name, bytes)] in the pack's order
"""
from __future__ import annotations

import io
import json
import re
from pathlib import Path
from typing import Any

V5 = Path(__file__).resolve().parents[3]
DIR = V5 / "data" / "dealdesk" / "agent_sample"
TEXTS = DIR / "sample_texts.json"
MONO = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"


def texts() -> dict[str, Any]:
    return json.loads(TEXTS.read_text())


def manifest() -> dict[str, Any]:
    t = texts()
    docs = []
    for d in t["documents"]:
        p = DIR / d["name"]
        docs.append({"name": d["name"], "kind": kind(d["name"]), "bytes": p.stat().st_size if p.exists() else 0})
    return {k: t[k] for k in ("customer", "market", "request_type", "relationship", "rm", "ask", "notes", "note")} | {"documents": docs}


def files() -> list[tuple[str, bytes]]:
    return [(d["name"], (DIR / d["name"]).read_bytes()) for d in texts()["documents"] if (DIR / d["name"]).exists()]


def kind(name: str) -> str:
    low = name.lower()
    if "scanned" in low and low.endswith(".pdf"):
        return "scanned PDF"
    return {".pdf": "PDF", ".docx": "Word", ".xlsx": "Excel", ".jpg": "photo", ".png": "image"}.get(Path(low).suffix, "file")


# ------------------------------------------------------------------------------------------------ writers
def _docx(text: str) -> bytes:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(9.5)
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if " | " in ln:  # a run of "a | b | c" lines is a Word table
            rows = []
            while i < len(lines) and " | " in lines[i]:
                rows.append([c.strip() for c in lines[i].split(" | ")])
                i += 1
            n = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=n)
            t.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c in range(n):
                    t.cell(r, c).text = cells[c] if c < len(cells) else ""
            continue
        p = doc.add_paragraph(ln.rstrip())
        p.paragraph_format.space_after = Pt(2)
        if ln.isupper() and len(ln.strip()) > 3:
            for r in p.runs:
                r.bold = True
        i += 1
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


def _pages(text: str) -> list[list[str]]:
    """Split at the printed page ends ('Page 3', '--- Page 2 of 2 ---', a rule of dashes)."""
    out: list[list[str]] = [[]]
    for ln in text.splitlines():
        if re.match(r"^\s*-{3} Page \d+ of \d+ -{3}\s*$", ln):
            if any(x.strip() for x in out[-1]):
                out.append([])
            continue
        if re.match(r"^-{20,}$", ln.strip()):
            out.append([])
            continue
        out[-1].append(ln)
        if re.search(r"\bPage \d+\s*$", ln) and not re.search(r"Page \d+ of", ln):
            out.append([])
    return [p for p in out if any(x.strip() for x in p)]


def _pdf(text: str) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    try:
        pdfmetrics.registerFont(TTFont("Mono", MONO))
        font = "Mono"
    except Exception:  # noqa: BLE001
        font = "Courier"
    pages = _pages(text)
    width = max(len(x) for p in pages for x in p)
    size_w, size_h = landscape(A4) if width > 110 else A4
    fs = max(5.2, min(9.0, (size_w - 64) / (0.602 * width)))
    lh = fs * 1.28
    b = io.BytesIO()
    c = canvas.Canvas(b, pagesize=(size_w, size_h))
    c.setTitle(text.splitlines()[0][:80])
    for p in pages:
        y = size_h - 40
        for ln in p:
            if y < 36:
                c.showPage()
                y = size_h - 40
            c.setFont(font, fs)
            c.drawString(32, y, ln)
            y -= lh
        c.showPage()
    c.save()
    return b.getvalue()


def _image(text: str, photo: bool, seed: int = 1):
    from PIL import Image, ImageDraw, ImageFilter, ImageFont
    lines = [ln for ln in text.splitlines()]
    fs = 22
    font = ImageFont.truetype(MONO, fs)
    cw = font.getlength("M")
    w = int(max(len(x) for x in lines) * cw) + 120
    h = int(len(lines) * fs * 1.45) + 140
    im = Image.new("RGB", (w, h), (252, 251, 247) if photo else (255, 255, 255))
    dr = ImageDraw.Draw(im)
    y = 70
    for ln in lines:
        dr.text((60, y), ln, font=font, fill=(28, 28, 32))
        y += int(fs * 1.45)
    if not photo:
        return im
    # a phone photo: a desk around the page, a slight angle, soft focus
    pad = 80
    bg = Image.new("RGB", (w + 2 * pad, h + 2 * pad), (138, 116, 92))
    bg.paste(im, (pad, pad))
    bg = bg.rotate(1.2 if seed % 2 else -1.0, resample=Image.BICUBIC, expand=False, fillcolor=(138, 116, 92))
    return bg.filter(ImageFilter.GaussianBlur(0.6))


def _scanned_pdf(text: str) -> bytes:
    imgs = [_image("\n".join(p), photo=False).convert("L") for p in _pages(text)]
    from PIL import ImageFilter
    imgs = [x.rotate(0.4, fillcolor=255).filter(ImageFilter.GaussianBlur(0.4)).convert("RGB") for x in imgs]
    b = io.BytesIO()
    imgs[0].save(b, "PDF", resolution=150.0, save_all=True, append_images=imgs[1:])
    return b.getvalue()


def _xlsx(text: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    sheets: list[tuple[str, list[str]]] = []
    head: list[str] = []
    for ln in text.splitlines():
        m = re.match(r"^SHEET (\d+): (.+?)(?: AS AT| \(|$)", ln)
        if m:
            sheets.append((m.group(2).title()[:28], [*head, ln] if not sheets else [ln]))
            continue
        if sheets:
            sheets[-1][1].append(ln)
        else:
            head.append(ln)
    if not sheets:
        sheets = [("Rating sheet", head)]
    for si, (name, lines) in enumerate(sheets):
        sh = ws if si == 0 else wb.create_sheet()
        sh.title = re.sub(r"[\[\]:*?/\\]", "", name) or f"Sheet{si + 1}"
        for r, ln in enumerate(lines, 1):
            cells = [c for c in re.split(r"\s{2,}", ln.strip())] if ln.strip() else []
            for ci, v in enumerate(cells, 1):
                num = v.replace(",", "")
                if re.fullmatch(r"-?\d+(\.\d+)?", num) and len(num) < 16 and not re.fullmatch(r"0\d+", num):
                    cell = sh.cell(row=r, column=ci, value=float(num) if "." in num else int(num))
                    if "," in v:
                        cell.number_format = "#,##0"
                else:
                    cell = sh.cell(row=r, column=ci, value=v)
                if ci == 1 and (ln.isupper() or re.match(r"^\d\. [A-Z ]+", ln.strip())):
                    cell.font = Font(bold=True)
        sh.column_dimensions["A"].width = 44
        for col in "BCDEFGHIJK":
            sh.column_dimensions[col].width = 15
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def make(name: str, text: str, seed: int = 1) -> bytes:
    k = kind(name)
    if k == "Word":
        return _docx(text)
    if k == "PDF":
        return _pdf(text)
    if k == "scanned PDF":
        return _scanned_pdf(text)
    if k == "Excel":
        return _xlsx(text)
    if k in ("photo", "image"):
        im = _image(text, photo=(k == "photo"), seed=seed)
        b = io.BytesIO()
        if name.lower().endswith(".jpg"):
            im.convert("RGB").save(b, "JPEG", quality=88)
        else:
            im.save(b, "PNG", optimize=True)
        return b.getvalue()
    return text.encode()


def build() -> list[Path]:
    out = []
    for i, d in enumerate(texts()["documents"]):
        p = DIR / d["name"]
        p.write_bytes(make(d["name"], d["text"], seed=i))
        out.append(p)
    return out


if __name__ == "__main__":
    for p in build():
        print(f"{p.stat().st_size:>9,}  {p.name}")
