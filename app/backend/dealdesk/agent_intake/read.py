"""Round A, step 1: read and classify each document into the intake schema, reconciled with R16's rule-based reader.

The intake schema (the Credit Request Agent's `intakePrompt`):
    {doc_type, checklist_ids, entity, period_or_date, expiry_date, summary, key_figures: [{label, value, page}], observations}

Two readers, one schema:
  rule reader   R16's reader (backend/dealdesk/intake: Word, PDF text, OCR for photos and scans; its classifier and fields)
                plus Round A's title rules for the documents R16 does not know (rating sheet, relationship profile, approval
                and covenant certificate, group structure, account conduct report) and a line reader for key figures
                (copied exactly from the text) and credit observations. Deterministic; the reader when there is no model key.
  model reader  the model reads the document's text (page by page) or its page images (photos, scanned PDFs) through the
                provider's image support, and answers in the schema (one tool call).

`reconcile(rule, model, text)`: classification agree → confirmed; the model finds more → offered (with the page); they
disagree → flagged. A key figure the model reports must be in the document's text (its numbers) or it is rejected; an
expiry date the model reports must be printed in the document.
"""
from __future__ import annotations

import base64
import csv
import datetime as dt
import io
import re
from pathlib import Path
from typing import Any

from backend.dealdesk.agent_intake import checklist as CK

TEXT_LIMIT = 60000
MAX_IMAGES = 3
IMAGE_EXT = (".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp")
SHEET_EXT = (".xlsx", ".xlsm", ".xls", ".csv")
ACCEPTED = (".pdf", ".docx", ".txt", ".md", *SHEET_EXT, *IMAGE_EXT)


# ================================================================================================ text of a file
def kind_of(name: str) -> str:
    low = name.lower()
    if low.endswith(".pdf"):
        return "pdf"
    if low.endswith(".docx"):
        return "word"
    if low.endswith(SHEET_EXT):
        return "sheet"
    if low.endswith(IMAGE_EXT):
        return "image"
    if low.endswith((".txt", ".md")):
        return "text"
    return ""


def sheet_text(name: str, data: bytes) -> str:
    """Excel (every sheet) or CSV as text: one row per line, cells joined by ' | '."""
    if name.lower().endswith(".csv"):
        raw = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(raw)))
        return "\n".join(" | ".join(c.strip() for c in r if c.strip()) for r in rows if any(c.strip() for c in r))
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[Sheet: {ws.title}]")
        for r in ws.iter_rows(values_only=True):
            cells = []
            for c in r:
                if c is None or str(c).strip() == "":
                    continue
                if isinstance(c, float) and c.is_integer():
                    c = int(c)
                cells.append(f"{c:,}" if isinstance(c, int) and abs(c) >= 10000 else str(c).strip())
            if cells:
                out.append(" | ".join(cells))
        out.append("")
    return "\n".join(out).strip()


def _jpeg(data: bytes, max_side: int = 1800) -> tuple[str, bytes]:
    from PIL import Image
    im = Image.open(io.BytesIO(data))
    im = im.convert("RGB")
    if max(im.size) > max_side:
        r = max_side / max(im.size)
        im = im.resize((int(im.size[0] * r), int(im.size[1] * r)))
    b = io.BytesIO()
    im.save(b, "JPEG", quality=85)
    return "image/jpeg", b.getvalue()


def file_text(name: str, data: bytes) -> dict[str, Any]:
    """{kind: text|image, pages: [str], images: [(media_type, bytes)], method}. Raises ValueError for an unreadable file."""
    k = kind_of(name)
    if k == "pdf":
        import pypdf
        rd = pypdf.PdfReader(io.BytesIO(data))
        pages = []
        for p in rd.pages:
            try:
                t = p.extract_text(extraction_mode="layout") or ""
            except Exception:  # noqa: BLE001
                t = p.extract_text() or ""
            pages.append("\n".join(x.rstrip() for x in t.splitlines() if x.strip()))
        if sum(len(re.sub(r"\W", "", x)) for x in pages) >= 30 * max(len(rd.pages), 1):
            return {"kind": "text", "pages": pages, "images": [], "method": "PDF text"}
        imgs = []
        for p in rd.pages[:MAX_IMAGES]:
            try:
                got = list(p.images)
            except Exception:  # noqa: BLE001
                got = []
            if got:
                imgs.append(_jpeg(max(got, key=lambda i: len(i.data)).data))
        return {"kind": "image", "pages": [], "images": imgs, "method": "scanned PDF (page images)"}
    if k == "word":
        from backend.dealdesk.intake.reader import _docx_lines
        return {"kind": "text", "pages": ["\n".join(_docx_lines(data))], "images": [], "method": "Word"}
    if k == "sheet":
        return {"kind": "text", "pages": [sheet_text(name, data)], "images": [], "method": "Excel" if not name.lower().endswith(".csv") else "CSV"}
    if k == "image":
        return {"kind": "image", "pages": [], "images": [_jpeg(data)], "method": "photo (page image)"}
    if k == "text":
        return {"kind": "text", "pages": [data.decode("utf-8", errors="replace")], "images": [], "method": "text"}
    raise ValueError(f"{name}: upload PDF, Word, Excel, CSV, text or a photo.")


def r16_name(name: str) -> str:
    return f"{name}.txt" if kind_of(name) == "sheet" else name


def r16_rendition(name: str, data: bytes) -> tuple[str, bytes]:
    """What R16's reader is given: the file itself, or a text rendition of a spreadsheet (R16 reads no Excel)."""
    if kind_of(name) == "sheet":
        return f"{name}.txt", sheet_text(name, data).encode()
    return name, data


# ================================================================================================ the rule reader
TITLE_RULES: list[tuple[str, str]] = [
    ("internal_rating", r"RATING SHEET|OBLIGOR RISK RATING|INTERNAL (?:RISK|OBLIGOR) RATING"),
    ("relationship_profile", r"CUSTOMER 360|RELATIONSHIP PROFILE|DEPOSIT AND RELATIONSHIP"),
    ("approval_history", r"CREDIT APPROVAL EXTRACT|COVENANT COMPLIANCE|PREVIOUS APPROVAL"),
    ("group_structure", r"GROUP STRUCTURE|CONNECTED PARTIES|GROUP CHART"),
    ("guarantor_networth", r"NET WORTH STATEMENT|PERSONAL NET[- ]WORTH|GUARANTOR NET WORTH"),
    ("bank_statements", r"ACCOUNT CONDUCT|BANK STATEMENT|STATEMENT OF ACCOUNT"),
    ("simah_company", r"COMMERCIAL CREDIT REPORT|SIMAH (?:COMPANY|COMMERCIAL)"),
    ("simah_owners", r"CONSUMER CREDIT REPORT|SIMAH (?:CONSUMER|INDIVIDUAL|OWNERS?)"),
    ("aecb_company", r"AECB.*(?:COMPANY|COMMERCIAL)"),
    ("aecb_owners", r"AECB.*(?:INDIVIDUAL|CONSUMER|OWNERS?)"),
    ("site_visit", r"SITE VISIT"),
    ("project_list", r"PROJECT LIST|CONTRACT BACKLOG|PROJECTS? AND BACKLOG"),
    ("ageing", r"\bAGE?ING\b"),
    ("request_letter", r"REQUEST LETTER|FACILITY REQUEST|REQUEST FOR (?:RENEWAL|NEW|ENHANCEMENT|CREDIT)"),
    ("cr", r"COMMERCIAL REGISTRATION(?: CERTIFICATE)?$|COMMERCIAL REGISTRATION CERTIFICATE|\bCR EXTRACT\b"),
    ("trade_licence", r"TRADE LICEN[CS]E"),
    ("moa", r"MEMORANDUM OF ASSOCIATION"),
    ("aoa", r"ARTICLES OF ASSOCIATION"),
    ("vat_cert", r"VAT REGISTRATION CERTIFICATE|TAX REGISTRATION CERTIFICATE"),
    ("zakat_cert", r"ZAKAT CERTIFICATE"),
    ("nitaqat", r"NITAQAT|GOSI CERTIFICATE"),
    ("vat_returns", r"VAT RETURNS?\b"),
    ("audited_fs", r"INDEPENDENT AUDITOR|AUDITED FINANCIAL STATEMENTS"),
    ("mgmt_accounts", r"MANAGEMENT ACCOUNTS|UNAUDITED (?:INTERIM )?FINANCIAL"),
    ("board_resolution", r"BOARD RESOLUTION|RESOLUTION TO BORROW|PARTNERS'? RESOLUTION"),
    ("other_bank_letters", r"FACILITY OFFER LETTER|SANCTION LETTER|OFFER LETTER"),
    ("collateral_docs", r"VALUATION REPORT|TITLE DEED|MORTGAGE DEED"),
    ("owner_ids", r"NATIONAL IDENTITY CARD|EMIRATES ID|PASSPORT COPY"),
]

FIG_WORDS = re.compile(
    r"(?i)revenue|profit|ebitda|equity|borrowing|debt|receivable|payable|cash|limit|outstanding|rating|\bpd\b|raroc|deposit|turnover|"
    r"utilisation|utilization|backlog|contract value|net worth|covenant|dscr|ratio|overdue|cheque|excess|share of wallet|lgd|score|"
    r"total|capital|assets|liabilit|margin|exposure|guarantee|advance payment|performance bond|retention|employees|commission|fee")
OBS = re.compile(
    r"(?i)\b(breach|overdue|arbitration|returned|excess|below hurdle|red flag|downgrade|past due|negative|watch|delay(?:ed)?|expired|"
    r"not received|concentration|emphasis of matter|unvalued|not verified|outstanding\b|uncertain|dispute)")
NUMTOK = re.compile(r"\(?-?\d[\d,]*(?:\.\d+)?\)?")
_MON = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def iso_dates(text: str) -> set[str]:
    """Every date printed in a text, as YYYY-MM-DD (2026-04-30, 30 April 2026, 30-Apr-2026, 30/04/2026)."""
    out = set(re.findall(r"\b(20\d\d-\d\d-\d\d)\b", text))
    for d, m, y in re.findall(r"\b(\d{1,2})[ \-]([A-Za-z]{3,9})[ \-,]+(20\d\d)\b", text):
        mm = _MON.get(m[:3].lower())
        if mm:
            try:
                out.add(dt.date(int(y), mm, int(d)).isoformat())
            except ValueError:
                pass
    for d, m, y in re.findall(r"\b(\d{1,2})/(\d{1,2})/(20\d\d)\b", text):
        try:
            out.add(dt.date(int(y), int(m), int(d)).isoformat())
        except ValueError:
            pass
    return out


def numbers(text: str) -> set[float]:
    """Every figure printed in a text (commas and brackets dropped; absolute value)."""
    out: set[float] = set()
    for t in NUMTOK.findall(text or ""):
        s = t.strip("()").replace(",", "")
        if not s or s in ("-", "."):
            continue
        try:
            out.add(abs(round(float(s), 6)))
        except ValueError:
            continue
    return out


def title_ids(name: str, lines: list[str]) -> list[str]:
    """Round A's title rules: a rule matches the file name, one of the first 6 lines (any case) or an UPPER-CASE line anywhere
    (a title or a part title, never a mention inside a sentence)."""
    heads = [Path(name).stem.upper()] + [x.upper() for x in lines[:6]] + [x for x in lines if x.strip() and x.upper() == x and re.search(r"[A-Z]{4}", x)]
    out = []
    for k, rx in TITLE_RULES:
        if any(re.search(rx, h) for h in heads):
            out.append(k)
    if "approval_history" in out and "bank_statements" in out and not re.search(r"(?i)account conduct|statement", Path(name).stem + " ".join(lines[:6])):
        out.remove("bank_statements")
    return out


def _cells(line: str) -> list[str]:
    parts = re.split(r"\s{2,}|\s\|\s|\t", line.strip())
    if len(parts) == 1 and ":" in line:
        a, b = line.split(":", 1)
        parts = [a, b]
    return [p.strip(" :|") for p in parts if p.strip(" :|")]


def line_figures(pages: list[str], limit: int = 30) -> list[dict[str, Any]]:
    """Key figures as printed: a label cell followed by a value cell holding a number (the first amount with thousands
    separators when there is one, so a note number is skipped). Ranked by credit relevance, then document order."""
    got: list[tuple[int, int, dict[str, Any]]] = []
    seen: set[str] = set()
    n = 0
    for pi, page in enumerate(pages, 1):
        for line in page.splitlines():
            cells = _cells(line)
            if len(cells) < 2 or not re.search(r"[A-Za-z]{3}", cells[0]) or len(cells[0]) > 90:
                continue
            vals = [c for c in cells[1:] if re.search(r"\d", c) and len(c) <= 48]
            if not vals:
                continue
            big = [c for c in vals if re.search(r"\d,\d{3}|\d\.\d|%|\bSAR\b|\bAED\b|x\b", c)]
            v = (big or vals)[0]
            lab = re.sub(r"\s+", " ", cells[0])[:80]
            key = f"{lab}|{v}"
            if key in seen or re.fullmatch(r"(?i)page|ref.*|tel.*|fax.*|p\.?o\.? box.*", lab):
                continue
            seen.add(key)
            score = 2 if FIG_WORDS.search(lab) else 0
            score += 1 if big else 0
            got.append((score, n, {"label": lab, "value": v, "page": pi}))
            n += 1
    got.sort(key=lambda x: (-x[0], x[1]))
    return [x[2] for x in sorted(got[:limit], key=lambda x: x[1])]


def _prose(page: str) -> tuple[list[str], list[str]]:
    """(sentences of the running text, table lines). A wrapped line of prose is joined to the next; a line with cells
    (runs of spaces or ' | ') stays a table line."""
    sents: list[str] = []
    table: list[str] = []
    buf = ""
    for raw in page.splitlines():
        ln = raw.strip()
        if not ln:
            continue
        if re.search(r"\S\s{3,}\S|\s\|\s", ln) and not re.search(r"[a-z]{3,}\s{3,}[a-z]{3,}", ln):
            table.append(ln)
            continue
        buf = f"{buf} {ln}".strip()
        if re.search(r"[.;:]$", ln):
            sents.append(buf)
            buf = ""
    if buf:
        sents.append(buf)
    out = []
    for p in sents:
        out += [x.strip() for x in re.split(r"(?<=[.;])\s+(?=[A-Z0-9])", p) if x.strip()]
    return out, table


def line_observations(pages: list[str], limit: int = 6) -> list[str]:
    """Credit-relevant points as printed: sentences with a warning word first, then remarks in table rows."""
    out: list[str] = []
    later: list[str] = []
    for page in pages:
        sents, table = _prose(page)
        for t in sents:
            if 8 <= len(t.split()) and len(t) <= 300 and OBS.search(t) and t not in out:
                out.append(t)
        for t in table:
            cells = [c for c in re.split(r"\s{2,}|\s\|\s", t) if c.strip()]
            words_ = [c for c in cells if re.search(r"[a-z]{3,}.*\s.*[a-z]{3,}", c)]
            if words_ and OBS.search(" ".join(words_)) and len(" ".join(words_).split()) >= 5:
                x = " · ".join(cells)[:220]
                if x not in later:
                    later.append(x)
    return (out + later)[:limit]


LEGAL = re.compile(r"(?i)\b(co\.?|company|llc|l\.l\.c\.?|est\.?|establishment|ltd\.?|limited|plc|corp(?:oration)?|group|holding|pjsc|psc|fze|fzco|w\.?l\.?l\.?)\b")


def company_name(x: Any) -> str:
    """A legal name, or '' when the text is not one (a table header read as a name, a sentence)."""
    t = re.sub(r"\s+", " ", str(x or "")).strip(" .,:;|")
    if not t or "|" in t or len(t) > 80 or len(t.split()) > 9 or not LEGAL.search(t) or not re.match(r"[A-Z]", t):
        return ""
    return t + ("." if re.search(r"(?i)\b(co|est|ltd)$", t) else "")


def name_in_texts(texts: list[str]) -> str:
    """The most printed company name across the documents."""
    seen: dict[str, int] = {}
    for t in texts:
        for m in re.finditer(r"\b([A-Z][A-Za-z&'\-]+(?: [A-Z][A-Za-z&'\-]+){0,5} (?:Co\.|Company|LLC|Est\.|Ltd\.?|Limited|PJSC|PSC|FZE|FZCO))", t):
            seen[m.group(1)] = seen.get(m.group(1), 0) + 1
    return max(seen, key=lambda k: (seen[k], len(k))) if seen else ""


def _fv(doc: dict[str, Any], k: str) -> str:
    x = (doc.get("fields") or {}).get(k)
    return str(x["value"]) if isinstance(x, dict) and x.get("value") not in (None, "") and not isinstance(x.get("value"), list) else ""


def rule_read(r16doc: dict[str, Any], pages: list[str], name: str, profile_name: str = "") -> dict[str, Any]:
    """The rule reader's intake record for one document (R16's classification and fields + Round A's title rules)."""
    lines = [x for p in pages for x in p.splitlines()] or list(r16doc.get("lines") or [])
    if not pages:
        pages = ["\n".join(r16doc.get("lines") or [])]
    code = r16doc.get("doc_type") or "unknown"
    ids = title_ids(name, lines)
    r16_ids = list(CK.R16_TO_CHECK.get(code, ()))
    if not ids:
        ids = r16_ids
    label = CK.LABEL.get(ids[0], "") if ids else ""
    doc_type = r16doc.get("doc_label") if (r16_ids and set(r16_ids) <= set(ids)) else (label or r16doc.get("doc_label") or "Document")
    if code in ("unknown", "unreadable") and not ids:
        doc_type = "Not recognised"
    period = ""
    fy = _fv(r16doc, "fiscal_year")
    if fy:
        period = "FY" + fy[:4]
    elif _fv(r16doc, "period_from") or _fv(r16doc, "period_to"):
        period = f"{_fv(r16doc, 'period_from')} to {_fv(r16doc, 'period_to')}".strip(" to")
    else:
        m = next((re.search(r"(?i)(?:as (?:at|of)|for the (?:year|six months|period) ended|period|date[d]?)\s*:?\s*([0-9]{1,2}[ \-][A-Za-z]{3,9}[ \-][0-9]{2,4}|20\d\d-\d\d-\d\d|[A-Za-z]{3}-?\d{2,4} to [A-Za-z]{3}-?\d{2,4})", x)
                  for x in lines[:40] if re.search(r"(?i)as (?:at|of)|ended|period|date", x)), None)
        period = m.group(1) if m else ""
    expiry = _fv(r16doc, "expiry_date") if re.fullmatch(r"\d{4}-\d\d-\d\d", _fv(r16doc, "expiry_date")) else ""
    if not expiry:
        m = next((re.search(r"(?i)(?:valid until|validity|expiry date)\s*[:|]?\s*(20\d\d-\d\d-\d\d|\d{1,2}[ \-][A-Za-z]{3,9}[ \-]20\d\d)", x) for x in lines
                  if re.search(r"(?i)valid until|validity|expiry date", x)), None)
        if m:
            got = sorted(iso_dates(m.group(1)))
            expiry = got[0] if got else ""
    entity = company_name(_fv(r16doc, "customer_name")) or _fv(r16doc, "full_name")
    if not entity and profile_name and profile_name.lower().split(" co")[0] in " ".join(lines[:30]).lower():
        entity = profile_name
    title = next((x.strip() for x in lines[:12] if x.strip() and x.upper() == x and re.search(r"[A-Z]{4}", x)), lines[0].strip() if lines else "")
    return {"doc_type": doc_type, "checklist_ids": ids, "entity": entity or None, "period_or_date": period or None,
            "expiry_date": expiry or None, "summary": re.sub(r"\s{2,}", " ", title)[:200],
            "key_figures": line_figures(pages), "observations": line_observations(pages), "r16_type": code, "r16_ids": r16_ids}


# ================================================================================================ the model reader
TOOL = {
    "name": "record_document",
    "description": "Record what this document is and what it says, in the intake schema.",
    "input_schema": {
        "type": "object",
        "properties": {
            "doc_type": {"type": "string", "description": "short name, e.g. Audited financial statements FY2024-FY2025"},
            "checklist_ids": {"type": "array", "items": {"type": "string"}, "description": "ids from the checklist that this document satisfies"},
            "entity": {"type": ["string", "null"]},
            "period_or_date": {"type": ["string", "null"]},
            "expiry_date": {"type": ["string", "null"], "description": "YYYY-MM-DD if the document states a validity or expiry date, else null"},
            "summary": {"type": "string", "description": "at most 2 sentences"},
            "key_figures": {"type": "array", "items": {"type": "object", "properties": {
                "label": {"type": "string"}, "value": {"type": "string", "description": "exact figure with currency, unit and period"},
                "page": {"type": "integer"}}, "required": ["label", "value"]}},
            "observations": {"type": "array", "items": {"type": "string"}, "description": "up to 6 credit-relevant points"},
        },
        "required": ["doc_type", "checklist_ids", "summary", "key_figures", "observations"],
    },
}


def intake_prompt(doc_id: str, name: str, market: str, as_of: str, image: bool) -> str:
    ids = "\n".join(f"{k} = {v}" for k, v in CK.ids_for(market))
    return (f"You are the document-intake step of a corporate RM's credit agent at a {market} bank. Read the document "
            f"{'in the attached image(s)' if image else 'below'} (reference {doc_id}, file \"{name}\") and record it with the "
            "record_document tool, in this schema:\n"
            '{"doc_type": "short name, e.g. Audited financial statements FY2024-FY2025",\n'
            ' "checklist_ids": ["ids from the list below that this document satisfies"],\n'
            ' "entity": "legal name or null",\n "period_or_date": "period covered or document date, or null",\n'
            ' "expiry_date": "YYYY-MM-DD if the document states a validity or expiry date, else null",\n'
            ' "summary": "at most 2 sentences",\n'
            ' "key_figures": [{"label": "...", "value": "exact figure with currency, unit and period", "page": 1}],\n'
            ' "observations": ["up to 6 credit-relevant points: weaknesses, anomalies, concentrations, conduct issues, inconsistencies"]}\n'
            f"Return up to 30 key_figures, each with the page it is printed on ([Page n] markers, or the image number). Copy figures "
            f"exactly as printed; never estimate or compute. Today is {as_of}.\nChecklist ids:\n{ids}")


def model_content(doc_id: str, name: str, market: str, as_of: str, ft: dict[str, Any]) -> list[dict[str, Any]]:
    image = ft["kind"] == "image" and bool(ft["images"])
    blocks: list[dict[str, Any]] = [{"type": "text", "text": intake_prompt(doc_id, name, market, as_of, image)}]
    if image:
        for mt, b in ft["images"][:MAX_IMAGES]:
            blocks.append({"type": "image", "source": {"type": "base64", "media_type": mt, "data": base64.b64encode(b).decode()}})
    else:
        text = "".join(f"\n[Page {i}]\n{p}" for i, p in enumerate(ft["pages"], 1))
        blocks.append({"type": "text", "text": f"DOCUMENT {doc_id}:\n{text[:TEXT_LIMIT]}"})
    return blocks


def clean_model(res: dict[str, Any], market: str) -> dict[str, Any]:
    valid = {k for k, _ in CK.ids_for(market)}
    figs = []
    for f in (res.get("key_figures") or [])[:40]:
        if isinstance(f, dict) and str(f.get("value") or "").strip():
            try:
                pg = int(f.get("page")) if f.get("page") not in (None, "") else None
            except (TypeError, ValueError):
                pg = None
            figs.append({"label": str(f.get("label") or "")[:120], "value": str(f["value"])[:80], "page": pg})
    exp = str(res.get("expiry_date") or "")
    return {"doc_type": str(res.get("doc_type") or "Document")[:120],
            "checklist_ids": [str(x) for x in (res.get("checklist_ids") or []) if str(x) in valid],
            "entity": res.get("entity") or None, "period_or_date": res.get("period_or_date") or None,
            "expiry_date": exp if re.fullmatch(r"\d{4}-\d\d-\d\d", exp) else None, "summary": str(res.get("summary") or "")[:400],
            "key_figures": figs, "observations": [str(x)[:300] for x in (res.get("observations") or [])[:8]]}


# ================================================================================================ reconcile
def _in_text(value: str, nums: set[float]) -> bool:
    toks = numbers(value)
    toks = {t for t in toks if not (1900 <= t <= 2100 and float(t).is_integer())}
    return all(t in nums for t in toks)


def _page_of(value: str, pages: list[str]) -> int | None:
    toks = [t for t in numbers(value) if not (1900 <= t <= 2100)]
    if not toks or not pages:
        return None
    for i, p in enumerate(pages, 1):
        pn = numbers(p)
        if all(t in pn for t in toks):
            return i
    return None


def reconcile(rule: dict[str, Any], model: dict[str, Any] | None, pages: list[str]) -> tuple[dict[str, Any], dict[str, Any]]:
    """(the reconciled intake record, the reconciliation note). `pages` is the text the figures are checked against
    (the document's own text; OCR text for photos and scans)."""
    if model is None:
        return dict(rule), {"status": "rule", "label": "Rule reader", "notes": [], "confirmed": 0, "offered": [], "rejected": []}
    text = "\n".join(pages)
    nums = numbers(text)
    dates = iso_dates(text)
    R, M = set(rule.get("checklist_ids") or []), set(model.get("checklist_ids") or [])
    notes: list[str] = []
    if R == M:
        status, ids = "confirmed", sorted(M, key=list(CK.LABEL).index)
    elif not R and M:
        status, ids = "offered", sorted(M, key=list(CK.LABEL).index)
        notes.append("The rule reader could not classify this file; the model read it as " + ", ".join(CK.LABEL[k] for k in ids) + ".")
    elif R < M:
        status, ids = "offered", sorted(M, key=list(CK.LABEL).index)
        notes.append("The model also found: " + ", ".join(CK.LABEL[k] for k in sorted(M - R, key=list(CK.LABEL).index)) + ".")
    elif M < R:
        status, ids = "flagged", sorted(R, key=list(CK.LABEL).index)
        notes.append("The model did not see: " + ", ".join(CK.LABEL[k] for k in sorted(R - M, key=list(CK.LABEL).index)) + " (kept from the rules).")
    else:
        status, ids = "flagged", sorted(M, key=list(CK.LABEL).index)
        notes.append("Readers disagree: rules read " + (", ".join(CK.LABEL[k] for k in sorted(R, key=list(CK.LABEL).index)) or "nothing")
                     + "; the model read " + ", ".join(CK.LABEL[k] for k in ids) + " (the model's reading is used; check it).")
    # the expiry date: printed in the document, else not used
    exp = rule.get("expiry_date")
    me = model.get("expiry_date")
    if me and me != exp:
        if me in dates or not text:
            if exp:
                notes.append(f"Expiry: rules read {exp}, the model {me} (printed; the model's is used).")
                status = "flagged"
            exp = me
        else:
            notes.append(f"Expiry {me} from the model is not printed in the document: not used.")
    # key figures: the model's must be in the text; the rule reader's are copied from it
    figs: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    offered: list[dict[str, Any]] = []
    confirmed = 0
    rule_vals = {tuple(sorted(numbers(f["value"]))): f for f in rule.get("key_figures") or []}
    for f in model.get("key_figures") or []:
        if text and not _in_text(f["value"], nums):
            rejected.append({**f, "why": "not in the document's text"})
            continue
        page = _page_of(f["value"], pages) or f.get("page")
        key = tuple(sorted(numbers(f["value"])))
        if key in rule_vals and key:
            confirmed += 1
            figs.append({**f, "page": page, "source": "both"})
        else:
            offered.append({**f, "page": page})
            figs.append({**f, "page": page, "source": "model"})
    have = {tuple(sorted(numbers(f["value"]))) for f in figs}
    for f in rule.get("key_figures") or []:
        if tuple(sorted(numbers(f["value"]))) not in have and len(figs) < 40:
            figs.append({**f, "source": "rule"})
    if rejected:
        notes.append(f"{len(rejected)} figure{'s' if len(rejected) != 1 else ''} from the model rejected: not in the document.")
    out = {"doc_type": model.get("doc_type") or rule.get("doc_type"), "checklist_ids": ids, "entity": model.get("entity") or rule.get("entity"),
           "period_or_date": model.get("period_or_date") or rule.get("period_or_date"), "expiry_date": exp,
           "summary": model.get("summary") or rule.get("summary"), "key_figures": figs,
           "observations": model.get("observations") or rule.get("observations") or []}
    label = {"confirmed": "Confirmed", "offered": "Model found more", "flagged": "Flagged"}[status]
    return out, {"status": status, "label": label, "notes": notes, "confirmed": confirmed, "offered": offered, "rejected": rejected,
                 "rule_ids": sorted(R), "model_ids": sorted(M)}


# ================================================================================================ market and relationship
def market_of(texts: list[str]) -> str:
    t = " ".join(texts)
    ksa = len(re.findall(r"SIMAH|ZATCA|Zakat|\bSAR\b|Riyadh|Jeddah|Dammam|Nitaqat|Saudi|Kingdom of Saudi", t))
    uae = len(re.findall(r"AECB|Emirates|\bAED\b|Dubai|Abu Dhabi|Sharjah|[Tt]rade [Ll]icen[cs]e|\bTRN\b", t))
    return "UAE" if uae > ksa else "KSA"


def existing_of(bank_existing: bool, ids: set[str], ask: str) -> bool:
    return bool(bank_existing) or bool(ids & {"relationship_profile", "approval_history"}) or bool(re.search(r"(?i)\brenew", ask or ""))
