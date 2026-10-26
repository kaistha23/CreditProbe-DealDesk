"""The 5-page Word report for a What-If scenario (PLAN_R2 6.5).

  1. Cover and executive summary
  2. Scenario definition and assumptions, and the likelihood
  3. ECL impact: the decomposition waterfall and the key figures
  4. Detail: sector, rating, stage migration and top movers
  5. AI interpretation, the model used (with its validation summary) and limitations

Charts are embedded as images (matplotlib, Agg). Styling is the CreditProbe
house style (backend.house_style.docx): a cover block on page 1, the running
header and classification, a footer with the reference and "Page X of Y", brand
headings and house tables. Every figure comes from the run (SAR million).
"""
from __future__ import annotations

import datetime as dt
import io
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

from backend.house_style import ACCENT as BRAND  # noqa: E402
from backend.house_style import MUTED, Meta  # noqa: E402
from backend.house_style import docx as hd  # noqa: E402

from backend.whatif_v5 import likelihood  # noqa: E402
from backend.whatif_v5.chat import commentary as cm  # noqa: E402
from backend.whatif_v5.models import cards, runner  # noqa: E402
from backend.whatif_v5.models.scenario import Scenario  # noqa: E402

UP, DOWN, TOTAL, MODEL = "#A92A25", "#12704B", "#" + BRAND, "#A16207"
BEFORE, AFTER = "#7FA7C9", "#C2410C"
hd.chart_style()


def _table(doc, header: list[str], rows: list[list[str]], widths: list[float] | None = None, right_from: int = 1):
    return hd.table(doc, header, rows, widths=widths, right_from=right_from, gap=False)


def _img(fig) -> io.BytesIO:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


def _waterfall_png(bars: list[dict]) -> io.BytesIO:
    fig, ax = plt.subplots(figsize=(7.2, 3.2))
    running = 0.0
    for i, b in enumerate(bars):
        if b["kind"] == "total":
            bottom, height, colour = 0.0, b["value"], TOTAL
            running = b["value"]
        else:
            v = b["value"]
            bottom = running if v >= 0 else running + v
            height = abs(v)
            colour = MODEL if b["kind"] == "model" else (UP if v >= 0 else DOWN)
            running += v
        ax.bar(i, height, bottom=bottom, color=colour, width=0.65)
        label = cm.money(b["value"]).replace("SAR ", "") if b["kind"] == "total" else f"{b['value']:+,.1f}"
        ax.text(i, bottom + height, label, ha="center", va="bottom", fontsize=6.5)
    levels, run_ = [], 0.0
    for b in bars:
        run_ = b["value"] if b["kind"] == "total" else run_ + b["value"]
        levels.append(run_)
    lo, hi = min(levels), max(levels)
    pad = max((hi - lo) * 0.25, hi * 0.02)
    ax.set_ylim(max(0.0, lo - pad * 2), hi + pad)
    ax.set_xticks(range(len(bars)))
    ax.set_xticklabels([b["label"] for b in bars], rotation=30, ha="right", fontsize=7)
    ax.set_ylabel("ECL, SAR million", fontsize=8)
    ax.tick_params(axis="y", labelsize=7)
    ax.spines[["top", "right"]].set_visible(False)
    return _img(fig)


def _pair_bars(labels: list[str], before: list[float], after: list[float], ylabel: str, horizontal: bool = False) -> io.BytesIO:
    fig, ax = plt.subplots(figsize=(7.2, 3.0 if not horizontal else max(2.4, 0.28 * len(labels) + 0.8)))
    import numpy as np
    x = np.arange(len(labels))
    if horizontal:
        ax.barh(x - 0.2, before, 0.4, color=BEFORE, label="Before")
        ax.barh(x + 0.2, after, 0.4, color=AFTER, label="After")
        ax.set_yticks(x)
        ax.set_yticklabels(labels, fontsize=7)
        ax.invert_yaxis()
        ax.set_xlabel(ylabel, fontsize=8)
    else:
        ax.bar(x - 0.2, before, 0.4, color=BEFORE, label="Before")
        ax.bar(x + 0.2, after, 0.4, color=AFTER, label="After")
        ax.set_xticks(x)
        ax.set_xticklabels(labels, fontsize=7)
        ax.set_ylabel(ylabel, fontsize=8)
    ax.tick_params(labelsize=7)
    ax.legend(fontsize=7, frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    return _img(fig)


def _h(doc, text: str, level: int = 1) -> None:
    hd.heading(doc, text, level)


def _para(doc, text: str, size: float = 10, bold: bool = False, colour: str | None = None, italic: bool = False):
    p = hd.para(doc, text, size=size, bold=bold, italic=italic, colour=colour)
    p.paragraph_format.space_after = Pt(4)
    return p


def _break(doc) -> None:
    hd.page_break(doc)


def _new_doc(meta: Meta, kind: str, subtitle: str, cover: bool | str):
    return hd.document(meta, kind=kind, subtitle=subtitle, cover=cover)


def build(sc: Scenario, method: str, title: str, owner: str, ref: str, generated_by: str = "") -> tuple[bytes, str]:
    result = runner.run(sc, method)
    lab = result["labels"]
    meta = Meta(title or "What-If scenario", book=lab["book"].lower(), period=result["period"],
                subject=f"{result['population']['label']} · {result['method_name']}", generated_by=generated_by,
                reference=f"{ref} · run {result['run_id']}")
    doc = _new_doc(meta, "What-If scenario report", "; ".join(s["restatement"] for s in result["steps"]), "block")
    _write(doc, sc, method, title, owner, result, cover_done=True)
    out = io.BytesIO()
    doc.save(out)
    safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in (title or "scenario"))[:50]
    return out.getvalue(), f"whatif_{sc.domain}_{result['period']}_{method}_{safe}.docx"


def _write(doc, sc: Scenario, method: str, title: str, owner: str, result: dict[str, Any], cover_done: bool = False) -> None:
    lik = likelihood.assess(sc)
    com = cm.commentary(result, lik)
    card = cards.card(sc.domain, method)
    c = result["cards"]
    lab = result["labels"]
    today = dt.date.today().strftime("%d %B %Y")

    # ---- page 1: cover and executive summary
    if not cover_done:
        hd.heading(doc, title or "; ".join(s["restatement"] for s in result["steps"]), 1)
        _para(doc, f"{lab['book']} book · as of {result['period']} · {result['population']['label']} · {result['method_name']}", size=10)
    _para(doc, f"Prepared {today}" + (f" for {owner}" if owner else ""), size=9, colour=MUTED)
    _h(doc, "Executive summary", 1)
    direction = "rises" if c["change"] > 0 else "falls" if c["change"] < 0 else "does not move"
    drivers = sorted([b for b in result["waterfall"] if b["kind"] == "driver" and b.get("active")], key=lambda b: -abs(b["value"]))
    bullets = [
        f"Scenario: {'; '.join(s['restatement'] for s in result['steps'])}.",
        f"ECL {direction} from {cm.money(c['start_ecl'])} to {cm.money(c['scenario_ecl'])} ({c['change_pct']:+.1f}%, {cm.money(c['change'])}).",
        (f"Largest driver: {drivers[0]['label']} ({cm.money(drivers[0]['value'])})." if drivers else "No driver moved ECL."),
        f"Stage moves: {c['moves']['1->2']:,} from Stage 1 to 2, {c['moves']['2->3']:,} from 2 to 3, {c['moves']['1->3']:,} from 1 to 3.",
        f"Coverage {c['coverage_before']:.2f}% → {c['coverage_after']:.2f}%; illustrative CET1 {c['capital_cet1_before_pct']:.1f}% → {c['capital_cet1_after_pct']:.2f}%.",
        f"Likelihood against history: {lik['label']}.",
    ]
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")
    _table(doc, ["Key figure", "Value"], [
        ["Reported ECL", cm.money(c["reported_ecl"])],
        [f"Baseline ECL ({result['method_name']})", cm.money(c["start_ecl"])],
        ["Scenario ECL", cm.money(c["scenario_ecl"])],
        ["Change", f"{cm.money(c['change'])} ({c['change_pct']:+.2f}%)"],
        ["EAD", f"{cm.money(c['ead_before'])} → {cm.money(c['ead_after'])}"],
        [f"{lab['exposures']} · {lab['owners']}", f"{c['exposures']:,} · {c['owners']:,}"],
    ], widths=[8, 8])
    _break(doc)

    # ---- page 2: definition, assumptions, likelihood
    _h(doc, "2. Scenario definition and assumptions", 1)
    _table(doc, ["#", "Step (as understood)", "Shocks", "Scope"],
           [[str(i + 1), s["restatement"], "; ".join(s["shocks"]), s["scope"]] for i, s in enumerate(result["steps"])],
           widths=[0.8, 7.2, 5, 4], right_from=99)
    _para(doc, "")
    _para(doc, "**Method assumptions.** " + " ".join(card["how_it_works"][:3] if method == "delta" else
                                                   [card["summary"]["purpose"], card["training"]["splits"]]), size=9)
    _para(doc, f"**Macroeconomic scenarios.** Base, upturn and downturn paths over five years, weighted "
               f"{int(result['scenario_weights']['base'] * 100)}/{int(result['scenario_weights']['upturn'] * 100)}/"
               f"{int(result['scenario_weights']['downturn'] * 100)}; MEV shocks move every path. Sources: IMF WEO, SAMA, GASTAT, EIA "
               "(see the model card).", size=9)
    _h(doc, "Likelihood against history", 2)
    _para(doc, f"**{lik['label']}.** {lik.get('note', '')}", size=9.5)
    for it in lik.get("items", []):
        doc.add_paragraph(it["text"], style="List Bullet")
    _break(doc)

    # ---- page 3: ECL impact
    _h(doc, "3. ECL impact", 1)
    _para(doc, "The bars go from the reported ECL to the start recomputed by the model (so model error is visible), then each "
               "driver's exact Shapley contribution, to the scenario ECL. They add up exactly.", size=9, colour=MUTED)
    doc.add_picture(_waterfall_png(result["waterfall"]), width=Cm(16.5))
    _table(doc, ["Step", "SAR million"], [[b["label"], (f"{b['value']:,.1f}" if b["kind"] == "total" else f"{b['value']:+,.1f}")]
                                          for b in result["waterfall"]], widths=[10, 6])
    _para(doc, f"Reconciliation: start + drivers − scenario = {result['reconciliation']['difference']:.1e}.", size=8, colour=MUTED)
    _break(doc)

    # ---- page 4: detail
    _h(doc, "4. Detail", 1)
    seg = sorted(result["by_segment"], key=lambda r: -r["ecl_after"])[:10]
    doc.add_picture(_pair_bars([r["segment"] for r in seg], [r["ecl_before"] for r in seg], [r["ecl_after"] for r in seg],
                               f"ECL by {lab['segment'].lower()}, SAR million", horizontal=True), width=Cm(12))
    rat = result["by_rating"]
    doc.add_picture(_pair_bars([r["grade"] for r in rat], [r["ead_before"] for r in rat], [r["ead_after"] for r in rat],
                               f"EAD by {lab['grade'].lower()}, SAR million"), width=Cm(12))
    mig = result["migration"]
    _table(doc, ["From \\ to", "Stage 1", "Stage 2", "Stage 3"],
           [[f"Stage {s}"] + [f"{mig['counts'][i][j]:,} · {cm.money(mig['ead'][i][j])}" for j in range(3)] for i, s in enumerate(mig["stages"])],
           widths=[3, 4.5, 4.5, 4.5])
    _para(doc, "")
    _table(doc, [lab["exposure"], lab["owner"], lab["segment"], "Stage", "ECL before", "ECL after"],
           [[m["exposure_id"], m["owner_name"] or m["owner_id"], m["segment"], f"{m['stage_before']}→{m['stage_after']}",
             cm.money(m["ecl_before"]), cm.money(m["ecl_after"])] for m in result["top_movers"][:8]],
           widths=[2.6, 4, 3.4, 1.5, 2.5, 2.5], right_from=4)
    _break(doc)

    # ---- page 5: interpretation, model, limitations
    _h(doc, "5. Interpretation, model and limitations", 1)
    for para in com["text"].split("\n\n"):
        _para(doc, para, size=9.5)
    _h(doc, f"Model used: {card['name']}", 2)
    val = card.get("validation") or {}
    if method == "xgboost":
        v = card["validation"]["pd"]
        _para(doc, f"PD model test R² {v['random']['test_probability']['r2']:.3f} (random split), "
                   f"{v['out_of_time']['test_probability']['r2']:.3f} (out-of-time); LGD test R² "
                   f"{card['validation']['lgd']['random']['test']['r2']:.3f}. Starting ECL {cm.money(c['start_ecl'])} against "
                   f"{cm.money(c['reported_ecl'])} reported ({c['model_error_pct']:+.2f}%).", size=9)
        val = card.get("engine_validation") or {}
    if val.get("zero_shock"):
        _para(doc, f"Validation: no shock reproduces the start exactly; ECL rises with PD and with notches "
                   f"({'pass' if val['pd_monotonic']['pass'] and val['rating_monotonic']['pass'] else 'see card'}); the waterfall adds up "
                   f"({val['shapley_additivity']['difference']:.1e}); same inputs give the same output.", size=9)
    _h(doc, "Limitations", 2)
    for lim in card["limitations"]:
        doc.add_paragraph(lim, style="List Bullet")


def pack(items: list[tuple[Scenario, str, str, str]], generated_by: str = "") -> bytes:
    """The report pack: a cover page, the list of scenarios, then each scenario's report pages in turn."""
    books = sorted({sc.domain for sc, *_ in items})
    meta = Meta("What-If report pack", book=books[0] if len(books) == 1 else "both", subject=f"{len(items)} scenario(s)",
                generated_by=generated_by)
    doc = _new_doc(meta, "Scenario report pack", f"{len(items)} scenario(s)", True)
    _h(doc, "Scenarios in this pack", 1)
    rows = []
    for sc, method, title, _ in items:
        res = runner.run(sc, method)
        cc = res["cards"]
        rows.append([title, res["labels"]["book"], res["method_name"], cm.money(cc["start_ecl"]), cm.money(cc["scenario_ecl"]),
                     f"{cc['change_pct']:+.1f}%"])
    _table(doc, ["Scenario", "Book", "Model", "ECL before", "ECL after", "Change"], rows, widths=[6, 2, 2.5, 2.5, 2.5, 1.8], right_from=3)
    for sc, method, title, _ref in items:
        _break(doc)
        _write(doc, sc, method, title, "", runner.run(sc, method))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
