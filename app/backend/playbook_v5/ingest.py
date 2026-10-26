"""Read an uploaded .docx, .pdf, .txt or .md into titled sections of text, for Playbook and Policy Docs."""
from __future__ import annotations

import io
import re
from typing import Any

from backend.playbook_v5 import store

KINDS = {".docx": "docx", ".pdf": "pdf", ".txt": "text", ".md": "text"}


def kind_of(name: str) -> str:
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    if ext not in KINDS:
        raise store.Invalid("Upload a .docx, .pdf, .txt or .md file.")
    return KINDS[ext]


def text_of(name: str, data: bytes) -> tuple[str, list[tuple[str, str]]]:
    """(plain text, [(heading, body)]) of a file. Headings come from Word heading styles, or numbered/short
    all-caps lines in a PDF or text file."""
    k = kind_of(name)
    parts: list[tuple[str, str]] = []  # (kind, text) with kind heading|para
    if k == "docx":
        from docx import Document

        try:
            d = Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise store.Invalid("That file could not be read as a Word document.") from exc
        for p in d.paragraphs:
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            parts.append(("heading" if style.startswith("heading") or style == "title" else "para", t))
        for tbl in d.tables:
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(("para", " | ".join(dict.fromkeys(cells))))
    else:
        if k == "pdf":
            from pypdf import PdfReader

            try:
                raw = "\n".join((pg.extract_text() or "") for pg in PdfReader(io.BytesIO(data)).pages)
            except Exception as exc:  # noqa: BLE001
                raise store.Invalid("That file could not be read as a PDF.") from exc
        else:
            raw = data.decode("utf-8", errors="replace")
        for line in raw.splitlines():
            t = line.strip().lstrip("#").strip()
            if not t:
                parts.append(("para", ""))
                continue
            is_head = (line.lstrip().startswith("#") or re.match(r"^(\d+(\.\d+)*\.?|[A-Z]{1,3}-\d+)\s+[A-Z]", t)
                       and len(t) < 90 and not t.endswith(".")) or (t.isupper() and 3 < len(t) < 70)
            parts.append(("heading" if is_head else "para", t))
    text = "\n".join(t for _, t in parts if t)
    sections: list[tuple[str, str]] = []
    head, body = "", []
    for kind, t in parts:
        if kind == "heading":
            if head or any(b.strip() for b in body):
                sections.append((head or "Introduction", _join(body)))
            head, body = t, []
        else:
            body.append(t)
    if head or any(b.strip() for b in body):
        sections.append((head or "Introduction", _join(body)))
    return text, sections


def _join(lines: list[str]) -> str:
    out, cur = [], []
    for ln in lines:
        if not ln.strip():
            if cur:
                out.append(" ".join(cur))
                cur = []
        else:
            cur.append(ln.strip())
    if cur:
        out.append(" ".join(cur))
    return "\n\n".join(out)


def clean(text: str) -> str:
    """Uploaded text must not carry banned labels into a report: replace them rather than refuse the file."""
    return store.BANNED.sub("[label removed]", text)


def to_sections(name: str, data: bytes, who: str) -> list[dict[str, Any]]:
    _, secs = text_of(name, data)
    out = []
    for head, body in secs[:40]:
        blocks = [{"kind": "paragraph", "text": clean(chunk)[: store.MAX_TEXT]}
                  for chunk in [c for c in body.split("\n\n") if c.strip()][: store.MAX_BLOCKS]]
        out.append(store.make_section(clean(head)[:200] or "Section", blocks, who, status="drafted"))
    if not out:
        raise store.Invalid("No text was found in that file.")
    return out
