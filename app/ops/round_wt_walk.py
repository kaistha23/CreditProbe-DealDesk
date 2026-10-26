#!/usr/bin/env python3
"""Rounds W/T: the browser walk of the living credit proposal, on the production build.

    .venv/bin/python ops/round_wt_walk.py [--ui http://127.0.0.1:6600] [--out var/walk/round-wt] [--deal DD-06] [--no-structure]

Signs in (ops/core_login.py), then: opens the deal → Recommendation → "Prepare the credit proposal" (the workspace opens: the
document left, Noor right) → three instructions (shorten, add, underline a key concern), each marked on its section → a figure
typed in text is refused → a number sent to the structure step (re-rendered from the engine) → Accept one change, Undo one →
the version list → a point's evidence → the Provenance view → Mark final → Word downloaded (checked for the provenance annex)
and PDF → the deal's More → Papers and the Deal Desk Recent work card reopen the same workspace. Screenshots and report.json
under --out. Exit 1 on any failure. No password printed.
"""
from __future__ import annotations

import argparse
import io
import json
import sys
import time
from pathlib import Path

V5 = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(V5))
sys.path.insert(0, str(V5 / "ops"))

from core_login import add_session  # noqa: E402

INSTRUCTIONS = [
    ("Shorten the relationship with the bank", None, "relationship"),
    ("Add to the background: The company has supplied the national oil company since its founding.", None, "background"),
    ("Underline the sentence about pricing exception as a key concern", "exceptions", "exceptions"),
]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ui", default="http://127.0.0.1:6600")
    ap.add_argument("--out", default=str(V5 / "var" / "walk" / "round-wt"))
    ap.add_argument("--deal", default="DD-06")
    ap.add_argument("--no-structure", action="store_true")
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    from playwright.sync_api import sync_playwright

    rep: dict = {"ui": a.ui, "deal": a.deal, "at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()), "steps": [], "failures": []}
    errors: list[str] = []

    def step(name: str, ok: bool, **info) -> None:
        rep["steps"].append({"step": name, "ok": ok, **info})
        print(("PASS " if ok else "FAIL ") + name + (f"  {info}" if info else ""))
        if not ok:
            rep["failures"].append(name)

    with sync_playwright() as pw:
        browser = pw.chromium.launch(args=["--disable-dev-shm-usage", "--disable-gpu", "--no-sandbox"])
        ctx = browser.new_context(viewport={"width": 1500, "height": 950}, accept_downloads=True)
        me = add_session(ctx)
        page = ctx.new_page()
        page.on("pageerror", lambda e: errors.append(str(e)[:200]))
        n = 0

        def shot(name: str) -> None:
            nonlocal n
            n += 1
            page.screenshot(path=str(out / f"{n:02d}_{name}.png"))

        try:
            # 1. the deal, its Recommendation step, "Prepare the credit proposal"
            page.goto(f"{a.ui}/deal-desk/{a.deal}?step=recommendation", wait_until="domcontentloaded", timeout=120_000)
            btn = page.locator("[data-testid='dealdesk-prepare-proposal']").first
            btn.wait_for(timeout=180_000)
            shot("deal_recommendation")
            step("open the deal on Recommendation", True)
            btn.click()
            page.wait_for_url(lambda u: "/deal-desk/paper/PW-" in u, timeout=240_000)
            page.wait_for_selector("[data-testid='pw-paper']", timeout=180_000)
            page.wait_for_selector("[data-testid='pw-conversation']", timeout=10_000)
            ws_url = page.url
            rep["workspace"] = ws_url.replace(a.ui, "")
            v0 = page.inner_text("[data-testid='pw-version-chip']")
            shot("workspace_opened")
            step("Prepare the credit proposal opens the workspace (document left, Noor right)", True, url=rep["workspace"], version=v0)

            def instruct(text: str, focus: str | None, timeout: int = 120_000) -> str:
                if focus:
                    page.locator(f"[data-testid='pw-section-{focus}'] h4").first.click()
                before = page.locator("[data-testid='pw-turn-noor']").count()
                page.fill("[data-testid='pw-input']", text)
                page.click("[data-testid='pw-send']")
                page.wait_for_function(f"document.querySelectorAll(\"[data-testid='pw-turn-noor']\").length > {before}", timeout=timeout)
                time.sleep(0.6)
                last = page.locator("[data-testid='pw-turn-noor']").last
                return f"{last.get_attribute('data-kind')}|{last.inner_text()}"

            # 2. three instructions, each marked on its section
            for text, focus, key in INSTRUCTIONS:
                r = instruct(text, focus)
                marked = page.locator(f"[data-testid='pw-section-{key}'] [data-testid='pw-mark']").count() > 0
                page.locator(f"[data-testid='pw-section-{key}']").scroll_into_view_if_needed()
                shot(f"instruction_{key}")
                step(f"instruction: {text[:50]}", r.startswith("edit|") and marked, reply=r[:160], marked=marked)
            # a before/after is visible on a mark
            rel = page.locator("[data-testid='pw-section-relationship'] [data-testid='pw-mark-diff']")
            step("before / after visible on the mark", rel.count() > 0 and page.locator("[data-testid='pw-section-relationship'] [data-testid='pw-mark-removed']").count() > 0)

            # 3. a figure typed in text is refused
            r = instruct('Replace "17.9" with "20" in the action requested', None)
            shot("figure_refused")
            step("a figure in text is refused (figures are the engine's)", r.startswith("rejected|"), reply=r[:160])

            # 4. a number goes to the structure step
            if not a.no_structure:
                r = instruct("Set the amount to SAR 20m", None, timeout=300_000)
                ok = r.startswith("structure|") and page.locator("[data-testid='pw-section-action'] [data-testid='pw-mark']").count() > 0
                page.locator("[data-testid='pw-section-action']").scroll_into_view_if_needed()
                shot("structure_step")
                step("Set the amount: sent to the structure step and re-rendered", ok, reply=r[:220])

            # 5. accept one, undo one
            page.locator("[data-testid='pw-section-relationship'] [data-testid='pw-accept']").click()
            page.wait_for_function("!document.querySelector(\"[data-testid='pw-section-relationship'] [data-testid='pw-mark']\")", timeout=30_000)
            step("Accept clears the mark", True)
            page.locator("[data-testid='pw-section-background'] [data-testid='pw-undo']").click()
            page.wait_for_function("!document.querySelector(\"[data-testid='pw-section-background'] [data-testid='pw-mark']\")", timeout=30_000)
            bg = page.inner_text("[data-testid='pw-section-background']")
            step("Undo reverts the section", "national oil company" not in bg)

            # 6. the version list
            page.click("[data-testid='pw-versions-toggle']")
            page.wait_for_selector("[data-testid='pw-versions']", timeout=10_000)
            nv = page.locator("[data-testid='pw-version']").count()
            shot("versions")
            step("version list (who, when, instruction)", nv >= 5, versions=nv, first=page.locator("[data-testid='pw-version']").first.inner_text()[:160])
            page.click("[data-testid='pw-versions-toggle']")

            # 7. a point's evidence
            pt = page.locator("[data-testid='pw-section-synopsis'] [data-testid='pw-point']").first
            pt.scroll_into_view_if_needed()
            pt.click()
            page.wait_for_selector("[data-testid='pw-evidence-item'], [data-testid='pw-evidence-none']", timeout=60_000)
            ev = page.locator("[data-testid='pw-evidence-item']").count()
            shot("point_evidence")
            step("click a point: its evidence", ev > 0, items=ev, first=page.locator("[data-testid='pw-evidence']").first.inner_text()[:200])

            # 8. the Provenance view
            page.click("[data-testid='pw-view-provenance']")
            page.wait_for_selector("[data-testid='pw-provenance']", timeout=90_000)
            counts = page.inner_text("[data-testid='pw-provenance-counts']")
            page.locator("[data-testid='pw-prov-section-exceptions'] summary").click()
            time.sleep(0.5)
            shot("provenance")
            step("Provenance view of the whole paper", "figures traced" in counts and page.locator("[data-testid='pw-prov-edits']").count() > 0, counts=counts)
            page.click("[data-testid='pw-view-document']")

            # 9. mark final, then Word (with the annex) and PDF
            page.click("[data-testid='pw-mark-final']")
            page.wait_for_selector("[data-testid='pw-final-chip']", timeout=30_000)
            step("Mark final sets the committee copy", True, chip=page.inner_text("[data-testid='pw-final-chip']"))
            with page.expect_download(timeout=180_000) as dl:
                page.click("[data-testid='pw-docx']")
            d = dl.value
            path = out / d.suggested_filename
            d.save_as(str(path))
            from docx import Document
            doc = Document(io.BytesIO(path.read_bytes()))
            text = "\n".join(p.text for p in doc.paragraphs)
            cells = " ".join(c.text for t in doc.tables[-3:] for r in t.rows for c in r.cells)
            ok = "Annex: Provenance" in text and "Underline the sentence" in cells and "dd_" in cells
            step("Word downloaded with the provenance annex", ok, file=path.name, bytes=path.stat().st_size)
            with page.expect_download(timeout=180_000) as dl:
                page.click("[data-testid='pw-pdf']")
            d = dl.value
            ppath = out / d.suggested_filename
            d.save_as(str(ppath))
            step("PDF downloaded", ppath.read_bytes()[:4] == b"%PDF", file=ppath.name, bytes=ppath.stat().st_size)
            shot("final")

            # 10. reopened from the deal's Papers and Recent work
            ws_id = rep["workspace"].split("/")[-1].split("?")[0]
            page.goto(f"{a.ui}/deal-desk/{a.deal}", wait_until="domcontentloaded", timeout=120_000)
            page.click("[data-testid='dealdesk-more-button']", timeout=120_000)
            page.click("[data-testid='dealdesk-action-papers']")
            page.wait_for_selector("[data-testid='pw-deal-paper']", timeout=60_000)
            shot("deal_papers")
            step("the deal's Papers list the workspace", ws_id in page.inner_html("[data-testid='pw-deal-papers']"))
            page.goto(f"{a.ui}/deal-desk", wait_until="domcontentloaded", timeout=120_000)
            card = page.locator(f"[data-testid='dealdesk-paper-card-{ws_id}']")
            card.wait_for(timeout=90_000)
            shot("recent_work")
            card.click()
            page.wait_for_url(lambda u: ws_id in u, timeout=60_000)
            page.wait_for_selector("[data-testid='pw-paper']", timeout=120_000)
            turns = page.locator("[data-testid='pw-turn-user']").count()
            step("Recent work reopens the thread where it was left", turns >= 4, turns=turns)
            shot("reopened")
        except Exception as e:  # noqa: BLE001
            try:
                shot("failure")
            except Exception:  # noqa: BLE001
                pass
            step(f"walk stopped: {type(e).__name__}: {str(e)[:200]}", False)
        rep["page_errors"] = errors[:20]
        rep["user"] = me.get("username")
        browser.close()
    if errors:
        rep["failures"].append(f"page errors: {errors[:3]}")
    (out / "report.json").write_text(json.dumps(rep, indent=1))
    ok = not rep["failures"]
    print(f"{'PASS' if ok else 'FAIL'}: {len(rep['steps'])} steps, {len(rep['failures'])} failures -> {out / 'report.json'}")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
