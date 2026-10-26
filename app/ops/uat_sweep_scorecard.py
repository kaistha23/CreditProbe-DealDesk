"""UAT sweep, Scorecard Validation (PLAN_R3 session S2): choose a scorecard, agree the plan, run it, review, report.

Called from ops/uat_sweep.py (`--only scorecard`). Every figure on screen is checked against the Scorecard
service's API, every button and Back path is clicked, the Word report and Excel export are downloaded and
checked in the house style, and each step ends with no console error, no page error and no failed request.
The Playbook hand-off is clicked end to end when Playbook (port 8506) is running; on a branch without it the
hand-off is covered by tests/v5/scorecard/test_playbook.py against the contract.
"""
from __future__ import annotations

import io
import json
import re
import urllib.error
import urllib.request
from typing import Any, Callable

import core_login

SC = "http://127.0.0.1:8507/api/v1/scorecard"
PB = "http://127.0.0.1:8506/api/v1/playbook"


def tid(t: str) -> str:
    return f'[data-testid="{t}"]'


def get_json(path: str, cookie: str = "") -> Any:
    req = urllib.request.Request(SC + path, headers={"Cookie": f"cp_session={cookie}"} if cookie else {})
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.loads(r.read())


def playbook_up() -> bool:
    try:
        with urllib.request.urlopen(PB + "/health", timeout=3):
            return True
    except Exception:  # noqa: BLE001
        return False


def scorecard_steps(s, page, UI: str, wait_for: Callable[..., Any]) -> None:
    from backend.house_style import check
    from uat_sweep import BANNED  # noqa: F401  (the label rule is applied by the Sweep on every step)

    cookie = core_login.login("anne")["cookie"]
    past = get_json("/validations", cookie)["validations"]
    made: list[str] = []          # the conversations this sweep starts, removed at the end so the UAT library stays as seeded

    # ---------------------------------------------------------------- 1. the home
    s.start("scorecard", "S2-1", "Sidebar → Scorecard Validation: chat box on top, past validations below")
    page.goto(UI + "/", wait_until="domcontentloaded")
    page.locator('nav[aria-label="Primary"]').get_by_role("link", name="Scorecard Validation").click()
    wait_for(page, r".*/scorecard-validation$")
    page.wait_for_selector(tid("sv-past-row"), timeout=90_000)
    s.check("the chat box is on the page", page.locator(tid("sv-ask")).count() == 1 and page.locator(tid("sv-ask-send")).count() == 1)
    box_y = page.locator(tid("sv-chat")).bounding_box()["y"]
    past_y = page.locator(tid("sv-past")).bounding_box()["y"]
    s.check("past validations are below the chat box", past_y > box_y, "below", f"{past_y} vs {box_y}")
    rows = page.locator(tid("sv-past-row"))
    s.check("one past validation per scorecard (seeded)", rows.count() >= 3 and {r["scorecard_id"] for r in past} >= {"RET-APP", "RET-BEH", "SME-ANN"}, ">= 3", rows.count())
    for r in past:
        row = page.locator(f'{tid("sv-past-row")}[data-id="{r["validation_id"]}"]')
        s.check(f"{r['scorecard_id']}: the row shows its verdict ({r['overall']}) and as-of {r['as_of']}",
                row.get_attribute("data-overall") == r["overall"] and r["as_of"] in row.inner_text(), r["overall"], row.inner_text()[:120])
    s.check("the examples are offered", page.locator(tid("sv-example")).count() >= 3)
    s.finish()

    # ---------------------------------------------------------------- 2. a past validation: results, every tab, charts
    sme = next(r for r in past if r["scorecard_id"] == "SME-ANN")
    detail = get_json(f"/validations/{sme['validation_id']}", cookie)["result"]
    s.start("scorecard", "S2-2", "Open the SME validation: every tab, red / amber / green verdicts against stated thresholds, Plotly charts and tables")
    page.locator(f'{tid("sv-past-row")}[data-id="{sme["validation_id"]}"] {tid("sv-past-open")}').click()
    wait_for(page, r".*/scorecard-validation/thread/sv-.*")
    page.wait_for_selector(tid("sv-validation"), timeout=90_000)
    s.check("the thread opens on the validation", page.locator(tid("sv-validation")).get_attribute("data-validation") == sme["validation_id"])
    counts = detail["summary"]["counts"]
    shown = page.locator(tid("sv-counts")).inner_text()
    s.check("the header counts equal the API's", f"{counts['green']} green · {counts['amber']} amber · {counts['red']} red" in shown, f"{counts}", shown)
    s.check("the overall verdict equals the API's", page.locator(tid("sv-overall")).get_attribute("data-status") == detail["summary"]["overall"])
    for sec in detail["sections"]:
        page.click(tid(f"sv-tab-{sec['id']}"))
        page.wait_for_selector(tid(f"sv-panel-{sec['id']}"))
        for t in sec["tests"]:
            row = page.locator(f'{tid("sv-panel-" + sec["id"])} {tid("sv-test-" + t["id"])}')
            s.check(f"{sec['title']}: {t['id']} shows verdict {t['status']}", row.count() == 1 and row.get_attribute("data-status") == t["status"], t["status"],
                    row.get_attribute("data-status") if row.count() else "missing")
            s.check(f"{sec['title']}: {t['id']} states its threshold", t["status"] == "na" or (row.count() == 1 and t["threshold"][:20] in row.inner_text()), "threshold text", row.inner_text()[:80] if row.count() else "")
        if sec["id"] == "variables" and page.locator(tid("sv-woe-toggle")).count():
            page.click(tid("sv-woe-toggle"))
        page.wait_for_function("() => [...document.querySelectorAll('[data-testid^=\"sv-chart-\"]:not([data-testid^=\"sv-chart-card\"])')].every(e => e.getAttribute('data-rendered') === 'true')", timeout=60_000)
        ids = [c["id"] for c in sec["charts"]]
        have = page.locator(f'{tid("sv-panel-" + sec["id"])} [data-testid^="sv-chart-"]:not([data-testid^="sv-chart-card"])').count()
        s.check(f"{sec['title']}: every chart is drawn ({len(ids)})", have == len(ids), len(ids), have)
        s.check(f"{sec['title']}: every Plotly chart has a plot", page.locator(f'{tid("sv-panel-" + sec["id"])} .js-plotly-plot').count() == len(ids), len(ids),
                page.locator(f'{tid("sv-panel-" + sec["id"])} .js-plotly-plot').count())
        for t in sec["tables"]:
            s.check(f"{sec['title']}: table '{t['title'][:40]}' has {len(t['rows'])} rows",
                    page.locator(f'{tid("sv-table-" + t["id"])} tbody tr').count() == len(t["rows"]), len(t["rows"]), page.locator(f'{tid("sv-table-" + t["id"])} tbody tr').count())
    page.click(tid("sv-tab-summary"))
    s.check("the summary lists every test", page.locator('[data-testid^="sv-summary-"]').count() == sum(len(x["tests"]) for x in detail["sections"]))
    s.finish()

    # ---------------------------------------------------------------- 3. downloads and Back
    s.start("scorecard", "S2-3", "Downloads in the house style (Word, Excel), Share through Mail, and Back")
    with page.expect_download() as d:
        page.click(tid("sv-word"))
    docx = d.value.path()
    body = open(docx, "rb").read()
    s.check("the Word report is in the house style", check.docx(body) == [], "no problems", check.docx(body)[:2])
    from docx import Document
    text = "\n".join(p.text for p in Document(io.BytesIO(body)).paragraphs)
    s.check("the report states the scorecard and the as-of", "SME Scorecard" in text and sme["as_of"] in text)
    s.check("the report carries the reviewer's conclusion", ("Not fit for use" in text) == (sme["review"]["conclusion"] == "not_fit"))
    with page.expect_download() as d:
        page.click(tid("sv-excel"))
    xl = open(d.value.path(), "rb").read()
    s.check("the Excel export is in the house style", check.xlsx(xl) == [], "no problems", check.xlsx(xl)[:2])
    page.click(tid("sv-share"))
    page.wait_for_selector(tid("share-dialog"))
    s.check("Share opens the Mail dialog", page.locator(tid("share-dialog")).count() == 1)
    page.keyboard.press("Escape")
    page.wait_for_selector(tid("share-dialog"), state="detached")
    if playbook_up():
        page.click(tid("scorecard-add-report"))
        page.wait_for_selector(tid("add-report-dialog"))
        page.click(tid("add-report-new"))
        page.click(tid("add-report-confirm"))
        page.wait_for_selector(tid("add-report-done") + ", " + tid("add-report-error"), timeout=60_000)
        s.check("Add to report: the section reached Playbook", page.locator(tid("add-report-done")).count() == 1, "added", page.locator(tid("add-report-error")).all_inner_texts())
        page.click(tid("add-report-close"))
    else:
        s.check("Add to report is offered (Playbook is not running on this branch)", page.locator(tid("scorecard-add-report")).count() == 1)
    page.click(tid("sv-back"))
    wait_for(page, r".*/scorecard-validation$")
    page.wait_for_selector(tid("sv-past-row"))
    s.check("Back returns to the home", page.locator(tid("sv-chat")).count() == 1)
    s.finish()

    # ---------------------------------------------------------------- 4. choose, agree the plan, run, review, report
    s.start("scorecard", "S2-4", "From the chat box: choose a scorecard, change the plan, run it, ask about the results")
    page.fill(tid("sv-ask"), "Validate the retail application scorecard")
    page.click(tid("sv-ask-send"))
    wait_for(page, r".*/scorecard-validation/thread/sv-.*")
    page.wait_for_selector(tid("sv-plan"), timeout=90_000)
    thread_id = page.url.rstrip("/").split("/thread/")[1].split("?")[0]
    made.append(thread_id)
    s.check("the assistant proposes a plan with every test group", page.locator('[data-testid^="sv-plan-group-"]').count() == 5)
    s.check("the plan reads the data first (readiness notes)", page.locator(tid("sv-plan-readiness") + " li").count() >= 3)

    def turns() -> int:
        return page.locator(tid("sv-turn")).count()

    def settled(n: int) -> None:
        page.wait_for_function("(n) => document.querySelectorAll('[data-testid=\"sv-turn\"]').length > n && !document.querySelector('[data-testid=\"sv-live-turn\"]')", arg=n, timeout=120_000)
        page.wait_for_load_state("networkidle")

    n = turns()
    page.click(tid("sv-plan-toggle-overrides"))
    settled(n)
    s.check("a toggle changes the plan: overrides are off", page.locator(f'{tid("sv-plan-group-overrides")}').last.get_attribute("data-enabled") == "false")
    n = turns()
    page.fill(tid("sv-composer"), "Cut the segments by region and as of 2025-08")
    page.click(tid("sv-send"))
    settled(n)
    s.check("the plan follows the words: region and 2025-08", page.locator(tid("sv-plan-segment")).last.inner_text() == "region" and page.locator(tid("sv-plan-asof")).last.inner_text() == "2025-08",
            "region / 2025-08", f"{page.locator(tid('sv-plan-segment')).last.inner_text()} / {page.locator(tid('sv-plan-asof')).last.inner_text()}")
    n = turns()
    page.click(tid("sv-plan-run"))
    page.wait_for_selector(tid("sv-working"), timeout=5_000)
    settled(n)
    page.wait_for_selector(tid("sv-validation"), timeout=90_000)
    vid = page.locator(tid("sv-validation")).last.get_attribute("data-validation")
    res = get_json(f"/validations/{vid}", cookie)
    s.check("the validation ran with the agreed plan", res["result"]["plan"]["segment_by"] == "region" and res["result"]["plan"]["as_of"] == "2025-08"
            and "overrides" not in [x["id"] for x in res["result"]["sections"] if x["tests"]], "region / 2025-08 / no overrides")
    shown = page.locator(tid("sv-counts")).last.inner_text()
    c = res["counts"]
    s.check("the verdict counts on screen equal the API's", f"{c['green']} green · {c['amber']} amber · {c['red']} red" in shown, c, shown)
    s.check("the thread is titled after the scorecard", "Retail Application Scorecard" in page.locator(tid("sv-thread-title")).inner_text())
    for q, want in (("What are the findings?", "sv-findings-block"), ("What is PSI?", "sv-answer")):
        n = turns()
        page.fill(tid("sv-composer"), q)
        page.click(tid("sv-send"))
        settled(n)
        s.check(f"'{q}' is answered", page.locator(tid(want)).count() >= 1)
    s.check("the PSI answer defines it from the glossary", "population" in page.locator(tid("sv-turn")).last.inner_text().lower() or "distribution" in page.locator(tid("sv-turn")).last.inner_text().lower())
    s.finish()

    s.start("scorecard", "S2-4b", "Conversation (PLAN_R4 A2): diagnosis in priority order; Gini by segment; PSI for one variable; overrides")
    asks = (("What is wrong with my scorecard?", ("in priority order", "Remedy")),
            ("Show the Gini by product", ("Gini by product",)),
            ("PSI for bureau score", ("CSI",)),
            ("Run the override analysis", ("Overrides",)))
    for q, words in asks:
        n = turns()
        page.fill(tid("sv-composer"), q)
        page.click(tid("sv-send"))
        settled(n)
        last = page.locator(tid("sv-turn")).last
        text = last.inner_text()
        s.check(f"'{q}' answered by the assistant", last.locator(tid("convo-markdown")).count() >= 1
                and all(w in text for w in words), words, text[:160])
    s.check("tables and charts from real queries", page.locator('[data-testid^="convo-table-"]').count() >= 3
            and page.locator('[data-testid^="convo-chart-"]').count() >= 2)
    s.check("the standard flow is unchanged (one validation in the thread)",
            len(get_json(f"/threads/{thread_id}", cookie)["validations"]) == 1)
    s.finish()

    s.start("scorecard", "S2-5", "Review: record the conclusion; the Word report carries it; reopen from the home; Back")
    page.wait_for_selector(tid("sv-review"))
    s.check("the review is pending, with a suggested conclusion", page.locator(tid("sv-review-status")).last.get_attribute("data-status") == "pending")
    page.locator(tid("sv-review-conditions")).last.check()
    page.locator(tid("sv-review-note")).last.fill("Approved for use with a quarterly recalibration.")
    page.locator(tid("sv-review-save")).last.click()
    page.wait_for_function("() => [...document.querySelectorAll('[data-testid=\"sv-review-status\"]')].some(e => e.getAttribute('data-status') === 'reviewed')", timeout=30_000)
    s.check("the review is recorded and shown", "Fit for use with conditions" in page.locator(tid("sv-review-conclusion")).last.inner_text())
    after = get_json(f"/validations/{vid}", cookie)
    s.check("the API holds the review", after["review"]["status"] == "reviewed" and after["review"]["conclusion"] == "conditions" and after["review"]["note"].startswith("Approved"))
    with page.expect_download() as d:
        page.click(tid("sv-word"))
    body = open(d.value.path(), "rb").read()
    text = "\n".join(p.text for p in Document(io.BytesIO(body)).paragraphs)
    s.check("the Word report carries the conclusion and the note", "Fit for use with conditions" in text and "quarterly recalibration" in text)
    s.check("the Word report is in the house style", check.docx(body) == [], "no problems", check.docx(body)[:2])
    page.go_back()
    page.wait_for_selector(tid("sv-past-row"), timeout=60_000)
    s.check("browser Back returns to the home, which lists the new validation", page.locator(f'{tid("sv-past-row")}[data-id="{vid}"]').count() == 1)
    page.locator(f'{tid("sv-past-row")}[data-id="{vid}"]').click()
    wait_for(page, rf".*/scorecard-validation/thread/{thread_id}.*")
    page.wait_for_selector(tid("sv-validation"))
    s.check("the row reopens the thread with its results", page.locator(tid("sv-review-status")).last.get_attribute("data-status") == "reviewed")
    s.finish()

    # ---------------------------------------------------------------- 5. the other two scorecards from the choice cards, links in
    s.start("scorecard", "S2-6", "Choose from the cards; a Mail or Playbook link (?validation=) opens the thread; an unknown thread is refused cleanly")
    page.goto(UI + "/scorecard-validation", wait_until="domcontentloaded")
    page.wait_for_selector(tid("sv-ask"))
    page.fill(tid("sv-ask"), "Which scorecards can you validate?")
    page.click(tid("sv-ask-send"))
    wait_for(page, r".*/scorecard-validation/thread/sv-.*")
    made.append(page.url.rstrip("/").split("/thread/")[1].split("?")[0])
    page.wait_for_selector(tid("sv-choose"), timeout=90_000)
    s.check("three scorecard cards are offered", page.locator('[data-testid^="sv-card-"]').count() == 3)
    n = turns()
    page.click(tid("sv-choose-RET-BEH"))
    settled(n)
    s.check("choosing a card opens that scorecard's plan", page.locator(tid("sv-plan")).count() >= 1 and "RET-BEH" in page.locator(tid("sv-plan")).last.inner_text())
    page.goto(UI + f"/scorecard-validation?validation={sme['validation_id']}", wait_until="domcontentloaded")
    wait_for(page, rf".*/scorecard-validation/thread/{sme['thread_id']}.*")
    page.wait_for_selector(tid("sv-validation"))
    s.check("?validation=<id> opens the validation's thread", page.locator(tid("sv-validation")).get_attribute("data-validation") == sme["validation_id"])
    s.allowed = [(404, "/threads/sv-nothere")]
    page.goto(UI + "/scorecard-validation/thread/sv-nothere", wait_until="domcontentloaded")
    page.wait_for_selector(tid("sv-error"), timeout=60_000)
    s.check("an unknown thread says so, with a way back", "does not exist" in page.locator(tid("sv-error")).inner_text() and page.locator(tid("sv-back")).count() == 1)
    s.allowed = []
    s.finish()
    # ---------------------------------------------------------------- people: personas, ownership, view-as
    from uat_sweep_core import as_user

    def call(method: str, path: str, ck: str, body: dict | None = None) -> tuple[int, Any]:
        req = urllib.request.Request(SC + path, method=method, data=json.dumps(body).encode() if body is not None else None,
                                     headers={"Cookie": f"cp_session={ck}", "Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                return r.status, json.loads(r.read() or b"{}")
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read() or b"{}")

    s.start("scorecard", "S2-7", "People: a persona sees the seeded validations and its own conversations; others do not; view-as is read only")
    browser = page.context.browser
    ctx, pg, _ = as_user(browser, s, "reem.alshehri")
    pg.goto(UI + "/scorecard-validation", wait_until="domcontentloaded")
    pg.wait_for_selector(tid("sv-past-row"), timeout=90_000)
    # PLAN_R5 A1: Reem's access is the retail book (Retail Credit policy), so she sees the retail validations only.
    seeded = sorted(r["validation_id"] for r in past if not str(r.get("scorecard_id", r["validation_id"])).upper().startswith(("SME", "VAL-SME")))
    shown = sorted(el.get_attribute("data-id") for el in pg.locator(tid("sv-past-row")).all())
    s.check("a persona sees the seeded validations of her book", shown == seeded and len(seeded) == 2, seeded, shown)
    pg.fill(tid("sv-ask"), "Validate the retail behavioural scorecard")
    pg.click(tid("sv-ask-send"))
    pg.wait_for_url(re.compile(r".*/scorecard-validation/thread/sv-.*"))
    pg.wait_for_selector(tid("sv-plan"), timeout=90_000)
    reem_thread = pg.url.rstrip("/").split("/thread/")[1].split("?")[0]
    made.append(reem_thread)
    ctx.close()
    khalid, anne = core_login.login("khalid.alotaibi")["cookie"], cookie
    s.check("another persona (not shared, not the team) cannot open it", call("GET", f"/threads/{reem_thread}", khalid)[0] == 404)
    s.check("and does not see it listed", reem_thread not in [t["thread_id"] for t in call("GET", "/threads", khalid)[1]["threads"]])
    s.check("the Administrator sees it", call("GET", f"/threads/{reem_thread}", anne)[0] == 200)
    s.check("the other persona cannot write to it either", call("POST", f"/threads/{reem_thread}/messages", khalid, {"text": "run it"})[0] == 404)
    vcookie = core_login.login("anne")["cookie"]
    core_login.api_post("/api/v1/auth/view-as", {"username": "reem.alshehri"}, vcookie)
    try:
        code, out = call("POST", "/threads", vcookie, {})
        s.check("viewing as someone, starting a conversation is refused as read only", code == 403 and out["detail"] == "Read only while viewing as Reem Al-Shehri.", 403, (code, out))
        s.check("viewing as someone, reading is allowed", call("GET", "/validations", vcookie)[0] == 200)
    finally:
        core_login.api_post("/api/v1/auth/view-as", {"username": None}, vcookie)
    s.finish()
    s.page = page

    for t in made:
        req = urllib.request.Request(SC + f"/threads/{t}", method="DELETE", headers={"Cookie": f"cp_session={cookie}"})
        urllib.request.urlopen(req, timeout=30).read()
    s.start("scorecard", "S2-8", "The sweep leaves the library as seeded: one validation per scorecard")
    after = get_json("/validations", cookie)["validations"]
    s.check("the past validations are the seeded ones again", sorted(r["validation_id"] for r in after) == sorted(r["validation_id"] for r in past), len(past), len(after))
    s.finish()
