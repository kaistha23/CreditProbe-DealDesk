"""R16 d: render one document model to Word (python-docx) and PDF (reportlab), and key pages to photo-style JPGs.

Document model (dict):
  title, subtitle?, issuer, issuer_lines [..], style (gov | bank | audit | letter | form | company | proposal),
  ref (one line under the title), fields [(label, printed value)] (the "Key details" block),
  blocks [ {"h": text} | {"h2": text} | {"p": text} | {"b": [bullets]} | {"n": [numbered items, number in the text]} | {"table": {"cols": [...], "rows": [[...]],
           "widths": [fractions]?, "bold_last": bool}} | {"sig": [(name, role)]} | {"pagebreak": 1} | {"stamp": text} ]
Every page carries the footer FOOTER.
"""
from __future__ import annotations

import io
import pathlib
import random
import subprocess
import tempfile
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageEnhance, ImageFilter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

FOOTER = "Fictional document (CreditProbe R16). All names and figures are invented."
PALETTE = {"gov": "#1f5a3d", "bank": "#7a5a1e", "audit": "#2b3a55", "letter": "#33415c", "form": "#4a4a4a",
           "company": "#264653", "proposal": "#1d3557", "bureau": "#0b5563", "other_bank": "#5b2a86"}


def _esc(s: Any) -> str:
    return (str(s) if s is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ------------------------------------------------------------------------------------------------ PDF
def pdf(doc: dict[str, Any], path: pathlib.Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    color = colors.HexColor(PALETTE.get(doc.get("style", "form"), "#333333"))
    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=9, leading=12)
    small = ParagraphStyle("s", parent=body, fontSize=7.6, leading=9.4)
    title = ParagraphStyle("t", parent=body, fontName="Helvetica-Bold", fontSize=14, leading=18, textColor=color, spaceAfter=2)
    sub = ParagraphStyle("st", parent=body, fontSize=9.5, textColor=colors.HexColor("#444444"))
    h = ParagraphStyle("h", parent=body, fontName="Helvetica-Bold", fontSize=10.5, leading=14, textColor=color, spaceBefore=8, spaceAfter=3)
    h2 = ParagraphStyle("h2", parent=body, fontName="Helvetica-Bold", fontSize=9.5, leading=12, spaceBefore=5, spaceAfter=2)
    right = ParagraphStyle("r", parent=small, alignment=TA_RIGHT)
    story: list[Any] = []
    story.append(Paragraph(_esc(doc["title"]), title))
    if doc.get("subtitle"):
        story.append(Paragraph(_esc(doc["subtitle"]), sub))
    if doc.get("ref"):
        story.append(Paragraph(_esc(doc["ref"]), small))
    story.append(Spacer(1, 6))
    if doc.get("fields"):
        rows = [[Paragraph(f"<b>{_esc(k)}:</b>", small), Paragraph(_esc(v), small)] for k, v in doc["fields"]]
        t = Table(rows, colWidths=[5.2 * cm, 12.3 * cm])
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.6, color), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f3f1ec")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.5),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 1.5), ("LINEBELOW", (0, 0), (-1, -2), 0.25, colors.HexColor("#cccccc"))]))
        story.append(Paragraph("Key details", h))
        story.append(t)
    width = A4[0] - 3.4 * cm
    for b in doc.get("blocks", []):
        if "h" in b:
            story.append(Paragraph(_esc(b["h"]), h))
        elif "h2" in b:
            story.append(Paragraph(_esc(b["h2"]), h2))
        elif "p" in b:
            story.append(Paragraph(b["p"] if b.get("html") else _esc(b["p"]), body))
            story.append(Spacer(1, 3))
        elif "b" in b:
            for x in b["b"]:
                story.append(Paragraph(("&bull; " + (x if b.get("html") else _esc(x))), ParagraphStyle("bl", parent=body, leftIndent=10, firstLineIndent=-7)))
            story.append(Spacer(1, 3))
        elif "n" in b:  # numbered list: the items carry their own number ("1. ...", "a) ..."), no bullet
            for x in b["n"]:
                story.append(Paragraph(x if b.get("html") else _esc(x), ParagraphStyle("nl", parent=body, leftIndent=14, firstLineIndent=-12)))
            story.append(Spacer(1, 3))
        elif "table" in b:
            tb = b["table"]
            cols = tb["cols"]
            n = len(cols)
            widths = [width * f for f in (tb.get("widths") or [1.0 / n] * n)]
            fs = tb.get("font", 7.4 if n > 6 else 8)
            cst = ParagraphStyle("c", parent=small, fontSize=fs, leading=fs + 2)
            cstr = ParagraphStyle("cr", parent=cst, alignment=TA_RIGHT)
            data = [[Paragraph(f"<b>{_esc(c)}</b>", cst) for c in cols]]
            br = set(tb.get("bold_rows", []))
            for ri, r in enumerate(tb["rows"]):
                data.append([Paragraph((f"<b>{_esc(v)}</b>" if ri in br else _esc(v)), cstr if _numeric(v) else cst) for v in r])
            t = Table(data, colWidths=widths, repeatRows=1)
            st = [("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#999999")), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9e6df")),
                  ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 1.2), ("BOTTOMPADDING", (0, 0), (-1, -1), 1.2)]
            if tb.get("bold_last"):
                st.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f3f1ec")))
            t.setStyle(TableStyle(st))
            story.append(t)
            story.append(Spacer(1, 5))
        elif "sig" in b:
            cells = [[Paragraph(f"<br/><br/>______________________<br/><b>{_esc(n)}</b><br/>{_esc(r)}", small) for n, r in b["sig"]]]
            t = Table(cells, colWidths=[width / max(1, len(b["sig"]))] * len(b["sig"]))
            story.append(KeepTogether([Spacer(1, 6), t]))
        elif "stamp" in b:
            t = Table([[Paragraph(f"<b>{_esc(b['stamp'])}</b>", ParagraphStyle("stp", parent=small, textColor=color))]], colWidths=[6 * cm])
            t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 1.2, color), ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
            story.append(Spacer(1, 6))
            story.append(t)
        elif "pagebreak" in b:
            story.append(PageBreak())

    issuer, lines = doc.get("issuer", ""), doc.get("issuer_lines", [])

    def page(c, d):
        c.saveState()
        c.setFillColor(color)
        c.rect(0, A4[1] - 2.2 * cm, A4[0], 2.2 * cm, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 12.5)
        c.drawString(1.7 * cm, A4[1] - 1.15 * cm, issuer)
        c.setFont("Helvetica", 7.5)
        for i, ln in enumerate(lines[:2]):
            c.drawString(1.7 * cm, A4[1] - 1.6 * cm - i * 0.33 * cm, ln)
        c.setFillColor(colors.HexColor("#777777"))
        c.setFont("Helvetica", 6.8)
        c.drawString(1.7 * cm, 0.9 * cm, FOOTER)
        c.drawRightString(A4[0] - 1.7 * cm, 0.9 * cm, f"Page {d.page}")
        c.restoreState()

    sd = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=1.7 * cm, rightMargin=1.7 * cm, topMargin=2.8 * cm, bottomMargin=1.6 * cm,
                           title=doc["title"], author=issuer, subject=doc.get("ref", ""))
    sd.build(story, onFirstPage=page, onLaterPages=page)
    del right


def _numeric(v: Any) -> bool:
    s = str(v).replace(",", "").replace("SAR", "").replace("%", "").rstrip("x").replace("(", "-").replace(")", "").strip()
    try:
        float(s)
        return True
    except ValueError:
        return False


# ------------------------------------------------------------------------------------------------ Word
def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill.lstrip("#"))
    tcPr.append(shd)


def docx(doc: dict[str, Any], path: pathlib.Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    color = PALETTE.get(doc.get("style", "form"), "#333333").lstrip("#")
    rgb = RGBColor.from_string(color.upper())
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.7)
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(9.5)
    # header band: issuer
    hp = sec.header.paragraphs[0]
    r = hp.add_run(doc.get("issuer", ""))
    r.bold, r.font.size, r.font.color.rgb = True, Pt(12), rgb
    for ln in doc.get("issuer_lines", [])[:2]:
        p = sec.header.add_paragraph()
        rr = p.add_run(ln)
        rr.font.size = Pt(7.5)
    fp = sec.footer.paragraphs[0]
    fr = fp.add_run(FOOTER)
    fr.font.size, fr.font.color.rgb = Pt(7), RGBColor(0x77, 0x77, 0x77)
    p = d.add_paragraph()
    r = p.add_run(doc["title"])
    r.bold, r.font.size, r.font.color.rgb = True, Pt(15), rgb
    if doc.get("subtitle"):
        d.add_paragraph(doc["subtitle"])
    if doc.get("ref"):
        q = d.add_paragraph()
        rr = q.add_run(doc["ref"])
        rr.font.size = Pt(8)
    if doc.get("fields"):
        _h(d, "Key details", rgb)
        t = d.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for k, v in doc["fields"]:
            row = t.add_row().cells
            row[0].text, row[1].text = f"{k}:", str(v)
            _shade(row[0], "#f3f1ec")
            for c in row:
                for pp in c.paragraphs:
                    for rr in pp.runs:
                        rr.font.size = Pt(8.5)
            row[0].paragraphs[0].runs[0].bold = True
        _widths(t, [5.2 / 17.5, 12.3 / 17.5])
    for b in doc.get("blocks", []):
        if "h" in b:
            _h(d, b["h"], rgb)
        elif "h2" in b:
            p = d.add_paragraph()
            rr = p.add_run(b["h2"])
            rr.bold = True
        elif "p" in b:
            txt = b["p"]
            if b.get("html"):
                txt = _strip(txt)
            d.add_paragraph(txt)
        elif "b" in b:
            for x in b["b"]:
                d.add_paragraph(_strip(x) if b.get("html") else x, style="List Bullet")
        elif "n" in b:
            for x in b["n"]:
                pn = d.add_paragraph(_strip(x) if b.get("html") else x)
                pn.paragraph_format.left_indent = Cm(0.6)
                pn.paragraph_format.first_line_indent = Cm(-0.5)
                pn.paragraph_format.space_after = Pt(2)
        elif "table" in b:
            tb = b["table"]
            t = d.add_table(rows=1, cols=len(tb["cols"]))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, c in enumerate(tb["cols"]):
                cell = t.rows[0].cells[i]
                cell.text = str(c)
                _shade(cell, "#e9e6df")
                cell.paragraphs[0].runs[0].bold = True
            for rrow in tb["rows"]:
                cells = t.add_row().cells
                for i, v in enumerate(rrow):
                    cells[i].text = str(v)
                    if _numeric(v):
                        cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            n = len(tb["cols"])
            _widths(t, tb.get("widths") or [1.0 / n] * n)
            fs = Pt(7.5 if len(tb["cols"]) > 6 else 8.5)
            for row in t.rows:
                for c in row.cells:
                    for pp in c.paragraphs:
                        pp.paragraph_format.space_after = Pt(0)
                        for rr in pp.runs:
                            rr.font.size = fs
            d.add_paragraph()
        elif "sig" in b:
            t = d.add_table(rows=1, cols=len(b["sig"]))
            for i, (n, role) in enumerate(b["sig"]):
                t.rows[0].cells[i].text = f"\n\n______________________\n{n}\n{role}"
        elif "stamp" in b:
            p = d.add_paragraph()
            rr = p.add_run(f"[ {b['stamp']} ]")
            rr.bold, rr.font.color.rgb = True, rgb
        elif "pagebreak" in b:
            d.add_page_break()
    d.core_properties.title = doc["title"]
    d.core_properties.author = doc.get("issuer", "")
    d.save(str(path))


def _widths(t, fracs, total_cm: float = 17.5) -> None:
    """Column widths that Word and LibreOffice both honour (grid columns + every cell), autofit off."""
    t.autofit = False
    for i, f in enumerate(fracs):
        w = Cm(total_cm * f)
        t.columns[i].width = w
        for cell in t.columns[i].cells:
            cell.width = w


def _h(d, text: str, rgb) -> None:
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run(text)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(11), rgb


def _strip(s: str) -> str:
    import re
    return re.sub(r"<[^>]+>", "", s).replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")


# ------------------------------------------------------------------------------------------------ images
def page_png(pdf_path: pathlib.Path, page: int, dpi: int = 140) -> Image.Image:
    with tempfile.TemporaryDirectory() as td:
        out = pathlib.Path(td) / "p"
        subprocess.run(["pdftoppm", "-r", str(dpi), "-f", str(page), "-l", str(page), "-png", str(pdf_path), str(out)], check=True)
        f = sorted(pathlib.Path(td).glob("p*.png"))[0]
        return Image.open(f).convert("RGB")


def photo(img: Image.Image, seed: str, mode: str = "photo") -> Image.Image:
    """Photo-style: the page on a desk, slightly rotated, with a soft shadow, uneven light and sensor noise.
    Scanner-style: grey cast, slight skew, speckle. Deterministic per seed."""
    rnd = random.Random(seed)
    w, h = img.size
    if mode == "scan":
        angle = rnd.uniform(-1.4, 1.4)
        page = ImageEnhance.Contrast(img).enhance(1.15)
        page = page.rotate(angle, resample=Image.BICUBIC, expand=True, fillcolor=(236, 236, 232))
        page = _noise(page, rnd, 10)
        page = ImageEnhance.Brightness(page).enhance(0.97)
        return page.filter(ImageFilter.GaussianBlur(0.45))
    angle = rnd.uniform(1.2, 3.4) * rnd.choice([-1, 1])
    pad = int(0.09 * w)
    W, H = w + 2 * pad, h + 2 * pad
    bg_top = (rnd.randint(88, 112), rnd.randint(72, 92), rnd.randint(56, 72))  # wooden desk
    bg = Image.new("RGB", (W, H), bg_top)
    dr = ImageDraw.Draw(bg)
    for y in range(0, H, 3):
        shade = int(10 * ((y / H) - 0.5)) + rnd.randint(-4, 4)
        dr.line([(0, y), (W, y)], fill=tuple(max(0, min(255, c + shade)) for c in bg_top))
    # paper slightly warm, uneven light
    paper = ImageEnhance.Color(img).enhance(0.9)
    paper = Image.blend(paper, Image.new("RGB", paper.size, (250, 244, 230)), 0.07)
    grad = Image.linear_gradient("L").resize(paper.size).rotate(rnd.choice([0, 90, 180, 270]), expand=False)
    dark = ImageEnhance.Brightness(paper).enhance(0.82)
    paper = Image.composite(paper, dark, grad.point(lambda v: 150 + v * 105 // 255))
    shadow = Image.new("L", (W, H), 0)
    ImageDraw.Draw(shadow).rectangle([pad + 14, pad + 18, pad + w + 14, pad + h + 18], fill=150)
    shadow = shadow.filter(ImageFilter.GaussianBlur(18))
    bg = Image.composite(Image.new("RGB", (W, H), (25, 20, 15)), bg, shadow)
    bg.paste(paper, (pad, pad))
    bg = bg.rotate(angle, resample=Image.BICUBIC, expand=False, fillcolor=bg_top)
    bg = _noise(bg, rnd, 7).filter(ImageFilter.GaussianBlur(0.6))
    return bg


def _noise(img: Image.Image, rnd: random.Random, amp: int) -> Image.Image:
    import numpy as np
    a = np.asarray(img).astype("int16")
    g = np.random.default_rng(rnd.randint(0, 2**31)).integers(-amp, amp + 1, size=a.shape[:2])[..., None]
    return Image.fromarray(np.clip(a + g, 0, 255).astype("uint8"))


def save_jpg(img: Image.Image, path: pathlib.Path, quality: int = 82) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=quality, optimize=True)
    path.write_bytes(buf.getvalue())
