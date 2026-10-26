"""Write the five regulatory and audit query letters to demo_inputs/queries/ (PLAN_R9 Session F).

    python ops/make_query_letters.py

Each letter is a .docx with a reference line and numbered questions. The Queries service
(backend.queries_v5) reads the questions back out of whatever letter is uploaded; these five are the
prepared cases. Names of the firms are fictional.
"""
from __future__ import annotations

import pathlib

from docx import Document
from docx.shared import Pt

V5 = pathlib.Path(__file__).resolve().parents[1]
OUT = V5 / "demo_inputs" / "queries"

LETTERS = [
    ("Q1_Auditor_Stage1_Past_Due.docx", "Crescent & Partners, Chartered Accountants", "Statutory audit FY2026: interim procedures",
     "REF: CP-AUD-2026-031 (staging of past-due wholesale facilities)",
     ["Provide the list of wholesale facilities more than 30 days past due at the latest quarter-end that remain in Stage 1, "
      "with exposure, days past due and ECL.",
      "For each, provide the rebuttal of the 30 days past due presumption (CBUAE CRM Art 7.8) and the evidence relied upon.",
      "Confirm that each rebuttal was approved by both the CCO and the CRO as required by Art 7.10, and identify any "
      "missing sign-offs.",
      "Quantify the ECL impact had these facilities been transferred to Stage 2."]),
    ("Q2_CBUAE_Stage3_Floors.docx", "Central Bank of the UAE, Banking Supervision Department", "Thematic review of credit-impaired exposures",
     "REF: BSD/CRM/2026/118 (Stage 3 provisioning floors)",
     ["Provide the top 20 Stage 3 wholesale exposures at the latest quarter-end with exposure, ECL and coverage.",
      "Split each exposure between its secured and unsecured portions after haircuts.",
      "State the time each exposure has spent in Stage 3.",
      "Compare the provision held against the minimum provisioning floors of Art 9.23 and quantify any shortfall."]),
    ("Q3_Auditor_Construction_ECL.docx", "Crescent & Partners, Chartered Accountants", "Statutory audit FY2026: ECL movement analysis",
     "REF: CP-AUD-2026-044 (construction Stage 2 ECL)",
     ["Explain the increase in Stage 2 ECL for the construction sector between the last two quarter-ends.",
      "Provide the movement by driver (new business, repayments, stage transfers, PD, LGD, overlays).",
      "Identify the ten counterparties contributing most to the increase."]),
    ("Q4_CBUAE_CRWA_Floor.docx", "Central Bank of the UAE, Banking Supervision Department", "Quarterly prudential return review",
     "REF: BSD/PRV/2026/207 (Stage 1 and 2 provision floor)",
     ["Confirm that Stage 1 and Stage 2 provisions are at least 1.5% of credit risk weighted assets (Art 9.21).",
      "Provide the computation: CRWA, Stage 1 ECL, Stage 2 ECL and the resulting ratio.",
      "If there is a shortfall, state the amount and how it is reflected in the regulatory capital deduction."]),
    ("Q5_Internal_Audit_Restructured.docx", "Internal Audit, Credit Risk Assurance", "Audit of restructured facilities",
     "REF: IA-CR-2026-09 (restructuring, Art 8 and 9.28 to 9.35)",
     ["List the facilities restructured in the last 12 months, with restructuring type (distressed or non-distressed).",
      "Provide their classification and stage before and after restructuring.",
      "Show the probation period status for each and any facility leaving probation early.",
      "Identify any facility restructured three or more times."]),
]


def write(name: str, sender: str, subject: str, ref: str, questions: list[str]) -> pathlib.Path:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph(sender)
    doc.add_paragraph("13 October 2026")
    doc.add_paragraph("The Chief Risk Officer")
    doc.add_paragraph(ref).runs[0].bold = True
    doc.add_paragraph(f"Subject: {subject}")
    doc.add_paragraph("Dear Sir or Madam,")
    doc.add_paragraph("In connection with the above, please provide the following information within ten working days:")
    for i, q in enumerate(questions, 1):
        doc.add_paragraph(f"{i}. {q}")
    doc.add_paragraph("Please provide supporting schedules in Excel.")
    doc.add_paragraph("Yours faithfully,")
    doc.add_paragraph(sender)
    path = OUT / name
    doc.save(path)
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for spec in LETTERS:
        print(write(*spec).relative_to(V5))


if __name__ == "__main__":
    main()
