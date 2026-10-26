"""PLAN_R12 Session G: browser walkthrough of Parameters & policy and the bounded (i) popover, at 1440x900 and 1280x800.

  python ops/r12_params_walkthrough.py [--ui http://127.0.0.1:5500] [--deals DD-01,DD-06] [--shots docs/v5/r12_review/params]

Needs the UI, the Workspace API (identity) and the Deal Desk API running (the Deal Desk API on a scratch DEALDESK_VAR: the
walkthrough uploads and approves the sample policy). Mints a session in the identity store; no password or secret is read
or printed. Writes walkthrough.json and screenshots under --shots.

Checks:
  - Parameters page: Use the sample policy → the review (found / new / not found) → Approve → the (i) on DSCR reads
    "Set by Corporate Credit Policy v<n>, clause 4.2: …"; Policies tab lists it Active; the viewer opens; Excel downloads
    and passes check.xlsx.
  - Every (i) on the Parameters page and on each deal's steps opens a panel that is inside the viewport, opaque and on top
    (the element at its centre and corners belongs to the panel), at both sizes.
  - The RAROC waterfall's value labels are category labels (no text drawn above the bars).
"""
from __future__ import annotations

import argparse
import io
import json
import sys
import time
from pathlib import Path

V5 = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(V5))
sys.path.insert(0, str(V5 / "ops"))
from playwright.sync_api import Page, sync_playwright  # noqa: E402

SIZES = [(1440, 900), (1280, 800)]

PANEL = r"""
() => {
  const ps = [...document.querySelectorAll('[data-testid="dd-info-panel"]')];
  if (!ps.length) return {open: false};
  const p = ps[ps.length - 1];
  const r = p.getBoundingClientRect();
  const vw = Math.min(document.documentElement.clientWidth, document.body.clientWidth, innerWidth), vh = Math.min(document.documentElement.clientHeight, innerHeight);
  const pts = [[r.left + r.width / 2, r.top + Math.min(r.height / 2, 40)], [r.left + 4, r.top + 4], [r.right - 4, r.top + 4],
               [r.left + 4, r.bottom - 4], [r.right - 4, r.bottom - 4]];
  const top = pts.every(([x, y]) => { const e = document.elementFromPoint(x, y); return e && p.contains(e); });
  const bg = getComputedStyle(p).backgroundColor;
  const opaque = !/rgba\(.*,\s*0(\.\d+)?\)$/.test(bg) && bg !== 'transparent';
  return {open: true, left: r.left, top: r.top, right: r.right, bottom: r.bottom, vw, vh,
          inside: r.left >= 0 && r.top >= 0 && r.right <= vw + 0.5 && r.bottom <= vh + 0.5, top_most: top, opaque, bg,
          z: getComputedStyle(p).zIndex, side: p.dataset.side, scroll: p.scrollHeight > p.clientHeight};
}
"""

OVERFLOW = r"""
() => {
  const vw = window.innerWidth, over = [];
  for (const e of document.querySelectorAll('main *')) {
    if (e.offsetParent === null || e.closest('svg')) continue;
    const r = e.getBoundingClientRect();
    if (r.width && r.right > vw + 1) over.push(e.tagName + ' ' + (e.textContent || '').trim().slice(0, 40));
    if (over.length > 8) break;
  }
  const t = document.querySelector('main')?.innerText || '';
  const rough = /TODO|FIXME|lorem|undefined|NaN|\[object Object\]|\bnull\b/.exec(t);
  return {docWidth: document.documentElement.scrollWidth, vw, overflow: over, rough: rough ? rough[0] : null};
}
"""


TICKS = r"""
() => {
  const els = [...document.querySelectorAll('[data-testid="dealdesk-chart-waterfall"] .xtick text')];
  const r = els.map((e) => e.getBoundingClientRect());
  let n = 0;
  for (let i = 0; i < r.length; i++) for (let j = i + 1; j < r.length; j++) {
    const a = r[i], b = r[j];
    if (a.left < b.right - 1 && b.left < a.right - 1 && a.top < b.bottom && b.top < a.bottom) n++;
  }
  return [els.map((e) => e.textContent), n];
}
"""


def mint() -> str:
    from r10_chat_audit import session
    return session()


def idle(page: Page, ms: int = 600) -> None:
    try:
        page.wait_for_load_state("networkidle", timeout=60000)
    except Exception:  # noqa: BLE001
        pass
    page.wait_for_timeout(ms)


def every_info(page: Page, where: str, out: list, limit: int = 80) -> None:
    """Open every visible (i) in turn and check its panel."""
    btns = page.locator('[data-testid="dd-info"], [data-testid="param-info"]')
    n = min(btns.count(), limit)
    for i in range(n):
        b = btns.nth(i)
        try:
            if not b.is_visible():
                continue
            b.scroll_into_view_if_needed(timeout=3000)
            b.click(timeout=3000)
            page.wait_for_timeout(120)
            r = page.evaluate(PANEL)
            r.update({"where": where, "i": i, "label": b.get_attribute("aria-label") or ""})
            out.append(r)
            page.keyboard.press("Escape")
            page.wait_for_timeout(60)
        except Exception as e:  # noqa: BLE001 - a button that moved is reported, never silently skipped
            out.append({"where": where, "i": i, "error": str(e)[:120]})


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ui", default="http://127.0.0.1:5500")
    ap.add_argument("--deals", default="DD-01,DD-06")
    ap.add_argument("--shots", default=str(V5 / "docs" / "v5" / "r12_review" / "params"))
    a = ap.parse_args()
    shots = Path(a.shots)
    shots.mkdir(parents=True, exist_ok=True)
    token = mint()
    report: dict = {"steps": [], "popovers": [], "pages": []}

    def step(name: str, ok: bool, detail: str = "") -> None:
        report["steps"].append({"step": name, "ok": bool(ok), "detail": detail})
        print(("PASS " if ok else "FAIL ") + name + (f": {detail}" if detail else ""))

    with sync_playwright() as pw:
        br = pw.chromium.launch()
        for w, h in SIZES:
            ctx = br.new_context(viewport={"width": w, "height": h}, accept_downloads=True)
            ctx.add_cookies([{"name": "cp_session", "value": token, "domain": "127.0.0.1", "path": "/"}])
            page = ctx.new_page()
            errors: list[str] = []
            page.on("pageerror", lambda e, errs=errors: errs.append(str(e)[:160]))
            t0 = time.time()
            page.goto(f"{a.ui}/deal-desk/parameters")
            page.wait_for_selector('[data-testid="parameters-tab"] [data-testid="param-row"]', timeout=180000)
            idle(page)
            step(f"{w}x{h} parameters page opens", True, f"{time.time() - t0:.1f}s")
            page.screenshot(path=str(shots / f"g01_parameters_{w}.png"))
            report["pages"].append({"page": "parameters", "size": f"{w}x{h}", **page.evaluate(OVERFLOW)})
            if w == 1440:
                page.click('[data-testid="params-sample"]')
                page.wait_for_selector('[data-testid="policy-review"]', timeout=120000)
                idle(page)
                ans = page.locator('[data-testid="policy-review-answer"]').inner_text()
                step("upload: the review answers first", "parameters" in ans and "Nothing changes until you approve" in ans, ans[:120])
                rows = page.locator('[data-testid="policy-review-row"]')
                step("found rows listed", rows.count() >= 10, f"{rows.count()} found")
                page.screenshot(path=str(shots / "g02_review_found.png"), full_page=False)
                page.click('[data-testid="policy-review-filter-new"]')
                page.wait_for_timeout(300)
                page.screenshot(path=str(shots / "g03_review_new.png"))
                page.click('[data-testid="policy-review-filter-not_found"]')
                page.wait_for_timeout(300)
                step("not found rows default to keep", page.locator('[data-testid="policy-review-row"][data-status="not_found"]').count() >= 20)
                page.click('[data-testid="policy-review-filter-found"]')
                every_info(page, f"review {w}", report["popovers"], limit=20)
                page.click('[data-testid="policy-review-approve"]')
                page.wait_for_selector('[data-testid="policy-review-decided"]', timeout=60000)
                idle(page)
                msg = page.locator('[data-testid="params-msg"]').inner_text()
                step("approve creates versions", "approved" in msg and "v2" in msg, msg[:160])
                page.screenshot(path=str(shots / "g04_approved.png"))
                page.click('[data-testid="policy-review-close"]')
                dscr = page.locator('[data-testid="param-row"][data-id="risk_appetite.dscr"]')
                dscr.scroll_into_view_if_needed()
                dscr.locator('[data-testid="param-info"]').click()
                page.wait_for_timeout(200)
                prov = page.locator('[data-testid="dd-info-panel"] [data-testid="param-provenance"]').inner_text()
                step("(i) shows the policy, version and clause", "Set by Corporate Credit Policy v" in prov and "clause 4.2" in prov, prov[:140])
                page.screenshot(path=str(shots / "g05_provenance.png"))
                page.keyboard.press("Escape")
                with page.expect_download(timeout=60000) as dl:
                    page.click('[data-testid="params-excel"]')
                body = Path(dl.value.path()).read_bytes()
                import openpyxl

                from backend.house_style import check
                wb = openpyxl.load_workbook(io.BytesIO(body))
                step("Excel: four sheets, house style", wb.sheetnames[1:] == ["Parameters", "Risk appetite", "Policy register", "Change log"]
                     and check.xlsx(body) == [], ", ".join(wb.sheetnames))
                page.click('[data-testid="params-tab-policies"]')
                page.wait_for_selector('[data-testid="policy-library-row"]', timeout=30000)
                st = page.locator('[data-testid="policy-library-row"] [data-testid="policy-status"]').first.inner_text()
                step("library lists the policy as active", st == "Active", st)
                page.click('[data-testid="policy-view"]')
                page.wait_for_selector('[data-testid="policy-clause"]', timeout=30000)
                step("viewer shows the clauses", page.locator('[data-testid="policy-clause"]').count() >= 40)
                page.screenshot(path=str(shots / "g06_library_viewer.png"))
                with page.expect_download(timeout=60000) as dl2:
                    page.click('[data-testid="policy-docx"]')
                step("Word download", Path(dl2.value.path()).read_bytes()[:2] == b"PK", dl2.value.suggested_filename)
                page.click('[data-testid="params-tab-changes"]')
                page.wait_for_selector('[data-testid="params-change"]', timeout=30000)
                page.screenshot(path=str(shots / "g07_change_log.png"))
                page.click('[data-testid="params-tab-parameters"]')
                page.wait_for_timeout(300)
            every_info(page, f"parameters {w}", report["popovers"])
            for pid in [x for x in a.deals.split(",") if x]:
                for stp in ("proposal", "analysis", "review", "stress", "structure", "recommendation"):
                    page.goto(f"{a.ui}/deal-desk/{pid}?step={stp}")
                    try:
                        page.wait_for_selector('[data-testid="dealdesk-stepper"]', timeout=180000)
                    except Exception:  # noqa: BLE001
                        step(f"{w}x{h} {pid} {stp} opens", False)
                        continue
                    idle(page, 1200)
                    if stp == "analysis":
                        wf = page.locator('[data-testid="dealdesk-chart-waterfall"]')
                        if wf.count():
                            wf.first.scroll_into_view_if_needed()
                            page.wait_for_timeout(600)
                            n_text = page.evaluate("() => document.querySelectorAll('[data-testid=\"dealdesk-chart-waterfall\"] .bartext, [data-testid=\"dealdesk-chart-waterfall\"] g.point text').length")
                            ticks, overlaps = page.evaluate(TICKS)
                            step(f"{w}x{h} {pid} waterfall values under the axis, no label overlaps", n_text == 0 and len(ticks) >= 5 and overlaps == 0,
                                 f"{len(ticks)} labels, {overlaps} overlaps")
                            wf.first.screenshot(path=str(shots / f"g08_waterfall_{pid}_{w}.png"))
                    report["pages"].append({"page": f"{pid} {stp}", "size": f"{w}x{h}", **page.evaluate(OVERFLOW)})
                    every_info(page, f"{pid} {stp} {w}", report["popovers"], limit=40)
            step(f"{w}x{h} no page errors", not errors, "; ".join(errors[:3]))
            ctx.close()
        br.close()
    pops = [p for p in report["popovers"] if p.get("open")]
    bad = [p for p in pops if not (p["inside"] and p["top_most"] and p["opaque"])]
    report["summary"] = {"popovers_checked": len(pops), "popovers_bad": len(bad), "errors": [p for p in report["popovers"] if p.get("error")][:10],
                         "steps_failed": [s for s in report["steps"] if not s["ok"]]}
    print(f"popovers checked {len(pops)}, out of bounds / covered / see-through {len(bad)}")
    for b in bad[:10]:
        print("  BAD", b)
    (shots / "walkthrough.json").write_text(json.dumps(report, indent=1, default=str))
    return 0 if not bad and not report["summary"]["steps_failed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
